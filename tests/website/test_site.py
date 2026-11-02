"""Checks of the website generator, ``python -m ostram site``.

The site is built into temporary folders. The tests compare two builds byte for
byte, read every value the site shows back from the setup manifests and their
authority files, check that every value carries a source entry read from the
repository's records, check every internal link and anchor and every external
host, and build again with registries that leave ``lac`` out, register a fixture
``lac`` setup, or give that fixture a loss set per node.
"""

from __future__ import annotations

import csv
import filecmp
import html
import json
from pathlib import Path
import posixpath
import re
import tempfile
import unittest
from unittest import mock

import yaml
from openpyxl import load_workbook

from ostram.paths import resolve_paths
from ostram.setups import load_manifest, load_setup
from ostram.site.blocks import format_value
from ostram.site.build import build_site
from ostram.site.sources import git_commit

ROOT = Path(__file__).resolve().parents[2]
FIXTURE = Path(__file__).resolve().parent / "fixtures" / "lac" / "setup.yaml"
ALLOWED_SCRIPT_HOSTS = {"cdn.jsdelivr.net"}
REPOSITORY_URL = "https://github.com/Climate-Lead-Group/OSTRAM"
KINDS = {"external", "assumption", "computed", "setting"}
SPAN = re.compile(r'<span class="v" data-v="([^"]+)"(?: data-f="([^"]*)")? data-s="([^"]*)" title="[^"]*">([^<]*)</span>')
FORBIDDEN = re.compile(
    r"(?<![A-Za-z])[A-Za-z]:[\\/](?!/)|[\\/]Users[\\/]|OneDrive|Desktop|\b\w+!\$?[A-Z]{1,3}\$?\d+|\b(?:handover|chat|hub)\b",
    re.IGNORECASE,
)


def _files(folder: Path) -> list[str]:
    return sorted(p.relative_to(folder).as_posix() for p in folder.rglob("*") if p.is_file())


def _pages(folder: Path) -> dict[str, str]:
    return {name: (folder / name).read_text(encoding="utf-8") for name in _files(folder) if name.endswith(".html")}


def _taxonomy(path: Path) -> list[tuple[str, str]]:
    with open(path, encoding="utf-8-sig", newline="") as handle:
        return [(row[0], row[1]) for row in list(csv.reader(handle))[1:] if len(row) > 1 and row[0]]


def _normal(text: str) -> str:
    """Return text without Markdown marks and comment signs, on one line."""

    text = re.sub(r"<(https?://[^>\s]+)>", r"\1", text)
    text = re.sub(r"(?m)^\s*#\s?", "", text)
    return " ".join(text.replace("*", "").replace("`", "").split())


def _record_text(url: str) -> str:
    """Return the normalised text of the repository file a record permalink points to."""

    path = ROOT / url.split("/blob/", 1)[1].split("/", 1)[1]
    if path.suffix == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        cells = [str(cell) for sheet in workbook.worksheets for row in sheet.iter_rows(values_only=True) for cell in row if cell]
        workbook.close()
        return "\n".join(_normal(cell) for cell in cells)
    return _normal(path.read_text(encoding="utf-8"))


def _github_anchors(path: Path) -> set[str]:
    """Return the anchors GitHub gives the headings of a Markdown file."""

    return {
        re.sub(r"[^\w\- ]", "", heading.strip().lower()).replace(" ", "-")
        for heading in re.findall(r"^#{1,6} (.+)$", path.read_text(encoding="utf-8"), re.MULTILINE)
    }


def _build_with(setups: dict[str, str]) -> tuple[dict[str, str], dict]:
    """Build the site with a registry that lists ``setups``; return its pages and values."""

    paths = resolve_paths(project_root=ROOT)
    with tempfile.TemporaryDirectory() as tmp:
        registry = Path(tmp) / "registry.json"
        registry.write_text(json.dumps({"schema": "ostram-setup-registry-v1", "setups": setups}), encoding="utf-8")
        out = Path(tmp) / "site"
        build_site(out, paths=paths, registry_path=registry)
        return _pages(out), json.loads((out / "values.json").read_text(encoding="utf-8"))


def expected_values(setup_id: str, paths, manifest=None) -> dict[str, object]:
    """Read the headline values of one setup from its files, independently of the generator."""

    manifest = manifest or load_setup(setup_id, paths=paths)
    sources = manifest.source_paths(paths, require_exists=False)
    expected: dict[str, object] = {}
    rows = _taxonomy(sources["interconnector_taxonomy"])
    expected["nodes"] = len({tech[6:] for category, tech in rows if category == "TRANSMISSION_PLANNED_NONRENO"})
    expected["interconnectors"] = sum(1 for category, _ in rows if category == "INTERCONNECTORS")
    expected["internal_grid_technologies"] = sum(
        1 for category, _ in rows
        if category.startswith("TRANSMISSION_") and category != "TRANSMISSION_SUM"
    )
    compilation = yaml.safe_load(sources["compilation_config"].read_text(encoding="utf-8"))
    expected["first_year"] = int(compilation["initial_year"])
    expected["last_year"] = int(compilation["final_year"])
    expected["years"] = expected["last_year"] - expected["first_year"] + 1
    text = sources["compilation_config"].read_text(encoding="utf-8")
    timeslices = re.search(r"^(\s*)Timeslices:\n((?:\1- .*\n)+)", text, re.MULTILINE)
    expected["timeslices"] = len(timeslices.group(2).strip().splitlines())
    registry_path = sources["scenario_registry"]
    raw = (
        json.loads(registry_path.read_text(encoding="utf-8"))
        if registry_path.suffix == ".json"
        else yaml.safe_load(registry_path.read_text(encoding="utf-8"))
    )
    registry = raw.get("scenario_registry", raw)
    expected["support_scenarios"] = len(registry["support_scenarios"])
    expected["decision_scenarios"] = len(registry["decision_scenarios"])
    expected["root_scenarios"] = len(registry["root_scenarios"])
    expected["scenarios"] = expected["support_scenarios"] + expected["decision_scenarios"]
    country = yaml.safe_load(sources["country_config"].read_text(encoding="utf-8"))
    for name, value in (country.get("internal_transmission") or {}).items():
        expected[f"internal.{name}"] = value
    return {f"{setup_id}.{key}": value for key, value in expected.items()}


class SiteBuildTest(unittest.TestCase):
    """One build of the repository's site, shared by the checks below."""

    @classmethod
    def setUpClass(cls):
        cls.paths = resolve_paths(project_root=ROOT)
        cls.tmp = tempfile.TemporaryDirectory()
        cls.out = Path(cls.tmp.name) / "first"
        cls.page_list = build_site(cls.out, paths=cls.paths)
        cls.pages = _pages(cls.out)
        cls.values = json.loads((cls.out / "values.json").read_text(encoding="utf-8"))

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_second_build_is_identical(self):
        second = Path(self.tmp.name) / "second"
        build_site(second, paths=self.paths)
        self.assertEqual(_files(self.out), _files(second))
        _, mismatch, errors = filecmp.cmpfiles(self.out, second, _files(self.out), shallow=False)
        self.assertEqual((mismatch, errors), ([], []))

    def test_every_page_of_the_list_is_written(self):
        for path in self.page_list:
            self.assertIn(path, self.pages)
        for path in ("index.html", "glossary.html", "applications/south_asia.html", "applications/lac.html"):
            self.assertIn(path, self.pages)

    def test_values_equal_the_authority_files(self):
        for setup_id in ("south_asia", "unescap_training"):
            for key, value in expected_values(setup_id, self.paths).items():
                with self.subTest(key=key):
                    self.assertIn(key, self.values)
                    self.assertEqual(self.values[key]["value"], value)
        manifest = load_setup("unescap_training", paths=self.paths)
        self.assertEqual(self.values["unescap_training.nodes"]["value"], len(manifest.metadata["country_regions"]))
        self.assertEqual(self.values["unescap_training.timeslices"]["value"], manifest.metadata["timeslices"]["count"])
        self.assertEqual(self.values["unescap_training.years"]["value"], manifest.metadata["year_range"]["count"])

    def test_repository_counts(self):
        self.assertEqual(self.values["repo.tests"]["value"], len(list((ROOT / "tests").rglob("test_*.py"))))
        self.assertEqual(self.values["repo.docs_pages"]["value"], len(list((ROOT / "docs").rglob("*.md"))))
        rules = [p for p in (ROOT / "ostram/pipeline/scenarios/rules").glob("*.py") if p.name != "__init__.py"]
        self.assertEqual(self.values["repo.rule_modules"]["value"], len(rules))
        self.assertEqual(self.values["repo.workflows"]["value"], len(list((ROOT / ".github/workflows").glob("*.yml"))))
        registry = json.loads((ROOT / "config/setups/registry.json").read_text(encoding="utf-8"))
        self.assertEqual(self.values["repo.setups"]["value"], len(registry["setups"]))
        scenarios = ROOT / "ostram/pipeline/scenarios"
        lines = sum(len(p.read_text(encoding="utf-8").splitlines()) for p in scenarios.rglob("*.py"))
        self.assertEqual(self.values["repo.stage.A3.lines"]["value"], lines)

    def test_workbook_and_manifest_values(self):
        manifest = load_setup("south_asia", paths=self.paths)
        raw = yaml.safe_load(manifest.path.read_text(encoding="utf-8"))
        for tech, ceiling in raw["vre_ceilings"]["ceilings_gw"].items():
            self.assertEqual(self.values[f"south_asia.vre_ceiling.{tech}"]["value"], float(ceiling))
        cap = raw["policies"]["link_import_cap"]
        self.assertEqual(self.values["south_asia.link_import_cap.first_year"]["value"], cap["first_year"])
        workbook = load_workbook(
            manifest.source_paths(self.paths, require_exists=False)["scenario_workbook"], read_only=True, data_only=True
        )
        rows = list(workbook["Transmission_Cost_Basis"].iter_rows(values_only=True))
        header = list(rows[0])
        checked = 0
        for row in rows[1:]:
            record = dict(zip(header, row))
            if isinstance(record["Technology"], str) and record["Technology"].startswith("TRN"):
                tech = record["Technology"]
                self.assertEqual(self.values[f"south_asia.corridor.{tech}.capital"]["value"], record["Capital_USD 2023_per_kW"])
                self.assertEqual(self.values[f"south_asia.corridor.{tech}.fixed"]["value"], record["Fixed_USD 2023_per_kW_year"])
                checked += 1
        workbook.close()
        self.assertEqual(checked, self.values["south_asia.interconnectors"]["value"])

    def test_every_value_on_a_page_matches_its_source(self):
        seen = 0
        for name, text in self.pages.items():
            for key, kind, _, shown in SPAN.findall(text):
                with self.subTest(page=name, key=key):
                    self.assertIn(key, self.values)
                    self.assertEqual(html.unescape(shown), format_value(self.values[key]["value"], kind))
                seen += 1
        self.assertGreater(seen, 100)

    def test_internal_links_and_anchors_resolve(self):
        ids = {name: set(re.findall(r'\sid="([^"]+)"', text)) for name, text in self.pages.items()}
        for name, text in self.pages.items():
            duplicates = [i for i in ids[name] if len(re.findall(rf'\sid="{re.escape(i)}"', text)) > 1]
            self.assertEqual(duplicates, [], name)
            for attribute, target in re.findall(r'\s(href|src)="([^"]+)"', text):
                target = html.unescape(target)
                if re.match(r"^[a-z][a-z0-9+.-]*:", target):
                    continue
                path, _, anchor = target.partition("#")
                resolved = posixpath.normpath(posixpath.join(posixpath.dirname(name), path)) if path else name
                with self.subTest(page=name, link=target):
                    self.assertTrue((self.out / resolved).is_file(), f"missing {resolved}")
                    if anchor and resolved.endswith(".html"):
                        self.assertIn(anchor, ids[resolved])

    def test_only_the_allowed_host_is_loaded(self):
        for name, text in self.pages.items():
            loaded = re.findall(r'<(?:script|img|iframe|source)[^>]*\ssrc="(https?://[^"/]+)', text)
            loaded += re.findall(r'<link[^>]*\shref="(https?://[^"/]+)', text)
            loaded += re.findall(r'import [^;]* from "(https?://[^"/]+)', text)
            for origin in loaded:
                self.assertIn(origin.split("//", 1)[1], ALLOWED_SCRIPT_HOSTS, name)
        for name in _files(self.out):
            if name.endswith((".css", ".js")):
                text = (self.out / name).read_text(encoding="utf-8")
                self.assertNotRegex(text, r"@import|url\(\s*['\"]?https?:", name)

    def test_no_em_dash_in_the_site_text(self):
        package = ROOT / "ostram" / "site"
        for path in sorted(package.rglob("*")):
            if path.is_file() and path.suffix in (".html", ".yaml", ".css", ".js", ".py"):
                self.assertNotIn("\u2014", path.read_text(encoding="utf-8"), path.name)

    def test_lac_is_a_real_application(self):
        page = self.pages["applications/lac.html"]
        self.assertNotIn('class="placeholder"', page)
        self.assertIn('data-v="lac.nodes"', page)
        self.assertIn('data-v="lac.', self.pages["index.html"])
        record = f'{REPOSITORY_URL}/blob/{git_commit(ROOT) or "main"}/areas/lac/data_sources.md'
        for name in ("applications/lac.html", "index.html", "paper.html"):
            with self.subTest(page=name):
                self.assertIn(f'href="{record}"', self.pages[name])
                self.assertIn('href="https://www.olade.org/"', self.pages[name])

    def test_bus_bar_note_states_the_model_rule(self):
        note = re.search(r'data-note-for="demand"[^>]*>(.*?)</div>', self.pages["model/node.html"], re.DOTALL).group(1)
        self.assertIn('data-v="south_asia.internal.transmission_loss"', note)
        self.assertRegex(
            note,
            r'href="' + re.escape(REPOSITORY_URL)
            + r'/blob/[^/"]+/docs/rules\.md#internal-transmission-losses-and-bus-bar-demand"',
        )
        self.assertIn("internal-transmission-losses-and-bus-bar-demand", _github_anchors(ROOT / "docs/rules.md"))

    def test_east_west_residual_cites_the_base_year_record(self):
        key = "south_asia.corridor.TRNINDEAINDWE.residual"
        entry = self.values[key]["entry"]
        record = (ROOT / "docs/reference/base_year_value_sources.md").read_text(encoding="utf-8")
        documents = dict(re.findall(r"^- \*\*(R\d+)\*\* .*?<(https?://[^>]+)>", record, re.MULTILINE))
        self.assertEqual(entry["kind"], "external")
        self.assertTrue(entry["record"].endswith("/docs/reference/base_year_value_sources.md"))
        self.assertEqual(
            [url for citation in entry["citations"] for url in citation["urls"]],
            [documents["R100"], documents["R101"], documents["R102"]],
        )
        self.assertEqual([sentence.split(".", 1)[0] for sentence in entry["assumptions"]], ["A14"])
        self.assertIn(f'data-v="{key}"', self.pages["transmission/interconnectors.html"])

    def test_every_printed_value_has_an_entry_and_its_page_lists_it(self):
        for name, text in self.pages.items():
            spans = SPAN.findall(text)
            self.assertEqual(len(spans), text.count('<span class="v"'), name)
            listed = set(re.findall(r'<li id="src-(\d+)">', text))
            self.assertEqual(bool(spans), 'id="sources-on-this-page"' in text, name)
            for key, _, numbers, _ in spans:
                with self.subTest(page=name, key=key):
                    self.assertIn(self.values[key]["entry"]["kind"], KINDS)
                    self.assertTrue(numbers.split())
                    self.assertLessEqual(set(numbers.split()), listed)
        for key, item in self.values.items():
            self.assertIn(item["entry"]["kind"], KINDS, key)

    def test_external_entries_carry_a_url_or_doi(self):
        for key, item in self.values.items():
            entry = item["entry"]
            if entry["kind"] == "external":
                with self.subTest(key=key):
                    urls = [url for citation in entry["citations"] for url in citation["urls"]]
                    self.assertTrue(urls)
                    for url in urls:
                        self.assertRegex(url, r"^https?://[^\s]+$")
        for name, text in self.pages.items():
            for listed in re.findall(r'<ol class="page-sources">(.*?)</ol>', text, re.DOTALL):
                for target in re.findall(r'href="([^"]+)"', listed):
                    self.assertNotRegex(html.unescape(target), r"[.,;:]$", name)

    def test_assumption_text_is_the_repository_sentence(self):
        checked = 0
        for key, item in self.values.items():
            entry = item["entry"]
            for sentence in entry.get("assumptions", []):
                with self.subTest(key=key):
                    record = _record_text(entry["record"])
                    self.assertNotIn("ASSUMPTION", sentence[:12])
                    if re.match(r"A\d+\. ", sentence):
                        # The base-year value record labels its assumptions A1, A2 and so on.
                        self.assertIn(_normal(sentence), record)
                    else:
                        self.assertIn("ASSUMPTION: " + _normal(sentence), record)
                checked += 1
        self.assertGreater(checked, 10)

    def test_computed_entries_name_inputs_with_entries(self):
        computed = [key for key, item in self.values.items() if item["entry"]["kind"] == "computed"]
        self.assertTrue(computed)
        for key in computed:
            for name in self.values[key]["entry"]["inputs"]:
                with self.subTest(key=key, input=name):
                    self.assertIn(name, self.values)
                    self.assertIn(self.values[name]["entry"]["kind"], KINDS)

    def test_entries_hold_no_local_or_internal_text(self):
        for key, item in self.values.items():
            with self.subTest(key=key):
                self.assertIsNone(FORBIDDEN.search(json.dumps(item["entry"], ensure_ascii=False)))
        for name, text in self.pages.items():
            for title in re.findall(r'<span class="v"[^>]* title="([^"]*)"', text):
                self.assertIsNone(FORBIDDEN.search(html.unescape(title).split("\n")[0]), name)

    def test_values_without_a_record_are_left_out(self):
        workbook = load_workbook(ROOT / "inputs/scenarios/OSTRAM_Scenario_Inputs.xlsx", read_only=True, data_only=True)
        rows = list(workbook["Interconnector_Params"].iter_rows(values_only=True))
        workbook.close()
        header = [str(cell) for cell in rows[0]]
        source, reference = header.index("Source"), header.index("IEEE Reference")
        first_year = next(i for i, name in enumerate(header) if re.fullmatch(r"\d{4}", name))
        record = (ROOT / "docs/reference/base_year_value_sources.md").read_text(encoding="utf-8")
        for row in rows[1:]:
            if row[0] == "BAU" and row[5] == "ResidualCapacity":
                key = f"south_asia.corridor.{row[2]}.residual"
                block = re.search(
                    rf"^### `{row[2]}` TotalAnnualMaxCapacity, {header[first_year]}(?: to \d{{4}})?\n\n- Values: ([\d.]+); GW;",
                    record, re.MULTILINE,
                )
                with self.subTest(key=key):
                    if row[source] or row[reference]:
                        self.assertIn(key, self.values)
                    elif block and float(block.group(1)) == float(row[first_year]):
                        self.assertIn(key, self.values)
                        self.assertTrue(self.values[key]["entry"]["record"].endswith("/base_year_value_sources.md"))
                    else:
                        self.assertNotIn(key, self.values)
                        self.assertFalse(any(f'data-v="{key}"' in text for text in self.pages.values()))
        self.assertNotIn("unescap_training.corridor.TRNBGDXXINDEA.residual", self.values)

    def test_repository_links_are_pinned_to_the_built_commit(self):
        commit = git_commit(ROOT)
        if commit is None:
            self.skipTest("the repository has no Git metadata")
        pinned = 0
        for name, text in self.pages.items():
            for target in re.findall(r'href="(' + re.escape(REPOSITORY_URL) + r'/(?:blob|tree)/[^"]+)"', text):
                with self.subTest(page=name, link=target):
                    self.assertTrue(target.startswith(f"{REPOSITORY_URL}/blob/{commit}/"))
                pinned += 1
            self.assertNotIn("Git was not available", text)
        self.assertGreater(pinned, 100)
        for item in self.values.values():
            for url in (item["entry"].get("url"), item["entry"].get("record")):
                if url:
                    self.assertTrue(url.startswith(f"{REPOSITORY_URL}/blob/{commit}/"), url)

    def test_data_page_has_a_source_column(self):
        self.assertIn("<th>Source</th>", self.pages["data.html"])


class NoGitTest(unittest.TestCase):
    """Without Git, repository links fall back to the main branch and the footer says so."""

    def test_links_fall_back_to_main(self):
        paths = resolve_paths(project_root=ROOT)
        with tempfile.TemporaryDirectory() as tmp, mock.patch("ostram.site.sources.git_commit", return_value=None):
            out = Path(tmp) / "site"
            build_site(out, paths=paths)
            pages = _pages(out)
        for name, text in pages.items():
            self.assertEqual(text.count("Git was not available when this site was built"), 1, name)
            for target in re.findall(r'href="(' + re.escape(REPOSITORY_URL) + r'/blob/[^"]+)"', text):
                self.assertTrue(target.startswith(f"{REPOSITORY_URL}/blob/main/"), target)


class LacPlaceholderTest(unittest.TestCase):
    """Without a ``lac`` setup in the registry, the LAC page is the placeholder."""

    def test_lac_is_a_placeholder_without_a_lac_setup(self):
        pages, _ = _build_with({"south_asia": "project:config/setups/south_asia.yaml"})
        page = pages["applications/lac.html"]
        self.assertIn('class="placeholder"', page)
        self.assertNotIn('data-v="lac.', page)
        self.assertNotIn('data-v="lac.', pages["index.html"])


class LossByNodeTest(unittest.TestCase):
    """The loss setting may be one number or a mapping of node to number.

    ASSUMPTION: the fixtures of tests/website/fixtures/lac_loss give the fixture ``lac``
    setup a loss mapping, the same 5% in both nodes or 4% and 6%; they are test values.
    """

    FIXTURES = "project:tests/website/fixtures/lac_loss"
    SOUTH_ASIA = "project:config/setups/south_asia.yaml"

    def test_a_mapping_with_one_number_prints_one_value(self):
        pages, values = _build_with({"south_asia": self.SOUTH_ASIA, "lac": f"{self.FIXTURES}/setup_equal.yaml"})
        self.assertEqual(values["lac.internal.transmission_loss"]["value"], 0.05)
        self.assertEqual(values["lac.internal.transmission_loss"]["entry"]["kind"], "assumption")
        self.assertAlmostEqual(values["lac.internal.efficiency"]["value"], 0.95)
        self.assertEqual(values["lac.internal.efficiency"]["entry"]["inputs"], ["lac.internal.transmission_loss"])
        self.assertFalse([k for k in values if k.startswith(("lac.internal.transmission_loss.", "lac.internal.efficiency."))])
        self.assertIn('data-v="lac.internal.transmission_loss"', pages["transmission/internal-grid.html"])
        self.assertIn('data-v="south_asia.internal.transmission_loss"', pages["index.html"])

    def test_a_mapping_with_different_numbers_prints_a_value_per_node(self):
        pages, values = _build_with({"south_asia": self.SOUTH_ASIA, "lac": f"{self.FIXTURES}/setup_by_node.yaml"})
        grid = pages["transmission/internal-grid.html"]
        self.assertNotIn("lac.internal.transmission_loss", values)
        self.assertNotIn("lac.internal.efficiency", values)
        for node, loss in {"AAAXX": 0.04, "BBBXX": 0.06}.items():
            with self.subTest(node=node):
                key = f"lac.internal.transmission_loss.{node}"
                self.assertEqual(values[key]["value"], loss)
                self.assertEqual(values[key]["entry"]["kind"], "assumption")
                efficiency = values[f"lac.internal.efficiency.{node}"]
                self.assertAlmostEqual(efficiency["value"], 1 - loss)
                self.assertEqual(efficiency["entry"]["kind"], "computed")
                self.assertEqual(efficiency["entry"]["inputs"], [key])
                self.assertIn(f'<code>{node}</code> <span class="v" data-v="{key}" data-f="pct"', grid)
        self.assertNotIn('data-v="lac.internal.transmission_loss"', grid)


class LacFixtureTest(unittest.TestCase):
    """With a ``lac`` setup registered, the same templates fill its pages."""

    def test_lac_pages_fill_with_a_lac_setup(self):
        paths = resolve_paths(project_root=ROOT)
        with tempfile.TemporaryDirectory() as tmp:
            registry = Path(tmp) / "registry.json"
            fixture = FIXTURE.relative_to(ROOT).as_posix()
            registry.write_text(
                json.dumps(
                    {
                        "schema": "ostram-setup-registry-v1",
                        "setups": {
                            "south_asia": "project:config/setups/south_asia.yaml",
                            "lac": f"project:{fixture}",
                        },
                    }
                ),
                encoding="utf-8",
            )
            out = Path(tmp) / "site"
            build_site(out, paths=paths, registry_path=registry)
            pages = _pages(out)
            values = json.loads((out / "values.json").read_text(encoding="utf-8"))
            manifest = load_manifest(FIXTURE, expected_setup="lac")
            for key, value in expected_values("lac", paths, manifest).items():
                with self.subTest(key=key):
                    if key.startswith("lac.internal."):
                        # The fixture's configuration records no source for these: left out.
                        self.assertNotIn(key, values)
                        self.assertFalse(any(f'data-v="{key}"' in text for text in pages.values()))
                    else:
                        self.assertEqual(values[key]["value"], value)
                        self.assertEqual(values[key]["entry"]["kind"], "setting")
            self.assertFalse([key for key in values if key.startswith("lac.internal.")])
            lac = pages["applications/lac.html"]
            self.assertNotIn('class="placeholder"', lac)
            self.assertIn('data-v="lac.nodes"', lac)
            self.assertIn('data-v="lac.nodes"', pages["index.html"])
            self.assertIn('id="families-lac"', pages["transmission/internal-grid.html"])
            self.assertIn('id="rpo-lac"', pages["transmission/repowering.html"])
            self.assertNotIn("applications/unescap_training.html", pages)


if __name__ == "__main__":
    unittest.main()
