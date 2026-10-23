"""Interconnectors sent from a bus-bar node draw (1 - loss) per unit of activity.

A bus-bar requirement equals net generation plus imports minus exports, so it
already holds the losses that export flows cause. The loss stage therefore sets
the input of every interconnector mode whose input fuel is a bus-bar node's
``ELC<node>03`` to 1 - loss, in both AR files and every year; a mode sent from a
node measured at consumption keeps 1.
"""
from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

import yaml
import re

from openpyxl import Workbook, load_workbook

from ostram.pipeline.scenarios.rules import apply_internal_tx_losses as losses


MODEL = Path(__file__).resolve().parents[2] / "model" / "osemosys_fast_preprocessed.txt"
YEARS = (2023, 2024, 2025)
LINK = "TRNBTNXXMDVXX"
PROJ_LINK_ROWS = [
    (1, LINK, "ELCBTNXX03", "Input", 1, 1, 1),
    (2, LINK, "ELCMDVXX03", "Input", 1, 1, 1),
    (1, LINK, "ELCMDVXX04", "Output", 0.95, 0.95, 0.95),
    (2, LINK, "ELCBTNXX04", "Output", 0.95, 0.95, 0.95),
    (1, "PWRHYDBTNXX", "HYDBTNXX", "Input", 1, 1, 1),
]
PROJ_INTERNAL_ROWS = [
    (1, "PWRTRNBTNXX", "ELCBTNXX01", "Input", 1, 1, 1),
    (1, "PWRTRNBTNXX", "ELCBTNXX02", "Output", 1, 1, 1),
    (1, "DSPTRNBTNXX", "ELCBTNXX02", "Input", 1, 1, 1),
    (1, "DSPTRNBTNXX", "ELCBTNXX03", "Output", 1, 1, 1),
]
BASE_LINK_ROWS = [
    (1, "ELCBTNXX03", 1, LINK, "ELCMDVXX04", 1),
    (2, "ELCMDVXX03", 1, LINK, "ELCBTNXX04", 1),
]
BASE_INTERNAL_ROWS = [
    (1, "ELCBTNXX01", 1, "PWRTRNBTNXX", "ELCBTNXX02", 1),
    (1, "ELCBTNXX02", 1, "DSPTRNBTNXX", "ELCBTNXX03", 1),
]


def _sheet(wb, title, header, rows, first=False):
    ws = wb.active if first else wb.create_sheet(title)
    ws.title = title
    ws.append(list(header))
    for row in rows:
        ws.append(list(row))


def _rows(path, sheet):
    ws = load_workbook(path, data_only=True)[sheet]
    return [tuple(r) for r in ws.iter_rows(min_row=2, values_only=True)]


class ExportLossRefundTests(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        root = Path(self._tmp.name)
        self.config = root / "Config_country_codes.yaml"
        self.config.write_text(yaml.safe_dump({"internal_transmission": {"transmission_loss": 0.03}}),
                               encoding="utf-8")
        self.template = root / "template.xlsx"
        wb = Workbook()
        _sheet(wb, "Demand_Projection", ("Demand/Share", "Fuel/Tech", "Name", "Measured_at") + YEARS, [
            ("Demand", "ELCBTNXX03", "x", "bus-bar", 10.0, 11.0, 12.0),
            ("Demand", "ELCMDVXX03", "x", "consumption", 5.0, 6.0, 7.0),
        ], first=True)
        wb.save(self.template)
        self.dir = root / "stage5"
        self.dir.mkdir()
        proj = Workbook()
        _sheet(proj, "Secondary", ("Mode.Operation", "Tech", "Fuel", "Direction") + YEARS,
               PROJ_LINK_ROWS, first=True)
        _sheet(proj, "Demand Techs", ("Mode.Operation", "Tech", "Fuel", "Direction") + YEARS,
               PROJ_INTERNAL_ROWS)
        proj.save(self.dir / losses.PROJ_FILE)
        base = Workbook()
        header = ("Mode.Operation", "Fuel.I", "Value.Fuel.I", "Tech", "Fuel.O", "Value.Fuel.O")
        _sheet(base, "Secondary", header, BASE_LINK_ROWS, first=True)
        _sheet(base, "Demand Techs", header, BASE_INTERNAL_ROWS)
        base.save(self.dir / losses.BASE_FILE)
        demand = Workbook()
        _sheet(demand, "Demand_Projection", ("Demand/Share", "Fuel/Tech", "Name") + tuple(map(str, YEARS)), [
            ("Demand", "ELCBTNXX03", "x", 10.0, 11.0, 12.0),
            ("Demand", "ELCMDVXX03", "x", 5.0, 6.0, 7.0),
        ], first=True)
        demand.save(self.dir / losses.DEMAND_FILE)

    def tearDown(self) -> None:
        self._tmp.cleanup()

    def _run(self):
        return losses.run(self.dir, skip_backup=True, config=self.config, template=self.template)

    def test_bus_bar_sender_input_is_one_minus_loss_in_both_files_every_year(self) -> None:
        log = self._run()
        proj = {(r[0], r[1], r[3]): r[4:] for r in _rows(self.dir / losses.PROJ_FILE, "Secondary")}
        self.assertEqual(proj[(1, LINK, "Input")], (0.97,) * len(YEARS))
        base = {(r[0], r[3]): r[2] for r in _rows(self.dir / losses.BASE_FILE, "Secondary")}
        self.assertEqual(base[(1, LINK)], 0.97)
        self.assertEqual(log["link_projections"]["cells"], len(YEARS))
        self.assertEqual(log["link_projections"]["edited"],
                         [{"tech": LINK, "mode": 1, "fuel": "ELCBTNXX03", "cells": len(YEARS)}])
        self.assertEqual(log["link_base_year"]["cells"], 1)

    def test_consumption_sender_input_stays_one(self) -> None:
        self._run()
        proj = {(r[0], r[1], r[3]): r[4:] for r in _rows(self.dir / losses.PROJ_FILE, "Secondary")}
        self.assertEqual(proj[(2, LINK, "Input")], (1,) * len(YEARS))
        base = {(r[0], r[3]): r[2] for r in _rows(self.dir / losses.BASE_FILE, "Secondary")}
        self.assertEqual(base[(2, LINK)], 1)

    def test_outputs_internal_families_and_other_inputs_unchanged(self) -> None:
        self._run()
        proj = _rows(self.dir / losses.PROJ_FILE, "Secondary")
        self.assertEqual([r for r in proj if r[3] == "Output"],
                         [r for r in PROJ_LINK_ROWS if r[3] == "Output"])
        self.assertIn(PROJ_LINK_ROWS[-1], proj)
        internal = {(r[1], r[3]): r[4:] for r in _rows(self.dir / losses.PROJ_FILE, "Demand Techs")}
        self.assertEqual(internal[("PWRTRNBTNXX", "Input")], (1,) * len(YEARS))
        self.assertEqual(internal[("PWRTRNBTNXX", "Output")], (0.97,) * len(YEARS))
        self.assertEqual(internal[("DSPTRNBTNXX", "Input")], (1,) * len(YEARS))
        self.assertEqual(internal[("DSPTRNBTNXX", "Output")], (1,) * len(YEARS))
        base = {(r[0], r[3]): r[5] for r in _rows(self.dir / losses.BASE_FILE, "Secondary")}
        self.assertEqual(base, {(1, LINK): 1, (2, LINK): 1})
        base_internal = {r[3]: (r[2], r[5]) for r in _rows(self.dir / losses.BASE_FILE, "Demand Techs")}
        self.assertEqual(base_internal, {"PWRTRNBTNXX": (1, 0.97), "DSPTRNBTNXX": (1, 1)})

    def test_switched_off_mode_stays_zero_and_other_values_raise(self) -> None:
        wb = load_workbook(self.dir / losses.PROJ_FILE)
        ws = wb["Secondary"]
        ws.cell(row=2, column=5, value=0)
        wb.save(self.dir / losses.PROJ_FILE)
        self._run()
        proj = {(r[0], r[1], r[3]): r[4:] for r in _rows(self.dir / losses.PROJ_FILE, "Secondary")}
        self.assertEqual(proj[(1, LINK, "Input")], (0, 0.97, 0.97))
        wb = load_workbook(self.dir / losses.PROJ_FILE)
        wb["Secondary"].cell(row=2, column=5, value=0.5)
        wb.save(self.dir / losses.PROJ_FILE)
        with self.assertRaises(ValueError) as ctx:
            self._run()
        self.assertIn(LINK, str(ctx.exception))

    def test_second_run_changes_nothing(self) -> None:
        self._run()
        before = {name: {s: _rows(self.dir / name, s) for s in ("Secondary", "Demand Techs")}
                  for name in (losses.PROJ_FILE, losses.BASE_FILE)}
        log = self._run()
        for key in ("projections", "base_year", "demand", "link_projections", "link_base_year"):
            self.assertEqual(log[key]["cells"], 0, key)
        after = {name: {s: _rows(self.dir / name, s) for s in ("Secondary", "Demand Techs")}
                 for name in (losses.PROJ_FILE, losses.BASE_FILE)}
        self.assertEqual(before, after)

    def test_restore_undoes_the_refund(self) -> None:
        backup = losses.make_backup(self.dir)
        self._run()
        losses.restore(self.dir, backup)
        proj = {(r[0], r[1], r[3]): r[4:] for r in _rows(self.dir / losses.PROJ_FILE, "Secondary")}
        self.assertEqual(proj[(1, LINK, "Input")], (1,) * len(YEARS))

    def test_self_test_passes(self) -> None:
        self.assertEqual(losses.run_self_test(), 0)


class ExportBoundTests(unittest.TestCase):
    def test_export_bounds_bound_link_activity(self) -> None:
        lines = [l for l in MODEL.read_text(encoding="utf-8").splitlines()
                 if re.match(r"s\.t\. D4_I1_TRADE_\d+_(BTN|NPL)_IND_ALL_EXPORT_", l)]
        self.assertEqual(len(lines), 12)
        for line in lines:
            self.assertNotIn("ActivityRatio", line)
            self.assertIn("RateOfActivity", line)


if __name__ == "__main__":
    unittest.main()
