"""Guard: no setup ships stored rule output in its Restrictions sheet.

Rows a scenario's rules generate are written only to disposable run states.
A scenario that inherits Restrictions (Control column
``inherit_restrictions_from``) must declare a ``restrictions`` dependency on
each source in its registry, so the materializer runs the source first and
hands its rows over from the same run.
"""

from __future__ import annotations

import unittest
from pathlib import Path

from openpyxl import load_workbook

from ostram.paths import resolve_paths
from ostram.pipeline.scenarios.registry import load_registry
from ostram.pipeline.scenarios.transformations.scenario_workbooks import (
    read_control_sheet,
)


REPO_ROOT = Path(__file__).resolve().parents[2]
SOUTH_ASIA = resolve_paths(project_root=REPO_ROOT, environ={})
TRAINING = REPO_ROOT / "examples" / "unescap_training"
SETUPS = {
    "south_asia": (SOUTH_ASIA.scenario_registry, SOUTH_ASIA.scenario_workbook),
    "unescap_training": (
        TRAINING / "config" / "scenarios" / "registry.json",
        TRAINING / "inputs" / "scenarios" / "OSTRAM_Scenario_Inputs.xlsx",
    ),
}


class NoStoredRuleOutputTests(unittest.TestCase):
    def test_restrictions_sheet_has_no_data_rows(self) -> None:
        for setup, (_, workbook_path) in SETUPS.items():
            with self.subTest(setup=setup):
                workbook = load_workbook(workbook_path, read_only=True, data_only=True)
                try:
                    data_rows = [
                        row
                        for row in workbook["Restrictions"].iter_rows(
                            min_row=2, values_only=True
                        )
                        if any(value is not None for value in row)
                    ]
                finally:
                    workbook.close()
                self.assertEqual(data_rows, [])

    def test_every_inheritance_is_a_declared_restrictions_dependency(self) -> None:
        for setup, (registry_path, workbook_path) in SETUPS.items():
            registry = load_registry(registry_path)
            roots = registry.roots_by_name
            inheriting = [
                config
                for config in read_control_sheet(workbook_path)
                if config.inherit_restrictions_from
            ]
            # Both setups inherit BAU's rows into A; an empty list would make
            # the check below pass without checking anything.
            self.assertTrue(inheriting, setup)
            for config in inheriting:
                with self.subTest(setup=setup, scenario=config.scenario):
                    self.assertIn(config.scenario, roots)
                    declared = roots[config.scenario].restriction_dependencies
                    for source in config.inherit_restrictions_from:
                        self.assertIn(source, declared)


if __name__ == "__main__":
    unittest.main()
