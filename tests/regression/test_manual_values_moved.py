"""The values the A3 stage 1 manual fixes wrote are read where their kind is read.

Cloned generators draw the fuel of their own country and family
(``clone_fuel_code`` in ``update_ao_from_extensions.py``), the names of the
cloned technologies come from ``Override_Tech.Name_AO`` of the
``AO_Extension_Decisions`` sheet, and the 0.95 output activity ratio of the two
Bangladesh links comes from ``Fixed_Horizon_Parameters``.
"""
import ast
from pathlib import Path
import unittest

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "ostram/pipeline/scenarios/transformations/update_ao_from_extensions.py"
WORKBOOK = ROOT / "inputs/scenarios/OSTRAM_Scenario_Inputs.xlsx"
namespace: dict = {}
tree = ast.parse(SOURCE.read_text(encoding="utf-8"))
exec(compile(ast.Module(body=[n for n in tree.body if isinstance(n, ast.FunctionDef)
                              and n.name == "clone_fuel_code"], type_ignores=[]),
             str(SOURCE), "exec"), namespace)
clone_fuel_code = namespace["clone_fuel_code"]

PRODUCED = {"GASBGD", "GASIND", "GASMDV", "GASINT", "OILBGD", "OILMDV", "OILNPL", "OILINT",
            "GEOINDNO", "HYDINDNO", "HYDINDEA"}
NAMES = {
    "PWRNGSBGDXX": "Natural Gas (Power generator) Bangladesh, region XX",
    "PWRNGSMDVXX": "Natural Gas (Power generator) Maldives, region XX",
    "PWROILMDVXX": "Oil (Power generator) Maldives, region XX",
    "PWROILNPLXX": "Oil (Power generator) Nepal, region XX",
    "PWRSHPINDEA": "Small Hydropower (Power generator) India, region EA",
    "PWRSHPINDNE": "Small Hydropower (Power generator) India, region NE",
    "PWRSHPINDNO": "Small Hydropower (Power generator) India, region NO",
    "PWRSHPINDSO": "Small Hydropower (Power generator) India, region SO",
    "PWRSHPINDWE": "Small Hydropower (Power generator) India, region WE",
    "PWRWOFMDVXX": "Offshore Wind (Power generator) Maldives, region XX",
    "PWRWONMDVXX": "Onshore Wind (Power generator) Maldives, region XX",
    "TRNBTNXXBGDXX": "Transmission interconnection from Bhutan, region XX to Bangladesh, region XX",
    "TRNBTNXXINDNE": "Transmission interconnection from Bhutan, region XX to India, region NE",
}


class CloneFuelTests(unittest.TestCase):
    def test_clones_draw_the_fuel_of_their_country_and_family(self):
        cases = [("GASIND", "PWRNGSINDEA", "PWRNGSBGDXX", "GASBGD"),
                 ("GASIND", "PWRNGSINDEA", "PWRNGSMDVXX", "GASMDV"),
                 ("OILBGD", "PWROILBGDXX", "PWROILMDVXX", "OILMDV"),
                 ("OILBGD", "PWROILBGDXX", "PWROILNPLXX", "OILNPL"),
                 ("HYDINDNO", "PWRHYDINDNO", "PWRGEOINDNO", "GEOINDNO")]
        for fuel, template, new, expected in cases:
            self.assertEqual(clone_fuel_code(fuel, template, new, PRODUCED), expected)

    def test_other_fuels_keep_their_code(self):
        self.assertEqual(clone_fuel_code("GASINT", "PWRNGSINDEA", "PWRNGSBGDXX", PRODUCED), "GASINT")
        # No SHPINDEA fuel is produced, so small hydro keeps the hydro resource.
        self.assertEqual(clone_fuel_code("HYDINDEA", "PWRHYDINDEA", "PWRSHPINDEA", PRODUCED), "HYDINDEA")
        self.assertEqual(clone_fuel_code("ELCBGDXX01", "PWRNGSINDEA", "PWRNGSBGDXX", PRODUCED), "ELCBGDXX01")
        self.assertEqual(clone_fuel_code("ELCBTNXX03", "TRNBTNXXINDEA", "TRNBTNXXBGDXX", PRODUCED), "ELCBTNXX03")


class MovedWorkbookValuesTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        wb = load_workbook(WORKBOOK, read_only=True, data_only=True)
        cls.decisions = {r[0]: r for r in wb["AO_Extension_Decisions"].iter_rows(min_row=2, values_only=True) if r[0]}
        cls.fixed = [r for r in wb["Fixed_Horizon_Parameters"].iter_rows(min_row=2, values_only=True)]
        wb.close()

    def test_cloned_technology_names(self):
        for tech, name in NAMES.items():
            self.assertEqual(self.decisions[tech][1], "Y", tech)
            self.assertEqual(self.decisions[tech][3], name, tech)

    def test_typed_link_ratios_and_no_unread_rows(self):
        oar = {r[3]: r[8] for r in self.fixed if r[6] == "OutputActivityRatio" and str(r[3]).startswith("TRN")}
        links = {t: v for t, v in oar.items() if len(t) == 13}
        self.assertEqual(links, {"TRNBTNXXBGDXX": 0.95, "TRNNPLXXBGDXX": 0.95,
                                 "TRNBTNXXINDNE": 0.95, "TRNMDVXXINDSO": 0.95})
        internal = ("PWRTRN", "RNWTRN", "TRNNLI", "TRNRPO", "RNWNLI", "RNWRPO")
        unread = [r[3:7] for r in self.fixed if isinstance(r[3], str)
                  and ((len(r[3]) == 11 and r[3][:6] in internal and r[6] in ("OperationalLife", "OutputActivityRatio"))
                       or (len(r[3]) == 13 and r[3].startswith("TRN") and r[6] == "OperationalLife"))]
        self.assertEqual(unread, [])


if __name__ == "__main__":
    unittest.main()
