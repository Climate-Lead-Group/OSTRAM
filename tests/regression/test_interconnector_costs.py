"""Interconnector costs: one central value per corridor, constant 2023 to 2050.

`Interconnector_Params` is the numeric authority for the CapitalCost and
FixedCost of the 18 interconnectors. Each row holds one value in every year.
Where the workbook has the `Transmission_Cost_Basis` sheet, its summary rows
name the same values and no cell holds a formula.
"""
from __future__ import annotations

import unittest
from pathlib import Path

from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parents[2]
WORKBOOK = REPO_ROOT / "inputs" / "scenarios" / "OSTRAM_Scenario_Inputs.xlsx"
TAXONOMY = REPO_ROOT / "config" / "scenarios" / "technology_types.csv"


def _interconnectors() -> set[str]:
    rows = [line.split(",") for line in TAXONOMY.read_text(encoding="utf-8-sig").splitlines()]
    return {r[1] for r in rows if len(r) > 1 and r[0] == "INTERCONNECTORS"}


class InterconnectorCostTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        book = load_workbook(WORKBOOK, read_only=True)
        sheet = book["Interconnector_Params"]
        header = [c for c in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
        cls.year_cols = [i for i, h in enumerate(header) if isinstance(h, (int, float)) and 2023 <= h <= 2050]
        cls.costs = {}
        for row in sheet.iter_rows(min_row=2, values_only=True):
            tech = next((v for v in row if isinstance(v, str) and v in _interconnectors()), None)
            param = next((v for v in row if v in ("CapitalCost", "FixedCost")), None)
            if tech and param:
                cls.costs[(tech, param)] = [row[i] for i in cls.year_cols]
        cls.basis = None
        if "Transmission_Cost_Basis" in book.sheetnames:
            cls.basis = list(book["Transmission_Cost_Basis"].iter_rows(values_only=True))
        book.close()

    def test_every_interconnector_has_one_cost_for_every_year(self) -> None:
        self.assertEqual(len(self.year_cols), 28)
        self.assertEqual({t for t, _ in self.costs}, _interconnectors())
        for key, values in self.costs.items():
            self.assertEqual(len(set(values)), 1, key)
            self.assertGreater(values[0], 0, key)

    def test_basis_sheet_names_the_same_values(self) -> None:
        if self.basis is None:
            self.skipTest("the workbook has no Transmission_Cost_Basis sheet")
        summary = {r[0]: (r[4], r[5]) for r in self.basis[1:19]}
        self.assertEqual(set(summary), _interconnectors())
        for tech, (capital, fixed) in summary.items():
            self.assertEqual(capital, self.costs[(tech, "CapitalCost")][0], tech)
            self.assertEqual(fixed, self.costs[(tech, "FixedCost")][0], tech)
        for row in self.basis:
            for value in row:
                self.assertFalse(isinstance(value, str) and value.startswith("="), value)


if __name__ == "__main__":
    unittest.main()
