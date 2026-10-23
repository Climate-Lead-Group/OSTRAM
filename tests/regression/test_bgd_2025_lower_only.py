"""Bangladesh 2025 generation has lower limits only, as in 2026.

The base-year pin removes the 2025 annual activity upper limits of the six
Bangladesh generators that have 2025 stock and keeps their 2025 lower limits
at 98% of BPDB's FY2024-25 references. The nuclear upper limit of 0 stays,
because Bangladesh has no nuclear stock in 2025.
"""
from __future__ import annotations

import csv
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook

from ostram.pipeline.scenarios.rules import apply_base_year_pin as PIN


REPO_ROOT = Path(__file__).resolve().parents[2]
PIN_CSV = REPO_ROOT / "config" / "scenarios" / "rules" / "pwr_min_2023_2026_pin.csv"

REMOVED = ("PWRCOABGDXX", "PWRHYDBGDXX", "PWRNGSBGDXX", "PWROILBGDXX", "PWRSPVBGDXX", "PWRWONBGDXX")
LOWER = REMOVED + ("PWRURNBGDXX",)
# BPDB Annual Report 2024-25, printed p. 14, in GWh; lower limit = GWh x 0.0036 x 0.98.
REFERENCE_GWH = {"PWRCOABGDXX": 27040, "PWRNGSBGDXX": 44613}


def _bgd_2025_rows() -> list[dict[str, str]]:
    with PIN_CSV.open(encoding="utf-8-sig", newline="") as stream:
        return [
            row for row in csv.DictReader(stream)
            if row["canonical_country"] == "BGD" and row["year"] == "2025"
            and row["parameter"] in (PIN.P_ACTIVITY_LOWER, PIN.P_ACTIVITY_UPPER)
        ]


class Bangladesh2025LowerOnlyTests(unittest.TestCase):
    def test_upper_limits_are_removed_and_lower_limits_kept(self) -> None:
        rows = {(r["technology"], r["parameter"]): r for r in _bgd_2025_rows()}
        for tech in REMOVED:
            upper = rows[(tech, PIN.P_ACTIVITY_UPPER)]
            self.assertEqual(upper["required_root_state"], "ABSENT", tech)
            self.assertEqual(upper["required_root_value"], "", tech)
        self.assertEqual(rows[("PWRURNBGDXX", PIN.P_ACTIVITY_UPPER)]["required_root_state"], "ZERO")
        for tech in LOWER:
            self.assertEqual(rows[(tech, PIN.P_ACTIVITY_LOWER)]["required_root_present"], "true", tech)
        for tech, gwh in REFERENCE_GWH.items():
            value = Decimal(rows[(tech, PIN.P_ACTIVITY_LOWER)]["required_root_value"])
            self.assertAlmostEqual(float(value), gwh * 0.0036 * 0.98, places=6)

    def test_applying_the_pin_clears_the_2025_upper_cells(self) -> None:
        rows = _bgd_2025_rows()
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            rules = tmp / "rules.csv"
            with rules.open("w", encoding="utf-8", newline="") as stream:
                writer = csv.DictWriter(stream, fieldnames=PIN.EXPECTED_FIELDS, lineterminator="\n")
                writer.writeheader()
                writer.writerows(rows)
            workbook = Workbook()
            sheets = {"Primary Techs": workbook.active}
            sheets["Primary Techs"].title = "Primary Techs"
            sheets["Secondary Techs"] = workbook.create_sheet("Secondary Techs")
            sheets["Capacities"] = workbook.create_sheet("Capacities")
            for sheet in sheets.values():
                for column, value in enumerate(
                    ("Tech", "Parameter", "Projection.Mode", "Projection.Parameter", "Timeslices", 2025), start=1
                ):
                    sheet.cell(row=1, column=column, value=value)
            secondary = sheets["Secondary Techs"]
            row_number = 2
            for tech in LOWER:
                for parameter in (PIN.P_ACTIVITY_LOWER, PIN.P_ACTIVITY_UPPER):
                    for column, value in enumerate((tech, parameter, "User defined", 0, None, 999.0), start=1):
                        secondary.cell(row=row_number, column=column, value=value)
                    row_number += 1
            workbook.save(tmp / PIN.PARAM_FILE)
            PIN.apply_pin_rules(tmp, "A_Calibrated_BAU", rules, skip_backup=True)
            result = load_workbook(tmp / PIN.PARAM_FILE)["Secondary Techs"]
            cells = {
                (result.cell(r, 1).value, result.cell(r, 2).value): result.cell(r, 6).value
                for r in range(2, result.max_row + 1)
            }
        for tech in REMOVED:
            self.assertIsNone(cells[(tech, PIN.P_ACTIVITY_UPPER)], tech)
            self.assertNotEqual(cells[(tech, PIN.P_ACTIVITY_LOWER)], 999.0, tech)
        self.assertEqual(cells[("PWRURNBGDXX", PIN.P_ACTIVITY_UPPER)], 0)


if __name__ == "__main__":
    unittest.main()
