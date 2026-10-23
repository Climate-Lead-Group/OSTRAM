"""The internal transmission residuals are computed in every run from that run's inputs.

``RNWTRN<node>`` and ``PWRTRN<node>`` get their base-year ``ResidualCapacity`` from
the node's 2023 peak time-slice demand and the renewable capacity available at that
peak, read from the delivered workbooks after the base-year pin. No stored per-node
file is read.
"""
from __future__ import annotations

import inspect
import tempfile
import unittest
from pathlib import Path

import yaml
from openpyxl import Workbook, load_workbook

from ostram.pipeline.scenarios.rules import apply_internal_transmission as itx


REPO_ROOT = Path(__file__).resolve().parents[2]
MODE = "Projection.Mode"
USER = "User defined"
SLICES = {"S1": 0.5, "S2": 0.5}
# Node AAAXX: 8,760 GWh a year (31.536 PJ), 60% in S1, so the peak is 1.2 GW in S1.
# Node BBBXX: 4,380 GWh a year (15.768 PJ), peak 0.5 GW, the Bhutan case where RE at peak exceeds the peak.
DEMAND = {"ELCAAAXX03": 31.536, "ELCBBBXX03": 15.768}
PROFILE = {("ELCAAAXX03", "S1"): 0.6, ("ELCAAAXX03", "S2"): 0.4,
           ("ELCBBBXX03", "S1"): 0.5, ("ELCBBBXX03", "S2"): 0.5}
RESIDUAL = [
    ("PWRSPVAAAXX", USER, 1.0),     # variable RE, CF 0.5 in S1
    ("PWRBIOAAAXX", USER, 0.1),     # firm RE, counted at CF 1
    ("PWRSHPAAAXX", USER, 0.2),     # small hydro, counted as non-RE
    ("PWRCOAAAAXX", USER, 5.0),     # non-RE
    ("PWRWONAAAXX", "EMPTY", 3.0),  # not compiled, so not counted
    ("PWRHYDBBBXX", USER, 1.0),     # variable RE, CF 0.9 in both slices
]
CF = {("PWRSPVAAAXX", "S1"): 0.5, ("PWRSPVAAAXX", "S2"): 0.0,
      ("PWRBIOAAAXX", "S1"): 0.3, ("PWRBIOAAAXX", "S2"): 0.3,
      ("PWRWONAAAXX", "S1"): 1.0, ("PWRWONAAAXX", "S2"): 1.0,
      ("PWRHYDBBBXX", "S1"): 0.9, ("PWRHYDBBBXX", "S2"): 0.9}
NODES = ("AAAXX", "BBBXX")


def _write_inputs(folder: Path, demand=DEMAND, cf=CF, residual=RESIDUAL) -> None:
    """Write the two delivered workbooks the residual part reads."""
    dem = Workbook()
    ws = dem.active
    ws.title = "Demand_Projection"
    ws.append(["Demand/Share", "Fuel/Tech", MODE, "Projection.Parameter", "2023", "2024"])
    for fuel, value in demand.items():
        ws.append(["Demand", fuel, USER, 0, value, value])
    ws = dem.create_sheet("Profiles")
    ws.append(["Timeslices", "Demand/Share", "Fuel/Tech", MODE, "Projection.Parameter", "2023"])
    for (fuel, ts), value in PROFILE.items():
        ws.append([ts, "Demand", fuel, USER, 0, value])
    dem.save(folder / itx.DEMAND_FILE_NAME)

    par = Workbook()
    ws = par.active
    ws.title = "Yearsplit"
    ws.append(["Timeslices", "Parameter", MODE, "2023"])
    for ts, value in SLICES.items():
        ws.append([ts, "YearSplit", USER, value])
    ws = par.create_sheet("Capacities")
    ws.append(["Timeslices", "Tech", "Parameter", MODE, "2023"])
    for (tech, ts), value in cf.items():
        ws.append([ts, tech, "CapacityFactor", USER, value])
    ws = par.create_sheet("Secondary Techs")
    ws.append(["Tech", "Parameter", MODE, "2023", "2024"])
    for tech, mode, value in residual:
        ws.append([tech, "ResidualCapacity", mode, value, value])
    ws = par.create_sheet("Demand Techs")
    ws.append(["Tech", "Parameter", MODE, 2023, 2024])
    for fam in itx.ALL_FAMILIES:
        for node in NODES:
            ws.append([f"{fam}{node}", "ResidualCapacity", USER, 5, 5])
    ws = par.create_sheet("Fixed Horizon Parameters")
    ws.append(["Tech", "Parameter", "Value"])
    par.save(folder / itx.PARAM_FILE_NAME)


def _write_config(folder: Path, margin=1.2) -> Path:
    block = {"esmap_capital_cost_usd2010_per_kw": 100, "muv_2010": 100, "muv_2022": 100,
             "muv_2023": 100, "fixed_cost_share": 0.04, "gec2_cost_rs_crore": 10,
             "gec2_capacity_mw": 1000, "inr_per_usd_2022": 1, "operational_life": 40}
    if margin is not None:
        block[itx.MARGIN_KEY] = margin
    path = folder / itx.CONFIG_NAME
    path.write_text(yaml.safe_dump({"internal_transmission": block}), encoding="utf-8")
    return path


class InternalTransmissionResidualTests(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name)

    def tearDown(self) -> None:
        self._tmp.cleanup()

    def test_formula_on_synthetic_inputs(self) -> None:
        _write_inputs(self.dir)
        got = itx.compute_residuals(self.dir, 1.2)
        self.assertEqual(sorted(got), list(NODES))
        a = got["AAAXX"]
        self.assertEqual(a["peak_timeslice"], "S1")
        self.assertAlmostEqual(a["peak_GW"], 1.2, places=12)
        # RE at peak: 1.0 x 0.5 (SPV) + 0.1 x 1 (BIO); SHP, COA and the EMPTY WON row do not count.
        self.assertAlmostEqual(a["RE_at_peak_GW"], 0.6, places=12)
        self.assertAlmostEqual(a["RNWTRN"], 0.72, places=12)
        self.assertAlmostEqual(a["PWRTRN"], 1.2 * 1.2 - 0.72, places=12)
        b = got["BBBXX"]
        self.assertAlmostEqual(b["peak_GW"], 0.5, places=12)
        self.assertAlmostEqual(b["RE_at_peak_GW"], 0.9, places=12)
        self.assertAlmostEqual(b["RNWTRN"], 0.6, places=12)
        self.assertEqual(b["PWRTRN"], 0.0)

    def test_residual_part_writes_both_families_and_zero_for_the_others(self) -> None:
        _write_inputs(self.dir)
        cfg = _write_config(self.dir)
        itx.run(self.dir, skip_backup=True, config=cfg, parts=("residuals",))
        wb = load_workbook(self.dir / itx.PARAM_FILE_NAME, data_only=True)
        got = {(r[0], r[1]): r[3] for r in wb["Demand Techs"].iter_rows(min_row=2, values_only=True)}
        wb.close()
        self.assertAlmostEqual(got[("RNWTRNAAAXX", "ResidualCapacity")], 0.72, places=12)
        self.assertAlmostEqual(got[("PWRTRNAAAXX", "ResidualCapacity")], 0.72, places=12)
        self.assertAlmostEqual(got[("RNWTRNBBBXX", "ResidualCapacity")], 0.6, places=12)
        self.assertEqual(got[("PWRTRNBBBXX", "ResidualCapacity")], 0.0)
        for fam in ("RNWNLI", "RNWRPO", "TRNNLI", "TRNRPO"):
            self.assertEqual(got[(f"{fam}AAAXX", "ResidualCapacity")], 0.0)

    def test_missing_demand_stops_naming_node_and_input(self) -> None:
        _write_inputs(self.dir, demand={"ELCAAAXX03": 31.536})
        with self.assertRaises(itx.MissingInputError) as ctx:
            itx.compute_residuals(self.dir, 1.2)
        self.assertIn("BBBXX", str(ctx.exception))
        self.assertIn("SpecifiedAnnualDemand", str(ctx.exception))

    def test_missing_peak_capacity_factor_uses_the_generator_mean(self) -> None:
        # PWRSPVAAAXX keeps only S2 (0.0); PWRHYDBBBXX has no row at all, so counts at 1.
        cf = {k: v for k, v in CF.items() if k[0] != "PWRHYDBBBXX" and k != ("PWRSPVAAAXX", "S1")}
        _write_inputs(self.dir, cf=cf)
        got = itx.compute_residuals(self.dir, 1.2)
        self.assertAlmostEqual(got["AAAXX"]["RE_at_peak_GW"], 0.1, places=12)
        self.assertAlmostEqual(got["BBBXX"]["RE_at_peak_GW"], 1.0, places=12)
        self.assertAlmostEqual(got["BBBXX"]["RNWTRN"], 0.6, places=12)
        self.assertEqual(got["BBBXX"]["PWRTRN"], 0.0)

    def test_missing_profile_stops_naming_node_and_input(self) -> None:
        _write_inputs(self.dir)
        wb = load_workbook(self.dir / itx.DEMAND_FILE_NAME)
        ws = wb["Profiles"]
        for row in range(ws.max_row, 1, -1):
            if ws.cell(row=row, column=3).value == "ELCAAAXX03" and ws.cell(row=row, column=1).value == "S2":
                ws.delete_rows(row)
        wb.save(self.dir / itx.DEMAND_FILE_NAME)
        with self.assertRaises(itx.MissingInputError) as ctx:
            itx.compute_residuals(self.dir, 1.2)
        self.assertIn("AAAXX", str(ctx.exception))
        self.assertIn("SpecifiedDemandProfile", str(ctx.exception))

    def test_missing_margin_stops(self) -> None:
        _write_inputs(self.dir)
        cfg = _write_config(self.dir, margin=None)
        with self.assertRaises(itx.MissingInputError) as ctx:
            itx.run(self.dir, skip_backup=True, config=cfg, parts=("residuals",))
        self.assertIn(itx.MARGIN_KEY, str(ctx.exception))

    def test_unread_projection_mode_stops(self) -> None:
        residual = RESIDUAL + [("PWRWOFAAAXX", "Growth", 1.0)]
        _write_inputs(self.dir, residual=residual)
        with self.assertRaises(itx.MissingInputError) as ctx:
            itx.compute_residuals(self.dir, 1.2)
        self.assertIn("PWRWOFAAAXX", str(ctx.exception))

    def test_build_reads_no_stored_residual_file(self) -> None:
        from ostram.pipeline.scenarios import transform

        self.assertEqual(list((REPO_ROOT / "config").rglob("internal_tx_residuals*")), [])
        for name in ("RESIDUALS_NAME", "resolve_residuals", "read_residuals"):
            self.assertFalse(hasattr(itx, name), name)
        for stage in (transform.stage_internal_transmission, transform.stage_internal_tx_residuals):
            source = inspect.getsource(stage)
            self.assertNotIn(".csv", source, stage.__name__)
            self.assertNotIn("--residuals", source, stage.__name__)

    def test_margin_is_set_in_the_south_asia_configuration(self) -> None:
        cfg = REPO_ROOT / "config" / "preparation" / "Config_country_codes.yaml"
        knobs = itx.read_config(cfg)
        self.assertIn(itx.MARGIN_KEY, knobs)


if __name__ == "__main__":
    unittest.main()
