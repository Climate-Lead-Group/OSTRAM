"""Cross-border forced builds by root scenario and A's existing-plus-committed limits.

Both read the workbook sheet Interconnector_Params in every run:
* a root scenario's TotalAnnualMinCapacityInvestment row replaces the BAU row
  of its link for that root and the scenarios derived from it;
* A_Calibrated_BAU's cross-border total capacity is ResidualCapacity plus the
  surviving forced builds, and its yearly build cap the larger of
  ResidualCapacity and the year's forced build.
"""
from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest import mock

import yaml
from openpyxl import Workbook, load_workbook

from ostram.pipeline.scenarios import orchestrator, transform
from ostram.pipeline.scenarios.registry import load_registry
from ostram.pipeline.scenarios.rules import relax_interconnectors as relax
from ostram.pipeline.scenarios.rules import relax_to_reference
from ostram.pipeline.scenarios.transformations import fix_trn_residuals
from ostram.pipeline.scenarios.transformations import interconnector_authority as authority
from ostram.pipeline.scenarios.transformations import scenario_workbooks
from ostram.pipeline.scenarios.transformations.cap_trn_to_residual import TRN_TECHS
from tests.regression.setup_file import SETUP_FILE


REPO_ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = REPO_ROOT / "inputs" / "scenarios" / "OSTRAM_Scenario_Inputs.xlsx"
TECH_TYPES = REPO_ROOT / "config" / "scenarios" / "technology_types.csv"
SCENARIO_CONFIG = REPO_ROOT / "config" / "scenarios"
YEARS = tuple(range(2023, 2051))
MIN = "TotalAnnualMinCapacityInvestment"
CROSS_BORDER = frozenset(tech for tech in TRN_TECHS if relax.is_cross_border(tech))
BAU_FORCED = {("TRNINDEANPLXX", 2027): 0.4, ("TRNINDNONPLXX", 2028): 1.15}
C_FORCED = {
    ("TRNINDEANPLXX", 2027): 1.4,
    ("TRNINDNONPLXX", 2028): 2.15,
    ("TRNBTNXXINDNE", 2030): 1.125,
    ("TRNBGDXXINDEA", 2031): 2.0,
    ("TRNBGDXXINDNE", 2031): 1.0,
    ("TRNINDSOLKAXX", 2039): 0.5,
}


def _forced(loaded: dict) -> dict:
    """{(tech, year): value} of the non-zero cross-border cells."""
    return {
        (tech, year): value
        for tech, profile in loaded.items() if tech in CROSS_BORDER
        for year, value in profile.items() if value
    }


def _authority_rows(workbook):
    ws = workbook["Interconnector_Params"]
    headers = {cell.value: cell.column for cell in ws[1] if cell.value is not None}
    return ws, headers


class ForcedBuildsByRootTests(unittest.TestCase):
    """Decision 74: committed builds in BAU rows, committed plus planned in C rows."""

    def test_bau_rows_reach_a_b_and_support_bau(self) -> None:
        for scenario in ("BAU", "A_Calibrated_BAU", "B_Optimised_VRE"):
            with self.subTest(scenario=scenario):
                loaded = authority.load_minimum_contribution_authority(TEMPLATE, scenario)
                self.assertEqual(set(loaded), set(authority.MINIMUM_CONTRIBUTION_TECHS))
                self.assertEqual(_forced(loaded), BAU_FORCED)

    def test_c_row_reaches_c_and_c_derived_scenarios(self) -> None:
        loaded = authority.load_minimum_contribution_authority(TEMPLATE, "C_Target_VRE")
        self.assertEqual(_forced(loaded), C_FORCED)
        # The India-internal BAU rows still apply to C.
        bau = authority.load_minimum_contribution_authority(TEMPLATE, "BAU")
        for tech in authority.MINIMUM_CONTRIBUTION_TECHS - CROSS_BORDER:
            self.assertEqual(loaded[tech], bau[tech])

        # The materialized C template carries the same merged rows.
        with tempfile.TemporaryDirectory() as directory:
            materialized = Path(directory) / "_materialized_C_Target_VRE.xlsx"
            scenario_workbooks.materialize_scenario_template(TEMPLATE, "C_Target_VRE", materialized)
            self.assertEqual(
                authority.load_minimum_contribution_authority(materialized, "C_Target_VRE"), loaded
            )

        # Derived scenarios are built from their root's A3 output: every
        # C-derived scenario takes C's rows.
        registry = load_registry(SETUP_FILE)
        derived = [d.name for d in registry.derived if d.base_scenario == "C_Target_VRE"]
        self.assertTrue(derived)
        for name in derived:
            base = registry.derived_by_name[name].base_scenario
            self.assertEqual(authority.load_minimum_contribution_authority(TEMPLATE, base), loaded)

    def test_stage_3_passes_the_root_scenario(self) -> None:
        calls = []

        def fake_subproc(module, args, **kwargs):
            calls.append((module, [str(a) for a in args]))

        with tempfile.TemporaryDirectory() as directory:
            s2, s3 = Path(directory) / "s2", Path(directory) / "s3"
            s2.mkdir()
            s3.mkdir()
            (s2 / "A-O_Parametrization_c2a_patched.xlsx").write_bytes(b"")
            with mock.patch.dict("os.environ", {"OSTRAM_TEMPLATE_PATH": str(TEMPLATE)}), \
                    mock.patch.object(transform, "run_subproc", side_effect=fake_subproc):
                transform.stage_3_fix_2(s2, s3, "C_Target_VRE")
        fix = next(args for module, args in calls if module.endswith("fix_trn_residuals"))
        self.assertEqual(fix[fix.index("--scenario") + 1], "C_Target_VRE")
        source = Path(orchestrator.__file__).read_text(encoding="utf-8")
        self.assertIn("dependencies.stage_3_fix_2(\n        stage2, stage3, plan.scenario\n    )", source)

    def _corrupted(self, mutate):
        workbook = load_workbook(TEMPLATE, data_only=False)
        ws, headers = _authority_rows(workbook)
        rows = [r for r in range(2, ws.max_row + 1)
                if ws.cell(r, headers["Parameter"]).value == MIN
                and ws.cell(r, headers["scenario"]).value == "C_Target_VRE"]
        self.assertTrue(rows)
        mutate(ws, headers, rows[0])
        return workbook

    def test_unregistered_scenario_name_stops_the_build(self) -> None:
        for name in ("C_Target_VRE_Clipped", "C_Target", "NDC"):
            with self.subTest(name=name):
                def mutate(ws, headers, row):
                    ws.cell(row, headers["scenario"]).value = name
                workbook = self._corrupted(mutate)
                with mock.patch.object(authority, "load_workbook", return_value=workbook):
                    with self.assertRaisesRegex(ValueError, "not a registered root scenario"):
                        authority.load_minimum_contribution_authority(TEMPLATE, "C_Target_VRE")
        with self.assertRaisesRegex(ValueError, "not a registered root scenario"):
            authority.load_minimum_contribution_authority(TEMPLATE, "C_Target_VRE_Clipped")

    def test_two_rows_for_one_root_and_link_stop_the_build(self) -> None:
        def mutate(ws, headers, row):
            ws.append([ws.cell(row, c).value for c in range(1, ws.max_column + 1)])
        workbook = self._corrupted(mutate)
        with mock.patch.object(authority, "load_workbook", return_value=workbook):
            with self.assertRaisesRegex(ValueError, "duplicate row"):
                authority.load_minimum_contribution_authority(TEMPLATE, "C_Target_VRE")

    def test_root_rows_only_for_forced_builds_and_with_a_source(self) -> None:
        def boundary_row(ws, headers, row):
            for r in range(2, ws.max_row + 1):
                if ws.cell(r, headers["Parameter"]).value == authority.MINIMUM_BOUNDARY_PARAMETER:
                    ws.cell(r, headers["scenario"]).value = "C_Target_VRE"
                    return
        workbook = self._corrupted(boundary_row)
        with mock.patch.object(authority, "load_workbook", return_value=workbook):
            with self.assertRaisesRegex(ValueError, "scenario must be 'BAU'"):
                authority.load_minimum_boundary_authority(TEMPLATE)

        def blank_source(ws, headers, row):
            ws.cell(row, headers["Source"]).value = " "
        workbook = self._corrupted(blank_source)
        with mock.patch.object(authority, "load_workbook", return_value=workbook):
            with self.assertRaisesRegex(ValueError, "Source must name"):
                authority.load_minimum_contribution_authority(TEMPLATE, "C_Target_VRE")


def _parametrization(path: Path, residual: dict, minimum: dict, life: int = 40) -> None:
    """A small Secondary Techs + Fixed Horizon Parameters workbook for the 18 TRN links."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Secondary Techs"
    ws.append(["Tech.ID", "Tech", "Tech.Name", "Parameter.ID", "Parameter", "Unit",
               "Projection.Mode", "Projection.Parameter", *YEARS])
    for tech in sorted(TRN_TECHS):
        res = residual.get(tech, 1.0)
        rows = {
            "ResidualCapacity": [res] * len(YEARS),
            MIN: [minimum.get((tech, y), 0.0) for y in YEARS],
            "TotalAnnualMaxCapacity": [res] * len(YEARS),
            "TotalAnnualMaxCapacityInvestment": [9999] * len(YEARS),
        }
        for parameter, values in rows.items():
            mode = "EMPTY" if parameter == "TotalAnnualMaxCapacity" else "User defined"
            ws.append([1, tech, tech, 1, parameter, None, mode, 0, *values])
    fixed = wb.create_sheet("Fixed Horizon Parameters")
    fixed.append(["Tech", "Parameter", "Value"])
    for tech in sorted(TRN_TECHS):
        fixed.append([tech, "OperationalLife", life])
    wb.save(path)
    wb.close()


def _rows(path: Path) -> dict:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb["Secondary Techs"]
        values = list(ws.values)
        header = list(values[0])
        tech, parameter, mode = (header.index(n) for n in ("Tech", "Parameter", "Projection.Mode"))
        years = {y: header.index(y) for y in YEARS}
        return {(r[tech], r[parameter]): ({y: r[c] for y, c in years.items()}, r[mode]) for r in values[1:]}
    finally:
        wb.close()


def _yaml(path: Path, content: dict) -> Path:
    path.write_text(yaml.safe_dump(content), encoding="utf-8")
    return path


A_CONFIG = {"mode": "multiplicative", "headroom_factor": 1.0,
            "cross_border_limit": "existing_plus_committed",
            "overrides": {"TRNINDEAINDNO": {2023: 22.53, 2029: 25.53}}}


class ExistingPlusCommittedTests(unittest.TestCase):
    """Decisions 72 and 75: A's cross-border total capacity is existing plus committed."""

    def run_a(self, directory: Path, residual: dict, minimum: dict, life: int = 40) -> dict:
        _parametrization(directory / relax.PARAM_FILE_NAME, residual, minimum, life)
        relax.run(directory, skip_backup=True, yaml_path=_yaml(directory / "a.yaml", A_CONFIG),
                  tech_types_path=TECH_TYPES)
        return _rows(directory / relax.PARAM_FILE_NAME)

    def test_a_limits_equal_residual_plus_committed(self) -> None:
        residual = {"TRNINDNONPLXX": 0.5, "TRNINDEANPLXX": 1.2, "TRNNPLXXBGDXX": 0.0}
        minimum = {("TRNINDNONPLXX", 2028): 1.15, ("TRNINDEANPLXX", 2027): 0.4,
                   ("TRNNPLXXBGDXX", 2030): 0.3}
        with tempfile.TemporaryDirectory() as tmp:
            rows = self.run_a(Path(tmp), residual, minimum, life=10)
        for tech in CROSS_BORDER:
            res = residual.get(tech, 1.0)
            cap, cap_mode = rows[(tech, "TotalAnnualMaxCapacity")]
            inv, _ = rows[(tech, "TotalAnnualMaxCapacityInvestment")]
            low, _ = rows[(tech, MIN)]
            self.assertEqual(cap_mode, "User defined")
            for year in YEARS:
                surviving = sum(v for (t, y), v in minimum.items() if t == tech and 0 <= year - y < 10)
                self.assertAlmostEqual(cap[year], res + surviving, msg=(tech, year))
                self.assertAlmostEqual(inv[year], max(res, minimum.get((tech, year), 0.0)), msg=(tech, year))
                # No committed build is cleared.
                self.assertAlmostEqual(low[year], minimum.get((tech, year), 0.0), msg=(tech, year))
        # The vintage of 2028 retires after ten years.
        cap = rows[("TRNINDNONPLXX", "TotalAnnualMaxCapacity")][0]
        self.assertAlmostEqual(cap[2037], 1.65)
        self.assertAlmostEqual(cap[2038], 0.5)
        # India-internal links keep their own rule.
        self.assertEqual(rows[("TRNINDEAINDNO", "TotalAnnualMaxCapacity")][0][2030], 9999)
        self.assertEqual(rows[("TRNINDSOINDWE", "TotalAnnualMaxCapacity")][0][2030], 1.0)

    def test_changed_workbook_minimum_changes_a_limit(self) -> None:
        """Interconnector_Params -> fix_trn_residuals -> relax_interconnectors, no other edit."""
        def a_limits(directory: Path, workbook_path: Path) -> dict:
            _parametrization(directory / "in.xlsx", {}, {})
            fix_trn_residuals.run_fix(directory / "in.xlsx", directory / relax.PARAM_FILE_NAME,
                                      authority_path=workbook_path, scenario="A_Calibrated_BAU")
            relax.run(directory, skip_backup=True, yaml_path=_yaml(directory / "a.yaml", A_CONFIG),
                      tech_types_path=TECH_TYPES)
            return _rows(directory / relax.PARAM_FILE_NAME)

        with tempfile.TemporaryDirectory() as tmp:
            base, edited = Path(tmp) / "base", Path(tmp) / "edited"
            base.mkdir()
            edited.mkdir()
            before = a_limits(base, TEMPLATE)
            workbook = load_workbook(TEMPLATE)
            ws, headers = _authority_rows(workbook)
            row = next(r for r in range(2, ws.max_row + 1)
                       if ws.cell(r, headers["Parameter"]).value == MIN
                       and ws.cell(r, headers["scenario"]).value == "BAU"
                       and ws.cell(r, headers["Tech"]).value == "TRNINDNONPLXX")
            ws.cell(row, headers[2028]).value = 0.65
            ws.cell(row, headers[2035]).value = 0.25
            copy = Path(tmp) / "inputs.xlsx"
            workbook.save(copy)
            workbook.close()
            after = a_limits(edited, copy)

        cap0 = before[("TRNINDNONPLXX", "TotalAnnualMaxCapacity")][0]
        cap1 = after[("TRNINDNONPLXX", "TotalAnnualMaxCapacity")][0]
        self.assertAlmostEqual(cap0[2028], 0.5 + 1.15)
        self.assertAlmostEqual(cap1[2028], 0.5 + 0.65)
        self.assertAlmostEqual(cap1[2035], 0.5 + 0.65 + 0.25)
        inv1 = after[("TRNINDNONPLXX", "TotalAnnualMaxCapacityInvestment")][0]
        self.assertAlmostEqual(inv1[2028], 0.65)
        self.assertAlmostEqual(inv1[2035], 0.5)
        for key in before:
            if key[0] != "TRNINDNONPLXX":
                self.assertEqual(before[key], after[key], key)

    def test_b_caps_untouched_where_b_is_looser(self) -> None:
        residual = {"TRNINDNONPLXX": 0.5, "TRNINDEANPLXX": 1.2, "TRNLKAXXMDVXX": 0.0}
        minimum = {("TRNINDNONPLXX", 2028): 1.15, ("TRNINDEANPLXX", 2027): 0.4}
        b_config = {"mode": "multiplicative", "headroom_factor": 1.5, "overrides": {
            "TRNINDNONPLXX": {2023: 0.5, 2025: 1.5, 2028: 2.65},
            "TRNINDEANPLXX": {2023: 1.35, 2027: 2.35, 2030: 4.35},
            "TRNLKAXXMDVXX": {2023: 0}}}
        with tempfile.TemporaryDirectory() as tmp:
            a_dir, b_dir = Path(tmp) / "a", Path(tmp) / "b"
            a_dir.mkdir()
            b_dir.mkdir()
            self.run_a(a_dir, residual, minimum)
            _parametrization(b_dir / relax.PARAM_FILE_NAME, residual, minimum)
            relax.run(b_dir, skip_backup=True, yaml_path=_yaml(b_dir / "b.yaml", b_config),
                      tech_types_path=TECH_TYPES)
            before = _rows(b_dir / relax.PARAM_FILE_NAME)
            log = relax_to_reference.run(b_dir, a_dir)
            after = _rows(b_dir / relax.PARAM_FILE_NAME)
        for tech in ("TRNINDNONPLXX", "TRNINDEANPLXX", "TRNLKAXXMDVXX"):
            for parameter in ("TotalAnnualMaxCapacity", "TotalAnnualMaxCapacityInvestment", MIN):
                self.assertEqual(before[(tech, parameter)], after[(tech, parameter)], (tech, parameter))
        changed = {(c["tech"], c["parameter"]) for c in log["changes"]}
        self.assertFalse(changed & {(t, p) for t in CROSS_BORDER for p in
                                    ("TotalAnnualMaxCapacity", "TotalAnnualMaxCapacityInvestment", MIN)})

    def test_switch_rejects_cross_border_overrides_and_unknown_values(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            directory = Path(tmp)
            _parametrization(directory / relax.PARAM_FILE_NAME, {}, {})
            bad = dict(A_CONFIG, overrides={"TRNBGDXXINDEA": {2023: 2.5}})
            with self.assertRaisesRegex(ValueError, "remove their overrides"):
                relax.run(directory, skip_backup=True, yaml_path=_yaml(directory / "a.yaml", bad),
                          tech_types_path=TECH_TYPES)
            with self.assertRaisesRegex(ValueError, "Invalid cross_border_limit"):
                relax.load_config(_yaml(directory / "c.yaml", dict(A_CONFIG, cross_border_limit="open")))

    def test_repository_yaml_files(self) -> None:
        a = relax.load_config(SCENARIO_CONFIG / "A_Calibrated_BAU" / "relax_interconnectors.yaml")
        self.assertEqual(a["cross_border_limit"], "existing_plus_committed")
        self.assertEqual(set(a["overrides"]),
                         {"TRNINDEAINDNO", "TRNINDEAINDSO", "TRNINDSOINDWE", "TRNINDEAINDNE"})
        self.assertEqual(a["headroom_factor"], 1.0)
        for root in ("B_Optimised_VRE", "C_Target_VRE"):
            config = relax.load_config(SCENARIO_CONFIG / root / "relax_interconnectors.yaml")
            self.assertIsNone(config["cross_border_limit"])
            self.assertTrue(CROSS_BORDER <= set(config["overrides"]), root)


if __name__ == "__main__":
    unittest.main()
