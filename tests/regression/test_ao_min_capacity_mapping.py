"""Exercise the mapping helpers without executing the file's pipeline driver."""
import ast
from pathlib import Path
import unittest

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill

SOURCE = Path(__file__).resolve().parents[2] / "ostram/pipeline/scenarios/transformations/update_ao_from_extensions.py"
NAMES = {"_norm_key_part", "color_row", "build_wv_lookup",
         "refresh_year_cells_in_place", "append_missing_min_capacity_rows"}
namespace = {"pd": pd, "PatternFill": PatternFill, "PARAM_REFRESH_COLOR": "B7D7E8"}
tree = ast.parse(SOURCE.read_text(encoding="utf-8"))
exec(compile(ast.Module(body=[n for n in tree.body if isinstance(n, ast.FunctionDef) and n.name in NAMES], type_ignores=[]), str(SOURCE), "exec"), namespace)


class MinCapacityMappingTests(unittest.TestCase):
    def setUp(self):
        self.ws = Workbook().active
        self.headers = ["Tech.ID", "Tech", "Tech.Name", "Parameter.ID", "Parameter",
                        "Unit", "Projection.Mode", "Projection.Parameter", 2027, 2028]
        self.ws.append(self.headers)
        self.ws.append([91, "SOLAR", "Existing solar", 10, "TotalAnnualMaxCapacity",
                        "GW", "User defined", None, .55, .70])

    def apply(self, lookup):
        df = pd.DataFrame(list(self.ws.values)[1:], columns=self.headers)
        return namespace["append_missing_min_capacity_rows"](self.ws, df, lookup, [2027, 2028])

    def test_positive_floor_reaches_native_sheet_and_cache_with_identity(self):
        lookup = {("SOLAR", "TotalAnnualMinCapacity"): {2027: 0, 2028: .3841}}
        df, count = self.apply(lookup)
        self.assertEqual(count, 1)
        row = dict(zip(self.headers, list(self.ws.values)[-1]))
        self.assertEqual((row["Tech.ID"], row["Tech"], row["Tech.Name"]), (91, "SOLAR", "Existing solar"))
        self.assertEqual((row["Parameter"], row["Unit"], row["Projection.Mode"]), ("TotalAnnualMinCapacity", "GW", "User defined"))
        self.assertEqual((row[2027], row[2028]), (0, .3841))
        self.assertEqual(df.iloc[-1][2028], .3841)
        self.assertEqual(self.apply(lookup)[1], 0)
        self.assertEqual(self.ws.max_row, 3)

    def test_blank_and_zero_defaults_other_parameters_and_unknown_tech_unchanged(self):
        original = list(self.ws.values)
        for values in [{2027: None, 2028: float("nan")}, {2027: 0, 2028: 0}]:
            self.assertEqual(self.apply({("SOLAR", "TotalAnnualMinCapacity"): values})[1], 0)
        self.assertEqual(self.apply({("UNKNOWN", "TotalAnnualMinCapacity"): {2028: .3841}})[1], 0)
        self.assertEqual(self.apply({("SOLAR", "TotalAnnualMinCapacityInvestment"): {2028: .272}})[1], 0)
        self.assertEqual(list(self.ws.values), original)

    def test_existing_stock_row_is_refreshed_without_duplicate(self):
        self.ws.append([91, "SOLAR", "Existing solar", None, "TotalAnnualMinCapacity", "GW", "User defined", None, 0, .1])
        df = pd.DataFrame(list(self.ws.values)[1:], columns=self.headers)
        lookup = {("SOLAR", "TotalAnnualMinCapacity"): {2027: 0, 2028: .3841}}
        *_, refreshed = namespace["refresh_year_cells_in_place"](self.ws, df, lookup, ["Tech", "Parameter"], [2027, 2028])
        result, count = namespace["append_missing_min_capacity_rows"](self.ws, refreshed, lookup, [2027, 2028])
        self.assertEqual((count, self.ws.max_row, result.iloc[-1][2028]), (0, 3, .3841))

    def test_absent_year_uses_standard_zero_without_copying_template_ceiling(self):
        df, count = self.apply({("SOLAR", "TotalAnnualMinCapacity"): {2028: .3841}})
        self.assertEqual((count, df.iloc[-1][2027], df.iloc[-1][2028]), (1, 0, .3841))


if __name__ == "__main__":
    unittest.main()
