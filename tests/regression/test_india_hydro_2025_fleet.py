"""The 2025 hydro fleet of India North and South can generate its pinned 2025 floor.

The base-year lock lowers an activity floor that the locked fleet cannot reach.
With the CEA monitored hydro capacity of December 2025 in the workbook, the most
the 2025 fleet can generate (ResidualCapacity x 31.536 x the sum over time slices
of CapacityFactor x YearSplit x AvailabilityFactor) is at or above the pin's 2025
lower limit, so the lock leaves the floor as it is.
"""
import csv
from pathlib import Path
import unittest

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
PIN = ROOT / "config/scenarios/rules/pwr_min_2023_2026_pin.csv"
WORKBOOK = ROOT / "inputs/scenarios/OSTRAM_Scenario_Inputs.xlsx"
TIMESLICES = ROOT / "inputs/scenarios/OSTRAM_Timeslice_Inputs.xlsx"
TECHS = ("PWRHYDINDNO", "PWRHYDINDSO")


class IndiaHydro2025FleetTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.cf, cls.floor = {}, {}
        with PIN.open(encoding="utf-8", newline="") as fh:
            for r in csv.DictReader(fh):
                if r["technology"] in TECHS and r["year"] == "2025":
                    if r["parameter"] == "CapacityFactor":
                        cls.cf[(r["technology"], r["timeslice"])] = float(r["required_root_value"])
                    elif r["parameter"] == "TotalTechnologyAnnualActivityLowerLimit":
                        cls.floor[r["technology"]] = float(r["required_root_value"])
        wb = load_workbook(TIMESLICES, read_only=True, data_only=True)
        rows = list(wb["YearSplit"].iter_rows(values_only=True))
        cls.split = {r[0]: float(r[3]) for r in rows[1:] if r[0]}
        wb.close()
        wb = load_workbook(WORKBOOK, read_only=True, data_only=True)
        rows = list(wb["Secondary_Techs"].iter_rows(values_only=True))
        hdr = list(rows[0])
        cls.values = {(r[hdr.index("Tech")], r[hdr.index("Parameter")]): dict(zip(hdr, r)) for r in rows[1:]}
        wb.close()

    def test_2025_fleet_reaches_the_floor(self):
        for tech in TECHS:
            capacity = self.values[(tech, "ResidualCapacity")][2025]
            availability = self.values.get((tech, "AvailabilityFactor"), {}).get(2025, 1.0)
            cf = sum(self.cf[(tech, ts)] * share for ts, share in self.split.items())
            most = capacity * 31.536 * cf * min(availability, 1.0)
            self.assertGreaterEqual(most, self.floor[tech], tech)


if __name__ == "__main__":
    unittest.main()
