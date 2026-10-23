"""The base-year build lock: plant-list builds only in the lock years.

The lock runs on the normalised CSVs of one scenario. These tests check that it
holds for every locked family, leaves the backstop, the internal grid and fuel
supply untouched, survives the reserve-margin repair, and that its two
pre-solve checks lower an unreachable activity floor and leave a link a trade
floor needs out of the lock.
"""
from __future__ import annotations

import csv
import io
import sys
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path
from unittest import mock

import yaml
from openpyxl import Workbook

from ostram.pipeline.execution import base_year_lock
from ostram.pipeline.execution.model_file_adjustments import reserve_margin_repair_xlsx


REPO_ROOT = Path(__file__).resolve().parents[2]
YEARS = range(2023, 2027)
GEN, STORAGE, LINK = "PWRCOAINDNO", "PWRLDSINDNO", "TRNINDNONPLXX"
OTHER_LINK = "TRNNPLXXBGDXX"
BACKSTOP, GRID, FUEL = "PWRBCKINDNO", "PWRTRNINDNO", "MINCOAIND"
TECHS = (GEN, STORAGE, LINK, OTHER_LINK, BACKSTOP, GRID, FUEL)


def _write(folder: Path, name: str, header: list[str], rows: list[list]) -> None:
    with (folder / f"{name}.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(header)
        writer.writerows(rows)


def _read(folder: Path, name: str) -> dict:
    with (folder / f"{name}.csv").open(encoding="utf-8") as handle:
        return {(r["TECHNOLOGY"], int(r["YEAR"])): float(r["VALUE"]) for r in csv.DictReader(handle)}


def _scenario(folder: Path, *, lower: float = 10.0, ceiling: float = 9999.0) -> None:
    """A small scenario: every technology open (9999) in every year, one coal
    plant in the 2024 plant list, a residual coal fleet of 1 GW, and one
    interconnector the 2025 trade floor needs."""
    years = list(range(2023, 2029))
    _write(folder, "TECHNOLOGY", ["VALUE"], [[t] for t in TECHS])
    _write(folder, "YearSplit", ["TIMESLICE", "YEAR", "VALUE"], [["S1", y, 1.0] for y in years])
    _write(folder, "CapacityToActivityUnit", ["REGION", "TECHNOLOGY", "VALUE"],
           [["GLOBAL", t, 31.536] for t in TECHS])
    _write(folder, "OperationalLife", ["REGION", "TECHNOLOGY", "VALUE"], [["GLOBAL", t, 30] for t in TECHS])
    _write(folder, "CapacityFactor", ["REGION", "TECHNOLOGY", "TIMESLICE", "YEAR", "VALUE"],
           [["GLOBAL", GEN, "S1", y, 0.5] for y in years])
    _write(folder, "AvailabilityFactor", ["REGION", "TECHNOLOGY", "YEAR", "VALUE"], [])
    _write(folder, "ResidualCapacity", ["REGION", "TECHNOLOGY", "YEAR", "VALUE"],
           [["GLOBAL", GEN, y, 1.0] for y in years] + [["GLOBAL", LINK, y, 0.5] for y in years])
    _write(folder, "TotalAnnualMinCapacityInvestment", ["REGION", "TECHNOLOGY", "YEAR", "VALUE"],
           [["GLOBAL", GEN, 2023, ""], ["GLOBAL", GEN, 2024, 0.25], ["GLOBAL", STORAGE, 2026, 0.1],
            ["GLOBAL", LINK, 2025, 0.2]])
    _write(folder, "TotalAnnualMaxCapacityInvestment", ["REGION", "TECHNOLOGY", "YEAR", "VALUE"],
           [["GLOBAL", t, y, ceiling if t == GEN and y == 2026 else 9999.0] for t in TECHS for y in years])
    _write(folder, "TotalTechnologyAnnualActivityLowerLimit", ["REGION", "TECHNOLOGY", "YEAR", "VALUE"],
           [["GLOBAL", GEN, 2025, lower]])
    _write(folder, "InputActivityRatio", ["REGION", "TECHNOLOGY", "FUEL", "MODE_OF_OPERATION", "YEAR", "VALUE"], [])
    _write(folder, "OutputActivityRatio", ["REGION", "TECHNOLOGY", "FUEL", "MODE_OF_OPERATION", "YEAR", "VALUE"],
           [["GLOBAL", OTHER_LINK, "ELCBGDXX01", 1, 2025, 0.98]])


MODEL = (
    's.t. D4_I1_TRADE_047_NPL_BGD_EXPORT_2025_Lower{r in REGION: D4HistoricalTradeActive = 1}: '
    'sum{l in TIMESLICE} (RateOfActivity[r,l,"TRNNPLXXBGDXX",1,2025]*YearSplit[l,2025]*'
    'OutputActivityRatio[r,"TRNNPLXXBGDXX","ELCBGDXX01",1,2025]) >= 0.05;\n'
)


def _lock(folder: Path, *, trade_active: bool = True) -> base_year_lock.LockReport:
    return base_year_lock.apply_base_year_lock(folder, "T_Lock", YEARS, trade_active=trade_active, model_text=MODEL)


class BaseYearLockTests(unittest.TestCase):
    def test_tracked_configuration_locks_2023_to_2026(self) -> None:
        params = yaml.safe_load((REPO_ROOT / "config/execution/Config_MOMF_T1_AB.yaml").read_text(encoding="utf-8"))
        self.assertEqual(base_year_lock.lock_years(params), range(2023, 2027))
        self.assertIsNone(base_year_lock.lock_years({}))

    def test_locked_families_follow_the_taxonomy(self) -> None:
        locked = base_year_lock.locked_technologies()
        self.assertEqual({locked[GEN], locked[STORAGE], locked[LINK]},
                         {"GENERATION", "STORAGE_LONG", "INTERCONNECTORS"})
        for tech in (BACKSTOP, GRID, "RNWTRNINDNO", "TRNNLIINDNO", FUEL, "RNWSPVINDNO"):
            self.assertNotIn(tech, locked)

    def test_lock_holds_for_every_locked_family(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            _scenario(folder)
            _lock(folder, trade_active=False)
            cap = _read(folder, "TotalAnnualMaxCapacityInvestment")
            plant = {(GEN, 2024): 0.25, (STORAGE, 2026): 0.1, (LINK, 2025): 0.2}
            for tech in (GEN, STORAGE, LINK, OTHER_LINK):
                for year in YEARS:
                    self.assertEqual(cap[(tech, year)], plant.get((tech, year), 0.0), (tech, year))
                for year in (2027, 2028):
                    self.assertEqual(cap[(tech, year)], 9999.0, (tech, year))

    def test_backstop_internal_grid_and_fuel_supply_are_untouched(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            _scenario(folder)
            before = _read(folder, "TotalAnnualMaxCapacityInvestment")
            _lock(folder)
            after = _read(folder, "TotalAnnualMaxCapacityInvestment")
            for tech in (BACKSTOP, GRID, FUEL):
                self.assertEqual({k: v for k, v in after.items() if k[0] == tech},
                                 {k: v for k, v in before.items() if k[0] == tech})

    def test_a_lower_ceiling_already_there_stays(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            _scenario(folder, ceiling=0.0)
            _lock(folder)
            self.assertEqual(_read(folder, "TotalAnnualMaxCapacityInvestment")[(GEN, 2026)], 0.0)

    def test_unreachable_activity_floor_is_lowered_to_the_locked_fleet(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            _scenario(folder, lower=50.0)
            report = _lock(folder)
            fleet = (1.0 + 0.25) * 31.536 * 0.5
            self.assertEqual([(GEN, 2025)], [(t, y) for t, y, _, _ in report.gaps])
            self.assertAlmostEqual(_read(folder, "TotalTechnologyAnnualActivityLowerLimit")[(GEN, 2025)], fleet)
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            _scenario(folder, lower=10.0)
            self.assertEqual(_lock(folder).gaps, [])
            self.assertEqual(_read(folder, "TotalTechnologyAnnualActivityLowerLimit")[(GEN, 2025)], 10.0)

    def test_link_a_trade_floor_needs_stays_out_of_the_lock(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            _scenario(folder)
            report = _lock(folder)
            self.assertEqual([OTHER_LINK], [link for _, _, links, _, _ in report.exempt for link in links])
            cap = _read(folder, "TotalAnnualMaxCapacityInvestment")
            for year in YEARS:
                self.assertEqual(cap[(OTHER_LINK, year)], 9999.0)
                self.assertEqual(cap[(LINK, year)], 0.2 if year == 2025 else 0.0)

    def test_reserve_margin_repair_does_not_reopen_a_locked_cell(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "fallbacks"
            sheet.append(["CR", "TECH_PREFIX", "YEAR", "TotalAnnualMaxCapacity", "TotalAnnualMaxCapacityInvestment"])
            sheet.append(["INDNO", "PWRNGS", "*", 5, 0.5])
            workbook.save(root / "fallbacks.xlsx")
            workbook.close()
            (root / "input.txt").write_text(
                "set YEAR := 2026 2027 ;\nset TECHNOLOGY := PWRNGSINDNO ;\n"
                "param default 0 : ResidualCapacity :=\nGLOBAL PWRNGSINDNO 2026 1\nGLOBAL PWRNGSINDNO 2027 1\n;\n"
                "param default 0 : TotalAnnualMinCapacity :=\n;\n"
                "param default 0 : TotalAnnualMinCapacityInvestment :=\n;\n"
                "param default -1 : TotalAnnualMaxCapacity :=\nGLOBAL PWRNGSINDNO 2026 9999\nGLOBAL PWRNGSINDNO 2027 9999\n;\n"
                "param default -1 : TotalAnnualMaxCapacityInvestment :=\nGLOBAL PWRNGSINDNO 2026 0\nGLOBAL PWRNGSINDNO 2027 0\n;\n",
                encoding="utf-8",
            )
            argv = ["reserve_margin_repair_xlsx", str(root / "input.txt"), "-o", str(root / "out.txt"),
                    "--fallback-xlsx", str(root / "fallbacks.xlsx"), "--xlsx-sheet", "fallbacks",
                    "--locked-years", "2023", "2026"]
            opened = []
            real_load = reserve_margin_repair_xlsx.load_workbook

            def tracked_load(*args, **kwargs):
                opened.append(real_load(*args, **kwargs))
                return opened[-1]

            try:
                with (
                    mock.patch.object(sys, "argv", argv),
                    mock.patch.object(reserve_margin_repair_xlsx, "load_workbook", side_effect=tracked_load),
                    redirect_stdout(io.StringIO()),
                    redirect_stderr(io.StringIO()),
                ):
                    self.assertEqual(reserve_margin_repair_xlsx.main(), 0)
            finally:
                for book in opened:
                    book.close()
            text = (root / "out.txt").read_text(encoding="utf-8")
            block = text.split("TotalAnnualMaxCapacityInvestment :=")[1].split(";")[0].split()
            self.assertEqual(block, ["GLOBAL", "PWRNGSINDNO", "2026", "0", "GLOBAL", "PWRNGSINDNO", "2027", "0.5"])


if __name__ == "__main__":
    unittest.main()
