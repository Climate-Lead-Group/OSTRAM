"""Interconnector_Params holds only parameters the model receives.

The interconnectors take capacity factor 1 and the OSeMOSYS Global losses,
so the sheet has no CapacityFactor or TransmissionLossFactor rows. Every
cell reference in the Transmission_Cost_Basis summary points at the row of
the same technology and parameter.
"""
from __future__ import annotations

import re
import unittest
from pathlib import Path

from openpyxl import load_workbook


REPO_ROOT = Path(__file__).resolve().parents[2]
WORKBOOK = REPO_ROOT / "inputs" / "scenarios" / "OSTRAM_Scenario_Inputs.xlsx"
REFERENCE = re.compile(r"Interconnector_Params!J(\d+):AK(\d+)")


class InterconnectorParamsRowsTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
        try:
            cls.params = list(workbook["Interconnector_Params"].iter_rows(values_only=True))
            cls.basis = list(workbook["Transmission_Cost_Basis"].iter_rows(min_row=1, max_row=19, values_only=True))
        finally:
            workbook.close()
        cls.header = {name: i for i, name in enumerate(cls.params[0]) if isinstance(name, str)}

    def test_no_capacity_factor_or_loss_rows(self) -> None:
        parameter = self.header["Parameter"]
        found = [row[parameter] for row in self.params[1:]
                 if row[parameter] in ("CapacityFactor", "TransmissionLossFactor")]
        self.assertEqual(found, [])

    def test_cost_basis_references_name_the_right_rows(self) -> None:
        tech, parameter = self.header["Tech"], self.header["Parameter"]
        columns = {name: i for i, name in enumerate(self.basis[0])}
        for row in self.basis[1:]:
            for column, expected in (("Capital_cells", "CapitalCost"), ("Fixed_cells", "FixedCost")):
                match = REFERENCE.fullmatch(row[columns[column]])
                self.assertIsNotNone(match, row[columns[column]])
                target = self.params[int(match.group(1)) - 1]
                self.assertEqual((target[tech], target[parameter]), (row[columns["Technology"]], expected))


if __name__ == "__main__":
    unittest.main()
