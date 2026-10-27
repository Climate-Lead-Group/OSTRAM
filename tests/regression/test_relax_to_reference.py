"""B_Optimised_VRE takes the looser of its own bounds and those of A_Calibrated_BAU.

The rule runs every time B is built, on the A3 output of A built in the same
run: an upper bound of B is never below A's, a lower bound never above A's.
"""

from __future__ import annotations

import os
import shutil
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from ostram.pipeline.scenarios.materializer import (
    MaterializationPaths,
    REQUIRED_AO_FILES,
    materialize_scenarios,
)
from ostram.pipeline.scenarios.registry import load_registry
from ostram.pipeline.scenarios.rules import relax_to_reference as rule


REPO_ROOT = Path(__file__).resolve().parents[2]
HEADER = ["Tech.ID", "Tech", "Tech.Name", "Parameter.ID", "Parameter", "Unit",
          "Projection.Mode", "Projection.Parameter", 2023, 2024]


def _workbook(folder: Path, rows: list[list]) -> Path:
    """Write A-O_Parametrization.xlsx with one 'Secondary Techs' sheet."""
    folder.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Secondary Techs"
    ws.append(HEADER)
    for row in rows:
        ws.append(row)
    path = folder / rule.PARAM_FILE_NAME
    wb.save(path)
    return path


def _row(tech: str, parameter: str, mode: str, v2023, v2024) -> list:
    return [1, tech, tech, 1, parameter, "GW", mode, None, v2023, v2024]


def _values(path: Path) -> dict[tuple[str, str], tuple]:
    ws = load_workbook(path)["Secondary Techs"]
    return {(r[1], r[4]): (r[6], r[8], r[9]) for r in ws.iter_rows(min_row=2, values_only=True)}


class RelaxToReferenceRuleTests(unittest.TestCase):
    def setUp(self) -> None:
        self._temp = tempfile.TemporaryDirectory()
        self.root = Path(self._temp.name)

    def tearDown(self) -> None:
        self._temp.cleanup()

    def test_tighter_bounds_end_at_the_reference_values(self) -> None:
        reference = self.root / "A"
        target = self.root / "B"
        _workbook(reference, [
            _row("PWRCOAINDNO", "TotalAnnualMaxCapacityInvestment", "User defined", 2.0, 3.0),
            _row("PWRCOAINDNO", "TotalAnnualMinCapacityInvestment", "User defined", 0.2, 0.0),
            _row("PWRNGSINDNO", "TotalAnnualMaxCapacityInvestment", "User defined", -1, 4.0),
            _row("PWRSDSINDNO", "TotalAnnualMinCapacityInvestment", "EMPTY", None, None),
            _row("PWRCOAINDNO", "CapitalCost", "User defined", 1000.0, 1000.0),
        ])
        path = _workbook(target, [
            _row("PWRCOAINDNO", "TotalAnnualMaxCapacityInvestment", "User defined", 1.0, 3.0),
            _row("PWRCOAINDNO", "TotalAnnualMinCapacityInvestment", "User defined", 0.5, 0.0),
            _row("PWRNGSINDNO", "TotalAnnualMaxCapacityInvestment", "User defined", 5.0, 2.5),
            _row("PWRSDSINDNO", "TotalAnnualMinCapacityInvestment", "User defined", 0.0, 0.1),
            _row("PWRCOAINDNO", "CapitalCost", "User defined", 900.0, 900.0),
        ])

        log = rule.run(target, reference)

        values = _values(path)
        self.assertEqual(values[("PWRCOAINDNO", "TotalAnnualMaxCapacityInvestment")][1:], (2.0, 3.0))
        self.assertEqual(values[("PWRCOAINDNO", "TotalAnnualMinCapacityInvestment")][1:], (0.2, 0.0))
        # -1 is no bound, so it is the looser value.
        self.assertEqual(values[("PWRNGSINDNO", "TotalAnnualMaxCapacityInvestment")][1:], (-1, 4.0))
        # A has no storage floor: B's floor row is no longer exported.
        self.assertEqual(values[("PWRSDSINDNO", "TotalAnnualMinCapacityInvestment")][0], "EMPTY")
        # A parameter that is not a bound is left as it is.
        self.assertEqual(values[("PWRCOAINDNO", "CapitalCost")][1:], (900.0, 900.0))
        self.assertEqual(len(log["changes"]), 5)

    def test_looser_bounds_stay_unchanged(self) -> None:
        reference = self.root / "A"
        target = self.root / "B"
        _workbook(reference, [
            _row("PWRCOAINDNO", "TotalAnnualMaxCapacityInvestment", "User defined", 2.0, 3.0),
            _row("PWRCOAINDNO", "TotalAnnualMinCapacityInvestment", "User defined", 0.2, 0.3),
            _row("PWRWONINDNO", "TotalAnnualMaxCapacity", "EMPTY", None, None),
        ])
        path = _workbook(target, [
            _row("PWRCOAINDNO", "TotalAnnualMaxCapacityInvestment", "User defined", 2.5, -1),
            _row("PWRCOAINDNO", "TotalAnnualMinCapacityInvestment", "User defined", 0.1, 0.3),
            _row("PWRWONINDNO", "TotalAnnualMaxCapacity", "EMPTY", None, None),
        ])
        before = path.read_bytes()

        log = rule.run(target, reference)

        self.assertEqual(log["changes"], [])
        self.assertEqual(path.read_bytes(), before)

    def test_a_blank_cell_that_differs_stops_the_rule(self) -> None:
        reference = self.root / "A"
        target = self.root / "B"
        _workbook(reference, [
            _row("PWRCOAINDNO", "TotalTechnologyAnnualActivityUpperLimit", "User defined", 10.0, None),
        ])
        _workbook(target, [
            _row("PWRCOAINDNO", "TotalTechnologyAnnualActivityUpperLimit", "User defined", 10.0, 8.0),
        ])
        with self.assertRaises(rule.RelaxationError):
            rule.run(target, reference)


class RelaxationDependencyTests(unittest.TestCase):
    def test_b_declares_a_as_its_reference_and_building_b_builds_a_first(self) -> None:
        registry = load_registry()
        roots = registry.roots_by_name
        self.assertEqual(roots["B_Optimised_VRE"].relaxation_dependencies, ("A_Calibrated_BAU",))
        for name in ("BAU", "A_Calibrated_BAU", "C_Target_VRE"):
            self.assertEqual(roots[name].relaxation_dependencies, ())
        self.assertEqual(
            registry.required_roots(registry.select("B_Optimised_VRE")),
            ("BAU", "A_Calibrated_BAU", "B_Optimised_VRE"),
        )
        # A derived B scenario inherits B's limits, so it needs the same roots.
        self.assertEqual(
            registry.required_roots(registry.select("B_Opt_SolarCapex130")),
            ("BAU", "A_Calibrated_BAU", "B_Optimised_VRE"),
        )

    def test_materialization_relaxes_b_to_the_a_of_the_same_run(self) -> None:
        registry = load_registry()
        with tempfile.TemporaryDirectory() as temp:
            outputs = Path(temp) / "A1_Outputs"
            for root in ("BAU", "A_Calibrated_BAU", "B_Optimised_VRE"):
                snapshot = outputs / f"_post_a2_snapshot_{root}"
                snapshot.mkdir(parents=True)
                for filename in REQUIRED_AO_FILES:
                    (snapshot / filename).write_bytes(b"fixture")
            paths = MaterializationPaths(
                preparation_workspace=Path(temp) / "preparation",
                a1_outputs=outputs,
                a3_entrypoint=REPO_ROOT / "ostram" / "pipeline" / "scenarios" / "transform.py",
                a3_process=REPO_ROOT / "ostram" / "pipeline" / "scenarios" / "transformations",
                soasia=REPO_ROOT / "inputs" / "scenarios" / "OSTRAM_Scenario_Inputs.xlsx",
            )
            events: list[tuple] = []

            def materialize_root(root: str, environment: dict[str, str]) -> None:
                del environment
                events.append(("root", root))
                shutil.copytree(outputs / f"_post_a2_snapshot_{root}", outputs / f"A1_Outputs_{root}")

            def relax(target: Path, reference: Path) -> dict:
                events.append(("relax", target.name, reference.name))
                return {"changes": [{}, {}]}

            record = materialize_scenarios(
                ("B_Optimised_VRE",),
                paths=paths,
                registry=registry,
                environment=dict(os.environ),
                root_materializer=materialize_root,
                direction_applier=lambda *_: {},
                relaxation_applier=relax,
                write_provenance=False,
            )

        self.assertEqual(events, [
            ("root", "BAU"),
            ("root", "A_Calibrated_BAU"),
            ("root", "B_Optimised_VRE"),
            ("relax", "A1_Outputs_B_Optimised_VRE", "A1_Outputs_A_Calibrated_BAU"),
        ])
        self.assertEqual(record["selected_scenarios"], ["B_Optimised_VRE"])
        self.assertEqual(
            record["roots"][2]["relaxation"],
            {"reference": "A_Calibrated_BAU", "changes": 2},
        )


if __name__ == "__main__":
    unittest.main()
