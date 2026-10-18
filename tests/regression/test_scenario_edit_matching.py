"""A scenario edit or VRE ceiling that matches nothing stops the run."""

from __future__ import annotations

import contextlib
import io
from pathlib import Path
import tempfile
import unittest
from unittest import mock

from openpyxl import Workbook, load_workbook
import yaml

from ostram.pipeline.scenarios import scenario_edits as edits


YEARS = list(range(2023, 2051))
HEADER = [
    "Tech.ID", "Tech", "Tech.Name", "Parameter.ID", "Parameter", "Unit",
    "Projection.Mode", "Projection.Parameter", *YEARS,
]
SCENARIO = "B_Opt_MyTest"


def _workbook():
    workbook = Workbook()
    secondary = workbook.active
    secondary.title = "Secondary Techs"
    secondary.append(HEADER)
    secondary.append([
        1, "PWRSOLBGDXX", "solar", 1, "TotalAnnualMaxCapacity", "GW",
        "User defined", 0, *[5.0] * len(YEARS),
    ])
    secondary.append([
        1, "PWRSOLBGDXX", "solar", 2, "CapitalCost", "USD", "User defined", 0,
        *[100.0] * len(YEARS),
    ])
    secondary.append([
        2, "PWRSPVBGDXX", "solar", 1, "TotalAnnualMaxCapacity", "GW",
        "User defined", 0, *[40.0] * len(YEARS),
    ])
    demand = workbook.create_sheet("Demand Techs")
    demand.append(HEADER)
    demand.append([
        3, "TRNNLIBGDXX", "backstop", 1, "ResidualCapacity", "GW",
        "User defined", 0, *[1.0] * len(YEARS),
    ])
    return workbook


def _log():
    return {"scenario": SCENARIO, "cells": [], "rows_created": [], "skipped": []}


def _apply(workbook, edit):
    sheet = workbook[edit["sheet"]]
    columns, years = edits.scan_columns(sheet)
    log = _log()
    edits.apply_edit(sheet, columns, years, edit, log)
    return log


class UnmatchedScenarioEditTests(unittest.TestCase):
    def assert_stops(self, edit, message):
        with self.assertRaises(edits.ScenarioEditError) as caught:
            _apply(_workbook(), edit)
        self.assertEqual(str(caught.exception), message)

    def test_misspelt_technology_stops_with_a_suggestion(self) -> None:
        self.assert_stops(
            {
                "sheet": "Secondary Techs", "tech": "PWRSOLBGXX",
                "param": "TotalAnnualMaxCapacity", "op": "set_flat", "value": 1.0,
            },
            f'Scenario {SCENARIO}: the edit to TotalAnnualMaxCapacity for '
            '"PWRSOLBGXX" changed nothing. No technology has that name. '
            'Did you mean PWRSOLBGDXX?',
        )

    def test_prefix_that_matches_nothing_stops(self) -> None:
        self.assert_stops(
            {
                "sheet": "Secondary Techs", "tech_prefix": "PWRSOLX",
                "param": "CapitalCost", "op": "multiply", "factor": 1.1,
            },
            f'Scenario {SCENARIO}: the edit to CapitalCost for "PWRSOLX" '
            'changed nothing. No technology starts with "PWRSOLX".',
        )

    def test_misspelt_parameter_stops_with_a_suggestion(self) -> None:
        self.assert_stops(
            {
                "sheet": "Secondary Techs", "tech": "PWRSOLBGDXX",
                "param": "TotalAnualMaxCapacity", "op": "set_flat", "value": 1.0,
            },
            f'Scenario {SCENARIO}: the edit to TotalAnualMaxCapacity for '
            '"PWRSOLBGDXX" changed nothing. Sheet "Secondary Techs" has no '
            'parameter "TotalAnualMaxCapacity". Did you mean '
            'TotalAnnualMaxCapacity?',
        )

    def test_create_if_absent_with_an_unknown_technology_stops(self) -> None:
        self.assert_stops(
            {
                "sheet": "Demand Techs", "tech": "TRNNLIBGXX",
                "param": "TotalAnnualMaxCapacity", "values": {"2030": 0.0},
                "create_if_absent": True,
            },
            f'Scenario {SCENARIO}: the edit to TotalAnnualMaxCapacity for '
            '"TRNNLIBGXX" changed nothing. No technology has that name. '
            'Did you mean TRNNLIBGDXX?',
        )

    def test_years_filter_with_no_matching_column_stops(self) -> None:
        self.assert_stops(
            {
                "sheet": "Secondary Techs", "tech": "PWRSOLBGDXX",
                "param": "CapitalCost", "op": "multiply", "factor": 1.5,
                "years": [2061, 2062],
            },
            f'Scenario {SCENARIO}: the edit to CapitalCost for "PWRSOLBGDXX" '
            'changed nothing. None of the years 2061, 2062 is a column in '
            'the workbook.',
        )

    def test_years_filter_with_one_missing_column_stops(self) -> None:
        self.assert_stops(
            {
                "sheet": "Secondary Techs", "tech": "PWRSOLBGDXX",
                "param": "CapitalCost", "op": "multiply", "factor": 1.5,
                "years": [2030, 2061],
            },
            f'Scenario {SCENARIO}: the edit to CapitalCost for "PWRSOLBGDXX" '
            'changed nothing. None of the years 2061 is a column in '
            'the workbook.',
        )

    def test_values_year_with_no_column_stops(self) -> None:
        self.assert_stops(
            {
                "sheet": "Secondary Techs", "tech": "PWRSOLBGDXX",
                "param": "CapitalCost", "values": {"2030": 90.0, "2061": 80.0},
            },
            f'Scenario {SCENARIO}: the edit to CapitalCost for "PWRSOLBGDXX" '
            'changed nothing. None of the years 2061 is a column in '
            'the workbook.',
        )

    def test_vre_ceiling_that_matches_nothing_stops(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            ceilings = Path(temp) / "south_asia.yaml"
            ceilings.write_text(
                "vre_ceilings:\n"
                "  sheet: Secondary Techs\n"
                "  param: TotalAnnualMaxCapacity\n"
                "  ceilings_gw:\n"
                "    PWRSPVBGXX: 1.0\n",
                encoding="utf-8",
            )
            with self.assertRaises(edits.ScenarioEditError) as caught:
                edits.apply_ceiling_layer(_workbook(), _log(), ceiling_path=ceilings)
        self.assertEqual(
            str(caught.exception),
            f'Scenario {SCENARIO}: the VRE ceiling for "PWRSPVBGXX" changed '
            'nothing. No technology has that name. Did you mean PWRSPVBGDXX?',
        )

    def test_edit_that_matches_rows_but_changes_no_value_passes(self) -> None:
        log = _apply(
            _workbook(),
            {
                "sheet": "Secondary Techs", "tech": "PWRSOLBGDXX",
                "param": "TotalAnnualMaxCapacity", "op": "set_flat", "value": 5.0,
            },
        )
        self.assertEqual(log["cells"], [])
        self.assertEqual(log["skipped"], [])

    def test_create_if_absent_with_known_names_creates_the_row(self) -> None:
        log = _apply(
            _workbook(),
            {
                "sheet": "Demand Techs", "tech": "TRNNLIBGDXX",
                "param": "TotalAnnualMaxCapacity", "values": {"2030": 0.0},
                "create_if_absent": True,
            },
        )
        self.assertEqual(
            [(row["tech"], row["param"]) for row in log["rows_created"]],
            [("TRNNLIBGDXX", "TotalAnnualMaxCapacity")],
        )


class CopyFromTechTests(unittest.TestCase):
    def test_copy_sets_every_year_to_the_source_row(self) -> None:
        workbook = _workbook()
        source = [100.0 + year - 2023 for year in YEARS]
        workbook["Secondary Techs"].append([
            4, "PWRSOLINDNO", "solar", 2, "CapitalCost", "USD", "User defined", 0,
            *source,
        ])
        log = _apply(
            workbook,
            {
                "sheet": "Secondary Techs", "techs": ["PWRSOLBGDXX"],
                "param": "CapitalCost", "op": "copy_from_tech",
                "source_tech": "PWRSOLINDNO",
            },
        )
        sheet = workbook["Secondary Techs"]
        columns, years = edits.scan_columns(sheet)
        row = edits.find_row(sheet, columns, "PWRSOLBGDXX", "CapitalCost")
        self.assertEqual(
            [sheet.cell(row, years[year]).value for year in YEARS], source
        )
        self.assertEqual(len(log["cells"]), len(YEARS) - 1)

    def test_missing_source_row_stops(self) -> None:
        with self.assertRaises(edits.ScenarioEditError) as caught:
            _apply(
                _workbook(),
                {
                    "sheet": "Secondary Techs", "tech": "PWRSOLBGDXX",
                    "param": "CapitalCost", "op": "copy_from_tech",
                    "source_tech": "PWRSOLINDNO",
                },
            )
        self.assertEqual(
            str(caught.exception),
            f'Scenario {SCENARIO}: the edit to CapitalCost for "PWRSOLBGDXX" '
            'has no source row. Sheet "Secondary Techs" has no CapitalCost row '
            "for PWRSOLINDNO.",
        )


class ReuseEditsFromTests(unittest.TestCase):
    """``reuse_edits_from`` applies another derived scenario's edits first."""

    def _build(self, scenarios, name):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            outputs, configs = root / "A1_Outputs", root / "configs"
            for base in ("B_Optimised_VRE", "A_Calibrated_BAU"):
                (outputs / f"A1_Outputs_{base}").mkdir(parents=True)
                _workbook().save(outputs / f"A1_Outputs_{base}" / edits.PARAM_FILE)
            (configs / "B_Optimised_VRE").mkdir(parents=True)
            for scenario, body in scenarios.items():
                (configs / scenario).mkdir(parents=True)
                (configs / scenario / "scenario.yaml").write_text(
                    yaml.safe_dump({"scenario": scenario, "apply_vre_ceiling_layer": False, **body}),
                    encoding="utf-8",
                )
            target = outputs / f"A1_Outputs_{name}"
            try:
                with mock.patch(
                    "ostram.pipeline.scenarios.rules.apply_base_year_pin.apply_pin_rules",
                    return_value={},
                ):
                    log = edits.build_scenario(
                        name, a1_outputs=outputs, configs=configs,
                        authority_path=root / "unused.xlsx",
                    )
            except edits.ScenarioEditError:
                self.assertFalse(target.exists())
                raise
            sheet = load_workbook(target / edits.PARAM_FILE)["Secondary Techs"]
            columns, years = edits.scan_columns(sheet)
            values = {
                param: sheet.cell(
                    edits.find_row(sheet, columns, "PWRSOLBGDXX", param), years[2030]
                ).value
                for param in ("CapitalCost", "TotalAnnualMaxCapacity")
            }
            return log, values

    COST = {
        "base_scenario": "B_Optimised_VRE",
        "edits": [{
            "sheet": "Secondary Techs", "tech": "PWRSOLBGDXX",
            "param": "CapitalCost", "op": "set_flat", "value": 7.0,
        }],
    }
    OWN = [{
        "sheet": "Secondary Techs", "tech": "PWRSOLBGDXX",
        "param": "TotalAnnualMaxCapacity", "op": "set_flat", "value": 3.0,
    }]

    def assert_refused(self, scenarios, message):
        with self.assertRaises(edits.ScenarioEditError) as caught:
            self._build(scenarios, "B_Opt_Both")
        self.assertIn(message, str(caught.exception))

    def test_reused_edits_are_applied_before_the_own_edits(self) -> None:
        log, values = self._build(
            {
                "B_Opt_Cost": self.COST,
                "B_Opt_Both": {
                    "base_scenario": "B_Optimised_VRE",
                    "reuse_edits_from": "B_Opt_Cost", "edits": self.OWN,
                },
            },
            "B_Opt_Both",
        )
        self.assertEqual(values, {"CapitalCost": 7.0, "TotalAnnualMaxCapacity": 3.0})
        self.assertEqual(
            [cell["param"] for cell in log["cells"]],
            ["CapitalCost"] * len(YEARS) + ["TotalAnnualMaxCapacity"] * len(YEARS),
        )
        self.assertEqual(log["reused_edits_from"], {"scenario": "B_Opt_Cost", "edits": 1})

    def test_a_scenario_that_is_not_derived_is_refused(self) -> None:
        self.assert_refused(
            {"B_Opt_Both": {
                "base_scenario": "B_Optimised_VRE",
                "reuse_edits_from": "B_Optimised_VRE", "edits": self.OWN,
            }},
            "reuse_edits_from names B_Optimised_VRE, which is not a derived scenario",
        )

    def test_a_scenario_with_another_base_is_refused(self) -> None:
        self.assert_refused(
            {
                "B_Opt_Cost": {**self.COST, "base_scenario": "A_Calibrated_BAU"},
                "B_Opt_Both": {
                    "base_scenario": "B_Optimised_VRE",
                    "reuse_edits_from": "B_Opt_Cost", "edits": self.OWN,
                },
            },
            "reuse_edits_from names B_Opt_Cost, whose base_scenario A_Calibrated_BAU "
            "is not B_Optimised_VRE.",
        )

    def test_a_scenario_that_itself_reuses_edits_is_refused(self) -> None:
        self.assert_refused(
            {
                "B_Opt_Cost": self.COST,
                "B_Opt_Middle": {
                    "base_scenario": "B_Optimised_VRE",
                    "reuse_edits_from": "B_Opt_Cost", "edits": [],
                },
                "B_Opt_Both": {
                    "base_scenario": "B_Optimised_VRE",
                    "reuse_edits_from": "B_Opt_Middle", "edits": self.OWN,
                },
            },
            "reuse_edits_from names B_Opt_Middle, which itself uses reuse_edits_from.",
        )


class StandaloneStopMessageTests(unittest.TestCase):
    def test_main_prints_the_stop_message_without_a_traceback(self) -> None:
        message = f"Scenario {SCENARIO}: the edit changed nothing."
        stderr = io.StringIO()
        with (
            mock.patch.object(
                edits, "build_scenario",
                side_effect=edits.ScenarioEditError(message),
            ),
            mock.patch("sys.argv", ["scenario_edits", "--scenario", SCENARIO]),
            contextlib.redirect_stderr(stderr),
        ):
            code = edits.main()
        self.assertEqual(code, 1)
        self.assertEqual(stderr.getvalue(), message + "\n")
        self.assertNotIn("Traceback", stderr.getvalue())


if __name__ == "__main__":
    unittest.main()
