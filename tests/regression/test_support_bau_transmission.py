"""Commissioning, retirement and inheritance guards for support BAU."""
import os
from pathlib import Path
import tempfile
import unittest
from unittest import mock

from openpyxl import Workbook, load_workbook

from ostram.pipeline.scenarios import transform
from ostram.pipeline.scenarios.transformations import support_bau_transmission as policy
from regression.test_a3_orchestration import A3TraceHarness, _load_a3


class SupportBAUTransmissionTests(unittest.TestCase):
    def fixture(self, path):
        wb = Workbook()
        ws = wb.active
        ws.title = "Secondary Techs"
        ws.append(["Tech", "Parameter", "Projection.Mode", 2023, 2024, 2025, 2026, 2027])
        for tech in sorted(policy.TRN_TECHS):
            for parameter, values in (
                ("ResidualCapacity", [5, 4, 3, 2, 1]),
                ("TotalAnnualMinCapacityInvestment", [0, 2, 0, 3, 0]),
                ("TotalAnnualMaxCapacity", [5, 4, 3, 2, 1]),
                ("TotalAnnualMaxCapacityInvestment", [9999] * 5),
                ("CapitalCost", [100] * 5),
            ):
                mode = "EMPTY" if parameter == "TotalAnnualMaxCapacity" else "User defined"
                ws.append([tech, parameter, mode, *values])
        ws.append(["PWRTEST", "TotalAnnualMaxCapacity", "EMPTY", 99, 98, 97, 96, 95])
        fixed = wb.create_sheet("Fixed Horizon Parameters")
        fixed.append(["Tech", "Parameter", "Value"])
        for tech in sorted(policy.TRN_TECHS):
            fixed.append([tech, "OperationalLife", 2])
        wb.save(path)
        wb.close()

    def values(self, path):
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            return {ws.title: list(ws.values) for ws in wb}
        finally:
            wb.close()

    def test_commissioning_expiry_and_unchanged_other_cells(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "A-O_Parametrization.xlsx"
            self.fixture(path)
            before = self.values(path)
            record = policy.apply_commitment_ceiling(path)
            after = self.values(path)
            for old, new in zip(before["Secondary Techs"], after["Secondary Techs"]):
                if old[0] in policy.TRN_TECHS and old[1] == "TotalAnnualMaxCapacity":
                    # A 2024 commitment lives in 2024/25, expires in 2026.
                    self.assertEqual(new[2], "User defined")
                    self.assertEqual(new[3:], (5, 6, 5, 5, 4))
                else:
                    self.assertEqual(old, new)
            self.assertEqual(before["Fixed Horizon Parameters"], after["Fixed Horizon Parameters"])
            self.assertEqual(len(record["cells"]), 18 * 5)
            self.assertTrue(all(row["previous"] == "EMPTY" and row["final"] == "User defined"
                                for row in record["projection_modes"]))
            policy.apply_commitment_ceiling(path)
            self.assertEqual(after, self.values(path))

    def test_invalid_lifetime_does_not_save_partial_changes(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "A-O_Parametrization.xlsx"
            self.fixture(path)
            wb = load_workbook(path)
            wb["Fixed Horizon Parameters"].cell(2, 3).value = 0
            wb.save(path)
            wb.close()
            before = path.read_bytes()
            with self.assertRaisesRegex(ValueError, "OperationalLife must be positive"):
                policy.apply_commitment_ceiling(path)
            self.assertEqual(before, path.read_bytes())

    def test_full_support_only(self):
        with mock.patch.object(policy, "apply_commitment_ceiling") as apply:
            with mock.patch.dict(os.environ, {"OSTRAM_PROFILE_POLICIES": '{"support_bau_committed_transmission": true}'}):
                for scenario in ("A_Calibrated_BAU", "B_Optimised_VRE", "C_Target_VRE"):
                    transform.finalize_support_bau(Path("output"), scenario)
                apply.assert_not_called()
                transform.finalize_support_bau(Path("output"), "BAU")
                apply.assert_called_once_with(Path("output/A-O_Parametrization.xlsx"))
            apply.reset_mock()
            with mock.patch.dict(os.environ, {"OSTRAM_PROFILE_POLICIES": '{}'}):
                transform.finalize_support_bau(Path("output"), "BAU")
            apply.assert_not_called()

    def test_finalization_follows_inheritance_export_and_delivery(self):
        module = _load_a3("support_bau_finalization")
        harness = A3TraceHarness(module)
        try:
            def finalize(output, scenario):
                names = [event[0] for event in harness.events]
                self.assertLess(names.index("stage_6_persist_restrictions"), names.index("deliver_outputs"))
                self.assertEqual(names[-1], "deliver_outputs")
                self.assertEqual(output, harness.output_dir)
                self.assertEqual(scenario, "BAU")
            with mock.patch.object(module, "finalize_support_bau", side_effect=finalize) as final:
                self.assertEqual(harness.run(scenario="BAU"), 0)
            final.assert_called_once()
        finally:
            harness.close()
