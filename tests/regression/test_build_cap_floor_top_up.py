"""The opt-in floor top-up of build_cap_rule.py.

With ``floor_top_up: true`` every GENERATION technology with a generation floor
gets build caps large enough to reach it: need = 1.01 x floor / output per GW,
and the cap of a year is raised by need - residual - the caps through that year
when that is positive. Blank or -1 caps are unbounded and skipped. With the key
off, the rule is unchanged.
"""
from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import Workbook, load_workbook

from ostram.pipeline.scenarios.rules import build_cap_rule as rule


YEARS = (2023, 2024, 2025)
# 31.536 PJ per GW x (0.5 x 0.5 + 0.5 x 0.5) x 1.0 = 15.768 PJ per GW
OUT_PER_GW = 31.536 * 0.5


def _workbook(path: Path, techs: dict) -> Path:
    """Write a parametrization workbook. ``techs`` maps a tech to its rows:
    {parameter: [value per year] or None for an EMPTY row}, plus optional
    'cf' (per time slice and year) and 'cau'."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Secondary Techs"
    ws.append(["Tech.ID", "Tech", "Tech.Name", "Parameter.ID", "Parameter", "Unit",
               "Projection.Mode", "Projection.Parameter", *YEARS])
    fixed = wb.create_sheet(rule.FIXED_HORIZON_SHEET)
    fixed.append(["Tech.Type", "Tech.ID", "Tech", "Tech.Name", "Parameter.ID", "Parameter", "Unit", "Value"])
    caps = wb.create_sheet(rule.CAPACITY_FACTOR_SHEET)
    caps.append(["Timeslices", "Tech.ID", "Tech", "Tech.Name", "Parameter.ID", "Parameter", "Unit",
                 "Projection.Mode", "Projection.Parameter", *map(str, YEARS)])
    split = wb.create_sheet(rule.YEAR_SPLIT_SHEET)
    split.append(["Timeslices", "Parameter.ID", "Parameter", "Unit", "Projection.Mode",
                  "Projection.Parameter", *map(str, YEARS)])
    for slice_ in ("S1D1", "S1D2"):
        split.append([slice_, 1, "YearSplit", None, "User defined", 0, *[0.5] * len(YEARS)])
    for tech, params in techs.items():
        for param in (rule.RES_PARAM, rule.FLOOR_PARAM, rule.MIN_INV_PARAM, rule.AVAILABILITY_PARAM,
                      rule.MAX_INV_PARAM):
            values = params.get(param)
            mode = "EMPTY" if values is None else "User defined"
            ws.append([1, tech, tech, 1, param, None, mode, 0, *(values or [None] * len(YEARS))])
        fixed.append(["Secondary", 1, tech, tech, 1, rule.CAPACITY_TO_ACTIVITY_PARAM, None, params.get("cau", 31.536)])
        fixed.append(["Secondary", 1, tech, tech, 2, rule.OPERATIONAL_LIFE_PARAM, None, 25])
        for slice_ in params.get("cf", ("S1D1", "S1D2")):
            caps.append([slice_, 1, tech, tech, 13, rule.CAPACITY_FACTOR_PARAM, None, "User defined", 0,
                         *[0.5] * len(YEARS)])
    wb.save(path)
    return path


def _caps(path: Path, tech: str) -> list:
    wb = load_workbook(path, data_only=True)
    ws = wb["Secondary Techs"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[1] == tech and row[4] == rule.MAX_INV_PARAM:
            wb.close()
            return list(row[8:])
    wb.close()
    raise KeyError(tech)


class FloorTopUpTests(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name)

    def tearDown(self) -> None:
        self._tmp.cleanup()

    def _top_up(self, techs: dict, generation=None) -> tuple[Path, list]:
        path = _workbook(self.dir / "p.xlsx", techs)
        wb = load_workbook(path)
        changes = rule.apply_floor_top_up(wb, wb["Secondary Techs"], generation or set(techs))
        wb.save(path)
        wb.close()
        return path, changes

    def test_deficit_raises_one_cap_by_exactly_the_deficit(self) -> None:
        floor_2024 = 2 * OUT_PER_GW
        path, changes = self._top_up({"PWRSPVAAAXX": {
            rule.RES_PARAM: [1.0, 1.0, 1.0],
            rule.FLOOR_PARAM: [10.0, floor_2024, floor_2024],
            rule.MAX_INV_PARAM: [0.0, 0.5, 0.5]}})
        deficit = 1.01 * floor_2024 / OUT_PER_GW - 1.0 - 0.5
        self.assertEqual(_caps(path, "PWRSPVAAAXX"), [0.0, 0.5 + deficit, 0.5])
        self.assertEqual([(c["year"], c["reason"]) for c in changes], [(2024, "floor_top_up")])
        self.assertAlmostEqual(changes[0]["deficit"], deficit, places=12)

    def test_reachable_floor_changes_nothing(self) -> None:
        path, changes = self._top_up({"PWRSPVAAAXX": {
            rule.RES_PARAM: [1.0, 1.0, 1.0],
            rule.FLOOR_PARAM: [10.0, 10.0, 15.0],
            rule.MAX_INV_PARAM: [0.0, 0.0, 0.2]}})
        self.assertEqual(changes, [])
        self.assertEqual(_caps(path, "PWRSPVAAAXX"), [0.0, 0.0, 0.2])

    def test_unbounded_cap_is_skipped_that_year_and_after(self) -> None:
        for cap in ([0.0, None, 0.0], [0.0, -1, 0.0]):
            with self.subTest(cap=cap):
                path, changes = self._top_up({"PWRSPVAAAXX": {
                    rule.RES_PARAM: [0.0, 0.0, 0.0],
                    rule.FLOOR_PARAM: [0.0, 50.0, 50.0],
                    rule.MAX_INV_PARAM: cap}})
                self.assertEqual(changes, [])
                self.assertEqual(_caps(path, "PWRSPVAAAXX"), cap)

    def test_absent_factor_counts_as_one_and_non_generation_is_untouched(self) -> None:
        techs = {"PWRWONAAAXX": {rule.FLOOR_PARAM: [0.0, 0.0, 10.0], rule.MAX_INV_PARAM: [0.0, 0.0, 0.0],
                                 "cf": ()},
                 "PWRSDSAAAXX": {rule.FLOOR_PARAM: [5.0, 5.0, 5.0], rule.MAX_INV_PARAM: [0.0, 0.0, 0.0]}}
        path, changes = self._top_up(techs, generation={"PWRWONAAAXX"})
        self.assertAlmostEqual(_caps(path, "PWRWONAAAXX")[2], 1.01 * 10.0 / 31.536, places=12)
        self.assertEqual(_caps(path, "PWRSDSAAAXX"), [0.0, 0.0, 0.0])
        self.assertEqual({c["tech"] for c in changes}, {"PWRWONAAAXX"})

    def test_technology_outside_the_allowed_set_is_topped_up(self) -> None:
        path = _workbook(self.dir / "p.xlsx", {
            "PWRSPVAAAXX": {rule.RES_PARAM: [2.0, 2.0, 2.0], rule.MAX_INV_PARAM: [0.0, 0.0, 0.0]},
            "PWRWONAAAXX": {rule.FLOOR_PARAM: [OUT_PER_GW, OUT_PER_GW, OUT_PER_GW],
                            rule.MAX_INV_PARAM: [0.0, 0.0, 0.0]}})
        with mock.patch.object(rule, "FLOOR_TOP_UP", True):
            log = rule.edit_parametrization(path, ["Secondary Techs"], {"PWRSPVAAAXX", "PWRWONAAAXX"}, {})
        sheet = log["sheets"][0]
        self.assertEqual(sheet["allowed_techs"], ["PWRSPVAAAXX"])
        self.assertEqual(_caps(path, "PWRWONAAAXX"), [1.01, 0.0, 0.0])
        self.assertIn("floor_top_up", {c["reason"] for c in sheet["changes"]})

    def test_key_off_leaves_the_rule_unchanged(self) -> None:
        techs = {"PWRWONAAAXX": {rule.FLOOR_PARAM: [OUT_PER_GW] * 3, rule.MAX_INV_PARAM: [0.0, 0.0, 0.0]}}
        path = _workbook(self.dir / "p.xlsx", techs)
        with mock.patch.object(rule, "FLOOR_TOP_UP", False), \
                mock.patch.object(rule, "apply_floor_top_up", side_effect=AssertionError("called")):
            log = rule.edit_parametrization(path, ["Secondary Techs"], set(techs), {})
        self.assertNotIn("floor_top_up", log["sheets"][0])
        self.assertEqual(_caps(path, "PWRWONAAAXX"), [0.0, 0.0, 0.0])

    def test_yaml_key_is_a_boolean(self) -> None:
        good = self.dir / "good.yaml"
        good.write_text("floor_top_up: true\n", encoding="utf-8")
        self.assertEqual(rule.load_config(good), {"floor_top_up": True})
        bad = self.dir / "bad.yaml"
        bad.write_text("floor_top_up: 1\n", encoding="utf-8")
        with self.assertRaises(ValueError):
            rule.load_config(bad)


if __name__ == "__main__":
    unittest.main()
