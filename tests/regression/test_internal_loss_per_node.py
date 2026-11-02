"""The internal transmission loss as one number or per node, and the root demand rows.

``internal_transmission.transmission_loss`` is one number for every node or a
mapping of node code to number. A mapping names every node with internal
families; each node's internal families, bus-bar demand rows and interconnector
modes sent from its ``ELC<node>03`` take that node's 1 - loss. A setup whose
scenario workbook has no ``Demand_Projection`` sheet gives its demand rows in its
root ``A-O_Demand.xlsx``, which the rule reads and never scales.
"""
from __future__ import annotations

import shutil
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest import mock

import yaml
from openpyxl import Workbook, load_workbook

from ostram.pipeline.scenarios.rules import apply_internal_tx_losses as losses


REPO_ROOT = Path(__file__).resolve().parents[2]
LAC_ROOT_DEMAND = REPO_ROOT / "areas" / "lac" / "inputs" / "scenarios" / "root_workbooks" / "A-O_Demand.xlsx"
LAC_CONFIG = REPO_ROOT / "areas" / "lac" / "config" / "preparation" / "Config_country_codes.yaml"
NODES = ("AAAXX", "BBBXX")


def _demand_workbook(path: Path, measured_at: bool) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = losses.DEMAND_SHEET
    ws.append(["Demand/Share", "Fuel/Tech", "Name", "2023", "2024"] + (["Measured_at"] if measured_at else []))
    ws.append(["Demand", "ELCAAAXX03", "x", 100.0, 200.0] + (["bus-bar"] if measured_at else []))
    ws.append(["Demand", "ELCBBBXX03", "x", 10.0, 20.0] + (["bus-bar"] if measured_at else []))
    wb.save(path)
    return path


def _stage(folder: Path) -> Path:
    """Write two-node AR and demand workbooks into ``folder`` and return it."""
    folder.mkdir(parents=True)
    wb = Workbook()
    ws = wb.active
    ws.title = losses.SHEET
    ws.append(["Mode.Operation", "Tech", "Fuel", "Direction", 2023, 2024])
    for node in NODES:
        for family in losses.INTERNAL_FAMILIES:
            ws.append([1, f"{family}{node}", f"ELC{node}02", "Output", 1, 1])
            ws.append([1, f"{family}{node}", f"ELC{node}00", "Input", 1, 1])
    link = wb.create_sheet(losses.LINK_SHEET)
    link.append(["Mode.Operation", "Tech", "Fuel", "Direction", 2023, 2024])
    link.append([1, "TRNAAAXXBBBXX", "ELCAAAXX03", "Input", 1, 1])
    link.append([2, "TRNAAAXXBBBXX", "ELCBBBXX03", "Input", 1, 1])
    link.append([1, "TRNAAAXXBBBXX", "ELCBBBXX04", "Output", 0.95, 0.95])
    wb.save(folder / losses.PROJ_FILE)
    wb = Workbook()
    ws = wb.active
    ws.title = losses.SHEET
    ws.append(["Tech", "Fuel.O", "Value.Fuel.O"])
    for node in NODES:
        for family in losses.INTERNAL_FAMILIES:
            ws.append([f"{family}{node}", f"ELC{node}02", 1])
    link = wb.create_sheet(losses.LINK_SHEET)
    link.append(["Mode.Operation", "Fuel.I", "Value.Fuel.I", "Tech", "Fuel.O", "Value.Fuel.O"])
    link.append([1, "ELCAAAXX03", 1, "TRNAAAXXBBBXX", "ELCBBBXX04", 0.95])
    link.append([2, "ELCBBBXX03", 1, "TRNAAAXXBBBXX", "ELCAAAXX04", 0.95])
    wb.save(folder / losses.BASE_FILE)
    _demand_workbook(folder / losses.DEMAND_FILE, measured_at=False)
    return folder


def _values(folder: Path) -> dict:
    """Return every cell value of the three workbooks, keyed by file, sheet and cell."""
    out = {}
    for name in (losses.PROJ_FILE, losses.BASE_FILE, losses.DEMAND_FILE):
        wb = load_workbook(folder / name, data_only=True)
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    out[(name, ws.title, cell.coordinate)] = cell.value
        wb.close()
    return out


class LossFormTests(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name)
        self.template = self.dir / "template.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = losses.DEMAND_SHEET
        ws.append(["Demand/Share", "Fuel/Tech", "Name", losses.MEASURED_AT_COL, 2023, 2024])
        ws.append(["Demand", "ELCAAAXX03", "x", "bus-bar", 100.0, 200.0])
        ws.append(["Demand", "ELCBBBXX03", "x", "bus-bar", 10.0, 20.0])
        wb.save(self.template)

    def tearDown(self) -> None:
        self._tmp.cleanup()

    def _config(self, loss, name="Config_country_codes.yaml") -> Path:
        path = self.dir / name
        path.write_text(yaml.safe_dump({"internal_transmission": {"transmission_loss": loss}}), encoding="utf-8")
        return path

    def _run(self, loss, folder="stage5"):
        stage = _stage(self.dir / folder)
        log = losses.run(stage, skip_backup=True, config=self._config(loss, f"{folder}.yaml"), template=self.template)
        return stage, log

    def test_number_and_mapping_forms_are_read(self) -> None:
        self.assertEqual(losses.read_losses(self._config(0.03)), 0.03)
        self.assertEqual(losses.read_losses(self._config({"AAAXX": 0.03, "BBBXX": 0.05})),
                         {"AAAXX": 0.03, "BBBXX": 0.05})
        self.assertEqual(losses.read_loss(self._config(0.03)), 0.03)
        with self.assertRaises(ValueError) as ctx:
            losses.read_loss(self._config({"AAAXX": 0.03}))
        self.assertIn("one number", str(ctx.exception))

    def test_mapping_gives_each_node_its_own_ratio(self) -> None:
        stage, log = self._run({"AAAXX": 0.03, "BBBXX": 0.05})
        self.assertEqual(log["oar"], {"AAAXX": 0.97, "BBBXX": 0.95})
        values = _values(stage)
        proj = load_workbook(stage / losses.PROJ_FILE, data_only=True)
        for row in proj[losses.SHEET].iter_rows(min_row=2, values_only=True):
            expected = (0.97 if row[1].endswith("AAAXX") else 0.95) if row[3] == "Output" else 1
            self.assertEqual(row[4:], (expected, expected), row)
        links = {(r[0], r[3]): r[4:] for r in proj[losses.LINK_SHEET].iter_rows(min_row=2, values_only=True)}
        proj.close()
        self.assertEqual(links[(1, "Input")], (0.97, 0.97))
        self.assertEqual(links[(2, "Input")], (0.95, 0.95))
        self.assertEqual(links[(1, "Output")], (0.95, 0.95))
        self.assertEqual(values[(losses.DEMAND_FILE, losses.DEMAND_SHEET, "D2")], 100.0 * 0.97)
        self.assertEqual(values[(losses.DEMAND_FILE, losses.DEMAND_SHEET, "E3")], 20.0 * 0.95)
        self.assertEqual(values[(losses.BASE_FILE, losses.LINK_SHEET, "C3")], 0.95)

    def test_mapping_with_equal_values_matches_the_number(self) -> None:
        number, _ = self._run(0.03, "number")
        mapping, _ = self._run({"AAAXX": 0.03, "BBBXX": 0.03}, "mapping")
        self.assertEqual(_values(number), _values(mapping))

    def test_missing_unknown_and_out_of_range_entries_stop_the_run(self) -> None:
        cases = [({"AAAXX": 0.03}, "missing: BBBXX"),
                 ({"AAAXX": 0.03, "BBBXX": 0.03, "CCCXX": 0.03}, "unknown: CCCXX"),
                 ({"AAAXX": 0.03, "BBBXX": 1.0}, "of BBBXX must be in [0,1)"),
                 ({"AAAXX": -0.01, "BBBXX": 0.03}, "of AAAXX must be in [0,1)"),
                 (1.2, "must be in [0,1)")]
        for i, (loss, message) in enumerate(cases):
            with self.subTest(loss=loss), self.assertRaises(ValueError) as ctx:
                self._run(loss, f"case{i}")
            self.assertIn(message, str(ctx.exception))


class RootDemandRowsTests(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name)
        self.roots = self.dir / "root_workbooks"
        self.roots.mkdir()
        _demand_workbook(self.roots / losses.DEMAND_FILE, measured_at=True)
        self.scenario = self.dir / "OSTRAM_Scenario_Inputs.xlsx"
        wb = Workbook()
        wb.active.title = "Control"
        wb.save(self.scenario)
        self.config = self.dir / "Config_country_codes.yaml"
        self.config.write_text(yaml.safe_dump({"internal_transmission": {"transmission_loss": 0.03}}),
                               encoding="utf-8")
        self.paths = SimpleNamespace(root_workbooks=self.roots, scenario_workbook=self.scenario)

    def tearDown(self) -> None:
        self._tmp.cleanup()

    def test_scenario_workbook_without_sheet_falls_back_to_root_demand(self) -> None:
        with mock.patch.object(losses, "resolve_paths", return_value=self.paths):
            self.assertEqual(losses.demand_workbook(self.scenario), self.roots / losses.DEMAND_FILE)
            self.assertEqual(losses.demand_workbook(), self.roots / losses.DEMAND_FILE)

    def test_demand_scaled_once_per_build(self) -> None:
        stage = _stage(self.dir / "stage5")
        shutil.copy(self.roots / losses.DEMAND_FILE, stage / losses.DEMAND_FILE)
        with mock.patch.object(losses, "resolve_paths", return_value=self.paths):
            first = losses.run(stage, skip_backup=True, config=self.config, template=self.scenario)
            second = losses.run(stage, skip_backup=True, config=self.config, template=self.scenario)
        self.assertEqual((first["demand"]["cells"], second["demand"]["cells"]), (4, 0))
        values = _values(stage)
        self.assertEqual(values[(losses.DEMAND_FILE, losses.DEMAND_SHEET, "D2")], 100.0 * 0.97)
        self.assertEqual(values[(losses.DEMAND_FILE, losses.DEMAND_SHEET, "E3")], 20.0 * 0.97)
        root = load_workbook(self.roots / losses.DEMAND_FILE, data_only=True)
        self.assertEqual(root.active["D2"].value, 100.0)
        root.close()

    def test_lac_marks_its_19_demand_rows_bus_bar_with_one_loss(self) -> None:
        rows = losses.read_demand_rows(LAC_ROOT_DEMAND)
        self.assertEqual(len(rows), 19)
        self.assertEqual({row["measured_at"] for row in rows.values()}, {losses.BUS_BAR})
        self.assertEqual(losses.read_losses(LAC_CONFIG), 0.03)


if __name__ == "__main__":
    unittest.main()
