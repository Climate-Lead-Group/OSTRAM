"""Bangladesh committed plants enter at their commissioning years.

Patuakhali (1,320 MW coal) and Rooppur unit 1 (1,200 MW nuclear) enter in
2026, Rooppur unit 2 (1,200 MW) and Rupsha (800 MW gas) in 2027. The two
placeholders without a named plant (5,000 MW of gas and 3,000 MW of solar)
are gone. The base-year pin never caps a build below the minimum investment
that the workbook gives the same technology and year.
"""
from __future__ import annotations

import csv
import unittest
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook


REPO_ROOT = Path(__file__).resolve().parents[2]
WORKBOOK = REPO_ROOT / "inputs" / "scenarios" / "OSTRAM_Scenario_Inputs.xlsx"
PIN_CSV = REPO_ROOT / "config" / "scenarios" / "rules" / "pwr_min_2023_2026_pin.csv"
MIN_INV = "TotalAnnualMinCapacityInvestment"
MAX_INV = "TotalAnnualMaxCapacityInvestment"

EXPECTED_GW = {
    "PWRCOABGDXX": {2026: 1.32},
    "PWRURNBGDXX": {2026: 1.2, 2027: 1.2},
    "PWRNGSBGDXX": {2027: 0.8},
    "PWRSPVBGDXX": {},
}


def _bau_min_investment() -> dict[str, dict[int, float]]:
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    try:
        rows = workbook["Secondary_Techs"].iter_rows(values_only=True)
        header = list(next(rows))
        col = {name: i for i, name in enumerate(header) if isinstance(name, str)}
        years = {h: i for i, h in enumerate(header) if isinstance(h, int)}
        out: dict[str, dict[int, float]] = {}
        for row in rows:
            if row[col["scenario"]] == "BAU" and row[col["Parameter"]] == MIN_INV:
                out[row[col["Tech"]]] = {y: row[i] for y, i in years.items() if row[i] not in (None, "")}
        return out
    finally:
        workbook.close()


class BangladeshDatedPlantsTests(unittest.TestCase):
    def test_plants_enter_at_their_commissioning_years(self) -> None:
        minimum = _bau_min_investment()
        for tech, expected in EXPECTED_GW.items():
            built = {y: v for y, v in minimum[tech].items() if v}
            self.assertEqual(built, expected, tech)

    def test_pin_never_caps_a_build_below_the_workbook_minimum(self) -> None:
        minimum = _bau_min_investment()
        with PIN_CSV.open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.DictReader(stream))
        pinned_min = {(r["technology"], int(r["year"])) for r in rows if r["parameter"] == MIN_INV}
        below = [
            (r["technology"], r["year"], r["required_root_value"])
            for r in rows
            if r["parameter"] == MAX_INV and r["required_root_present"] == "true"
            and (r["technology"], int(r["year"])) not in pinned_min
            and Decimal(r["required_root_value"])
            < Decimal(str(minimum.get(r["technology"], {}).get(int(r["year"]), 0)))
        ]
        self.assertEqual(below, [])


if __name__ == "__main__":
    unittest.main()
