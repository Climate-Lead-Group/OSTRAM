"""The six interconnector cost bounding scenarios differ from their parents only in costs.

Each of A_Calibrated_BAU, B_Optimised_VRE and C_Target_VRE has a TxCostLow and a
TxCostHigh variant. A variant takes no setting that its parent does not have and
sets only CapitalCost and FixedCost of the 18 interconnectors, flat over the
years. Fixed cost is ROUND(0.015 x capital, 4). The low capital is below the
central value of Interconnector_Params and the high capital above it, and the
three low (and the three high) variants carry the same values.
"""
from __future__ import annotations

import csv
import unittest
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import yaml
from openpyxl import load_workbook


REPO_ROOT = Path(__file__).resolve().parents[2]
SCENARIOS = REPO_ROOT / "config" / "scenarios"
WORKBOOK = REPO_ROOT / "inputs" / "scenarios" / "OSTRAM_Scenario_Inputs.xlsx"
ROOTS = ("A_Calibrated_BAU", "B_Optimised_VRE", "C_Target_VRE")
BOUNDS = ("TxCostLow", "TxCostHigh")


def _interconnectors() -> set[str]:
    with (SCENARIOS / "technology_types.csv").open(encoding="utf-8-sig", newline="") as stream:
        return {row[1] for row in csv.reader(stream) if len(row) >= 2 and row[0] == "INTERCONNECTORS"}


def _central() -> dict[str, float]:
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    try:
        rows = list(workbook["Interconnector_Params"].iter_rows(values_only=True))
    finally:
        workbook.close()
    header = {name: i for i, name in enumerate(rows[0]) if isinstance(name, str)}
    first_year = next(i for i, name in enumerate(rows[0]) if name == 2023)
    return {row[header["Tech"]]: row[first_year] for row in rows[1:]
            if row[header["Parameter"]] == "CapitalCost"}


def _costs(name: str) -> dict[tuple[str, str], float]:
    block = yaml.safe_load((SCENARIOS / name / "scenario.yaml").read_text(encoding="utf-8"))
    return {(edit["tech"], edit["param"]): edit["value"] for edit in block["edits"]}


class TxCostBoundScenarioTests(unittest.TestCase):
    def test_each_variant_changes_only_interconnector_costs(self) -> None:
        links = _interconnectors()
        self.assertEqual(len(links), 18)
        for root in ROOTS:
            for bound in BOUNDS:
                name = f"{root}_{bound}"
                block = yaml.safe_load((SCENARIOS / name / "scenario.yaml").read_text(encoding="utf-8"))
                self.assertEqual(block["scenario"], name)
                self.assertEqual(block["base_scenario"], root)
                self.assertIs(block["apply_vre_ceiling_layer"], False)
                self.assertEqual(set(block), {"scenario", "base_scenario", "description",
                                              "apply_vre_ceiling_layer", "edits"})
                pairs = [(e["sheet"], e["tech"], e["param"], e["op"]) for e in block["edits"]]
                self.assertEqual(len(pairs), 36)
                self.assertEqual({p[1] for p in pairs}, links)
                self.assertEqual({(p[0], p[2], p[3]) for p in pairs},
                                 {("Secondary Techs", "CapitalCost", "set_flat"),
                                  ("Secondary Techs", "FixedCost", "set_flat")})

    def test_values_bound_the_central_costs(self) -> None:
        central = _central()
        for bound in BOUNDS:
            costs = [_costs(f"{root}_{bound}") for root in ROOTS]
            self.assertEqual(costs[0], costs[1])
            self.assertEqual(costs[0], costs[2])
            for (tech, param), value in costs[0].items():
                if param == "FixedCost":
                    capital = Decimal(f"{costs[0][(tech, 'CapitalCost')]:.2f}")
                    expected = (Decimal("0.015") * capital).quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP)
                    self.assertEqual(Decimal(str(value)), expected, tech)
                elif bound == "TxCostLow":
                    self.assertLess(value, central[tech], tech)
                else:
                    self.assertGreater(value, central[tech], tech)


if __name__ == "__main__":
    unittest.main()
