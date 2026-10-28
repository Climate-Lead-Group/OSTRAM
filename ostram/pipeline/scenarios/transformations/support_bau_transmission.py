"""Final support BAU ceiling: existing capacity plus surviving commitments.

Required annual investment minima are committed additions. Their
year is the commissioning year; OSeMOSYS keeps vintage v alive in year y when
0 <= y - v < OperationalLife. Never export this policy as inherited Restrictions.

``relax_interconnectors`` applies the same ceiling to the cross-border links of
A_Calibrated_BAU, through ``techs``.
"""
from __future__ import annotations

import json
import math
from pathlib import Path

from openpyxl import load_workbook

from .cap_trn_to_residual import TRN_TECHS


def _number(value: object, label: str, *, blank_zero: bool = False) -> float:
    if value is None and blank_zero:
        return 0.0
    if isinstance(value, bool) or not isinstance(value, (float, int)):
        raise ValueError(f"{label}: expected a numeric value, got {value!r}")
    result = float(value)
    if not math.isfinite(result) or result < 0:
        raise ValueError(f"{label}: expected a finite nonnegative value")
    return result


def apply_commitment_ceiling(path: Path, techs=TRN_TECHS, scenario: str = "BAU",
                             record_name: str | None = "support_bau_transmission.json") -> dict:
    """Change only allowlisted TRN total-capacity ceilings, fail before saving.

    ``techs`` is a subset of the TRN allowlist; ``record_name=None`` writes no
    record file and only returns the record.
    """
    techs = frozenset(techs)
    if not techs or not techs <= TRN_TECHS:
        raise ValueError(f"commitment ceiling: techs outside the TRN allowlist: {sorted(techs - TRN_TECHS)}")
    workbook = load_workbook(path)
    try:
        sheet = workbook["Secondary Techs"]
        headers = {cell.value: cell.column for cell in sheet[1]}
        years = {int(value): column for value, column in headers.items()
                 if str(value).isdigit() and 1900 <= int(value) <= 2200}
        if not years:
            raise ValueError("support BAU: missing model years")
        parameters = {"ResidualCapacity", "TotalAnnualMinCapacityInvestment",
                      "TotalAnnualMaxCapacity"}
        rows = {}
        for row in sheet.iter_rows(min_row=2):
            tech = row[headers["Tech"] - 1].value
            parameter = row[headers["Parameter"] - 1].value
            if tech in techs and parameter in parameters:
                key = (tech, parameter)
                if key in rows:
                    raise ValueError(f"support BAU: duplicate parameter {key}")
                rows[key] = row[0].row
        fixed = workbook["Fixed Horizon Parameters"]
        columns = {cell.value: cell.column - 1 for cell in fixed[1]}
        lives = {}
        for row in fixed.iter_rows(min_row=2, values_only=True):
            tech = row[columns["Tech"]]
            if tech in techs and row[columns["Parameter"]] == "OperationalLife":
                if tech in lives:
                    raise ValueError(f"support BAU: duplicate OperationalLife for {tech}")
                lives[tech] = _number(row[columns["Value"]], f"{tech} OperationalLife")
                if lives[tech] <= 0:
                    raise ValueError(f"{tech}: OperationalLife must be positive")
        changes, evidence, projection_modes = [], [], []
        for tech in sorted(techs):
            if tech not in lives or any((tech, p) not in rows for p in parameters):
                raise ValueError(f"support BAU: incomplete transmission parameters for {tech}")
            mode = sheet.cell(rows[tech, "TotalAnnualMaxCapacity"], headers["Projection.Mode"])
            projection_modes.append(dict(technology=tech, previous=mode.value, final="User defined"))
            # B1 ignores populated values under EMPTY. Activate this final
            # support-only ceiling explicitly, including zero-capacity years.
            changes.append((mode, "User defined"))
            commitments = {
                year: _number(sheet.cell(rows[tech, "TotalAnnualMinCapacityInvestment"], col).value,
                              f"{tech} commitment {year}", blank_zero=True)
                for year, col in years.items()
            }
            for year, col in sorted(years.items()):
                residual = _number(sheet.cell(rows[tech, "ResidualCapacity"], col).value,
                                   f"{tech} residual {year}")
                surviving = math.fsum(value for vintage, value in commitments.items()
                                      if 0 <= year - vintage < lives[tech])
                ceiling = residual + surviving
                cell = sheet.cell(rows[tech, "TotalAnnualMaxCapacity"], col)
                evidence.append(dict(technology=tech, year=year, residual=residual,
                                     committed_addition=commitments[year],
                                     operational_life=lives[tech], surviving_commitments=surviving,
                                     previous_ceiling=cell.value, ceiling=ceiling))
                changes.append((cell, ceiling))
        for cell, value in changes:
            cell.value = value
        workbook.save(path)
    finally:
        workbook.close()
    record = dict(schema="ostram-support-bau-transmission-v1", scenario=scenario,
                  commitment_source="TotalAnnualMinCapacityInvestment (unchanged)",
                  survival_rule="0 <= year - commissioning_year < OperationalLife",
                  projection_modes=projection_modes,
                  cells=evidence)
    if record_name is not None:
        path.with_name(record_name).write_text(json.dumps(record, indent=2), encoding="utf-8")
    label = "Support BAU" if scenario == "BAU" else scenario
    print(f"{label}: {len(techs)} transmission ceilings include surviving commitments")
    return record
