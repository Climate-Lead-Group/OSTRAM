"""
relax_to_reference.py
=====================

Rule: no limit of a scenario is tighter than the same limit of a reference
scenario built in the same run. ``B_Optimised_VRE`` takes it with
``A_Calibrated_BAU`` as the reference, so that B can always reproduce A's
plan and its least-cost solution never costs more than A's.

RULE
----
For every cell of A-O_Parametrization.xlsx that holds a bound:

    upper bound  -> the target takes the larger of its value and the
                    reference's; -1 or a row the compiler does not export
                    counts as no bound
    lower bound  -> the target takes the smaller of its value and the
                    reference's; a row the compiler does not export counts
                    as 0

Every other cell is left as it is. A row is exported when its
Projection.Mode is set and is not EMPTY, Zero or According to demand (the
B1 compiler's rule). When the reference row is not exported and the target
row breaks it, the target row is not exported either. A blank year cell in
an exported row that differs from the other scenario cannot be compared, so
the rule stops with an error.

The reference folder is read only. The target workbook is edited in place.

USAGE
-----
    python -m ostram.pipeline.scenarios.rules.relax_to_reference \\
        --input-dir A1_Outputs/A1_Outputs_B_Optimised_VRE \\
        --reference-dir A1_Outputs/A1_Outputs_A_Calibrated_BAU
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

PARAM_FILE_NAME = "A-O_Parametrization.xlsx"

# The bounds of model/osemosys_fast_preprocessed.txt. An upper bound limits a
# variable from above, a lower bound or requirement from below.
UPPER_BOUNDS = (
    "TotalAnnualMaxCapacity",
    "TotalAnnualMaxCapacityInvestment",
    "TotalTechnologyAnnualActivityUpperLimit",
    "TotalTechnologyModelPeriodActivityUpperLimit",
    "AnnualEmissionLimit",
    "ModelPeriodEmissionLimit",
)
LOWER_BOUNDS = (
    "TotalAnnualMinCapacity",
    "TotalAnnualMinCapacityInvestment",
    "TotalTechnologyAnnualActivityLowerLimit",
    "TotalTechnologyModelPeriodActivityLowerLimit",
    "REMinProductionTarget",
    "ReserveMargin",
)
NOT_EXPORTED_MODES = ("EMPTY", "Zero", "", None, "According to demand")
EXPORTED_MODE = "User defined"
NO_BOUND = -1


class RelaxationError(ValueError):
    """A bound cell that the rule cannot compare."""


def _year(value) -> int | None:
    try:
        year = int(value)
    except (TypeError, ValueError):
        return None
    return year if 1900 <= year <= 2200 else None


def _layout(header: list) -> dict | None:
    """Return the key and year columns of a sheet, or None when it holds no parameters."""
    if "Parameter" not in header:
        return None
    tech = next((header.index(n) for n in ("Tech", "Technology") if n in header), None)
    mode = header.index("Projection.Mode") if "Projection.Mode" in header else None
    years = {y: i for i, h in enumerate(header) if (y := _year(h)) is not None}
    return {"parameter": header.index("Parameter"), "tech": tech, "mode": mode, "years": years}


def _bound_rows(ws, layout: dict) -> dict[tuple, tuple[int, tuple]]:
    """Return {(tech, parameter): (row number, row values)} of the bound rows of a sheet."""
    rows: dict[tuple, tuple[int, tuple]] = {}
    for r, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        parameter = values[layout["parameter"]] if layout["parameter"] < len(values) else None
        if parameter not in UPPER_BOUNDS and parameter not in LOWER_BOUNDS:
            continue
        tech = values[layout["tech"]] if layout["tech"] is not None else None
        key = (tech, parameter)
        if key in rows:
            raise RelaxationError(f"{ws.title}: two rows for {key}")
        rows[key] = (r, values)
    return rows


def _exported(sheet: str, layout: dict, r: int, values: tuple) -> bool:
    if layout["mode"] is None:
        return True
    mode = values[layout["mode"]] if layout["mode"] < len(values) else None
    if mode in NOT_EXPORTED_MODES:
        return False
    if mode != EXPORTED_MODE:
        raise RelaxationError(
            f"{sheet} row {r}: Projection.Mode {mode!r} is not supported for a bound"
        )
    return True


def _value(values: tuple, column: int):
    return values[column] if column < len(values) else None


def _looser(parameter: str, own, reference):
    """Return the looser of two bound values of the same cell."""
    if parameter in UPPER_BOUNDS:
        if own == NO_BOUND or reference == NO_BOUND:
            return NO_BOUND
        return max(own, reference)
    return min(own, reference)


def relax(target, reference) -> list[dict]:
    """Loosen the bounds of the target workbook to the reference workbook, in place.

    ``target`` is an openpyxl workbook opened for writing; ``reference`` one
    opened for reading. Returns one record per changed cell or row.
    """
    changes: list[dict] = []
    for name in target.sheetnames:
        if name not in reference.sheetnames:
            continue
        tws, rws = target[name], reference[name]
        layout = _layout(list(next(tws.iter_rows(min_row=1, max_row=1, values_only=True))))
        if layout is None:
            continue
        rlayout = _layout(list(next(rws.iter_rows(min_row=1, max_row=1, values_only=True))))
        if rlayout is None or rlayout["years"].keys() != layout["years"].keys():
            raise RelaxationError(f"{name}: the year columns differ from the reference")
        trows, rrows = _bound_rows(tws, layout), _bound_rows(rws, rlayout)
        for (tech, parameter), (r, values) in trows.items():
            if not _exported(name, layout, r, values):
                continue
            upper = parameter in UPPER_BOUNDS
            ref_row = rrows.get((tech, parameter))
            if ref_row is None or not _exported(name, rlayout, *ref_row):
                # The reference has no bound here: no upper limit, a lower limit of 0.
                cells = [_value(values, c) for c in layout["years"].values()]
                breaks = (any(v is not None and v != NO_BOUND for v in cells) if upper
                          else any(v is not None and v > 0 for v in cells))
                if breaks:
                    if layout["mode"] is None:
                        raise RelaxationError(f"{name}: {parameter} has no reference value")
                    tws.cell(row=r, column=layout["mode"] + 1).value = "EMPTY"
                    changes.append({"sheet": name, "tech": tech, "parameter": parameter,
                                    "year": None, "old": EXPORTED_MODE, "new": "EMPTY"})
                continue
            for year, c in layout["years"].items():
                own = _value(values, c)
                ref = _value(ref_row[1], rlayout["years"][year])
                if own == ref:
                    continue
                if own is None or ref is None:
                    raise RelaxationError(
                        f"{name}: {tech} {parameter} {year} is blank in one scenario "
                        f"(own {own!r}, reference {ref!r})"
                    )
                new = _looser(parameter, own, ref)
                if new != own:
                    tws.cell(row=r, column=c + 1).value = new
                    changes.append({"sheet": name, "tech": tech, "parameter": parameter,
                                    "year": year, "old": own, "new": new})
    return changes


def run(input_dir: Path, reference_dir: Path) -> dict:
    """Apply the rule to ``input_dir`` with ``reference_dir`` as the reference."""
    target_path = Path(input_dir) / PARAM_FILE_NAME
    reference_path = Path(reference_dir) / PARAM_FILE_NAME
    for path in (target_path, reference_path):
        if not path.is_file():
            raise FileNotFoundError(f"Parametrization file not found: {path}")
    reference = load_workbook(reference_path, read_only=True, data_only=True)
    target = load_workbook(target_path)
    try:
        changes = relax(target, reference)
        if changes:
            target.save(target_path)
    finally:
        reference.close()
        target.close()
    return {"file": str(target_path), "reference": str(reference_path), "changes": changes}


def print_summary(log: dict) -> None:
    """Print the number of changed cells by parameter and direction."""
    counts: dict[str, int] = {}
    for change in log["changes"]:
        counts[change["parameter"]] = counts.get(change["parameter"], 0) + 1
    print(f"relax_to_reference: {len(log['changes'])} change(s) in {log['file']}")
    for parameter, n in sorted(counts.items()):
        side = "upper" if parameter in UPPER_BOUNDS else "lower"
        print(f"  {parameter} ({side} bound): {n}")


def main() -> int:
    """Run the command-line interface and return its exit code."""
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--reference-dir", type=Path, required=True)
    parser.add_argument("--log", type=Path, default=None,
                        help="Write the per-cell change list to this JSON file.")
    args = parser.parse_args()
    try:
        log = run(args.input_dir, args.reference_dir)
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    if args.log is not None:
        args.log.write_text(json.dumps(log, indent=2, default=str), encoding="utf-8")
    print_summary(log)
    return 0


if __name__ == "__main__":
    sys.exit(main())
