"""Values the site shows, read from the registered setups and their authority files.

Every count and parameter on the site comes from here. Each value carries the
repository file it was read from, so a page can cite it and the tests can read
the same files again and compare.
"""

from __future__ import annotations

import ast
import csv
from dataclasses import dataclass, field
from pathlib import Path, PurePosixPath
import re
from typing import Any, Mapping

from ostram import yaml_io
from ostram.paths import ProjectPaths
from ostram.setups import (
    REGISTRY_SCHEMA,
    SetupManifest,
    load_manifest,
    resolve_authority_reference,
)

# Internal grid families in the interconnector taxonomy, in teaching order.
INTERNAL_CATEGORIES = (
    "TRANSMISSION_PLANNED_RE",
    "TRANSMISSION_UNPLANNED_RE",
    "TRANSMISSION_REPOT_RE",
    "TRANSMISSION_PLANNED_NONRENO",
    "TRANSMISSION_UNPLANNED_NR",
    "TRANSMISSION_REPOT_NR",
)
DISPATCH_CATEGORY = "TRANSMISSION_SUM"
INTERCONNECTOR_CATEGORY = "INTERCONNECTORS"
_NODE_COMMENT = re.compile(r'^\s*-\s*"?([A-Z]{3}[A-Z]{0,2})"?\s*#\s*(.+?)\s*$')


@dataclass(frozen=True)
class Value:
    """One value shown on the site and the repository file it is read from."""

    key: str
    value: Any
    source: str
    note: str = ""


@dataclass
class SetupInfo:
    """What the site shows about one registered setup."""

    setup_id: str
    title: str
    manifest: str
    nodes: list[str] = field(default_factory=list)
    node_labels: dict[str, str] = field(default_factory=dict)
    interconnectors: list[str] = field(default_factory=list)
    families: dict[str, list[str]] = field(default_factory=dict)
    dispatch: list[str] = field(default_factory=list)
    technologies_by_node: dict[str, list[str]] = field(default_factory=dict)
    years: list[int] = field(default_factory=list)
    timeslices: list[str] = field(default_factory=list)
    seasons: int = 0
    day_types: int = 0
    brackets: int = 0
    support: list[str] = field(default_factory=list)
    roots: list[dict] = field(default_factory=list)
    decision: list[str] = field(default_factory=list)
    derived: list[dict] = field(default_factory=list)
    internal: dict[str, Any] = field(default_factory=dict)
    family_parameters: dict[str, dict] = field(default_factory=dict)
    link_import_cap: dict[str, Any] = field(default_factory=dict)
    historical_trade_roots: list[str] = field(default_factory=list)
    vre_ceilings: dict[str, float] = field(default_factory=dict)
    corridor_costs: list[dict] = field(default_factory=list)
    corridor_residual: dict[str, float] = field(default_factory=dict)
    corridor_references: dict[str, str] = field(default_factory=dict)
    corridor_units: str = ""
    sources: dict[str, str] = field(default_factory=dict)
    values: dict[str, Value] = field(default_factory=dict)

    def add(self, name: str, value: Any, source: str, note: str = "") -> None:
        key = f"{self.setup_id}.{name}"
        self.values[key] = Value(key, value, source, note)


def relative(path: Path, root: Path) -> str:
    """Return ``path`` relative to the project root, in POSIX form."""

    try:
        return PurePosixPath(Path(path).resolve().relative_to(root.resolve())).as_posix()
    except ValueError:
        return Path(path).name


def registered_setups(
    paths: ProjectPaths,
    registry_path: Path | None = None,
) -> list[SetupManifest]:
    """Return the manifests of every setup in the registry, in registry order."""

    registry = (
        Path(registry_path).resolve()
        if registry_path is not None
        else (paths.config_root / "setups" / "registry.json").resolve()
    )
    raw = yaml_io.load_file(registry)
    if raw.get("schema") != REGISTRY_SCHEMA:
        raise ValueError(f"unsupported setup registry schema in {registry}")
    manifests = []
    for setup_id, entry in raw["setups"].items():
        reference = entry.get("manifest") if isinstance(entry, Mapping) else entry
        manifest_path = resolve_authority_reference(
            reference, paths=paths, setup_root=registry.parent
        )
        manifests.append(load_manifest(manifest_path, expected_setup=setup_id))
    return manifests


def _taxonomy(path: Path) -> list[tuple[str, str]]:
    with open(path, encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.reader(handle))
    return [(row[0].strip(), row[1].strip()) for row in rows[1:] if len(row) >= 2 and row[0].strip()]


def _node_of(code: str, prefix_length: int = 6) -> str:
    return code[prefix_length:]


def _node_labels(country_config: Path, nodes: list[str], countries: Mapping) -> dict[str, str]:
    """Return a label per node, from the comment beside it in the country list or the country name."""

    comments: dict[str, str] = {}
    in_list = False
    for line in country_config.read_text(encoding="utf-8").splitlines():
        if re.match(r"^countries:\s*$", line):
            in_list = True
            continue
        if in_list and line and not line.startswith((" ", "-", "#")):
            in_list = False
        if in_list:
            match = _NODE_COMMENT.match(line)
            if match:
                comments[match.group(1)] = match.group(2)
    labels = {}
    for node in nodes:
        name = (countries.get(node[:3]) or {}).get("english_name", node[:3])
        labels[node] = comments.get(node) or comments.get(node[:3] if node.endswith("XX") else node) or (
            name if node.endswith("XX") else f"{name} {node[3:]}"
        )
    return labels


def node_losses(setting: Any) -> dict[str, Any]:
    """Return the internal grid loss by node, from one number or a mapping of node to number.

    One number, or a mapping whose nodes all hold the same number, gives one value under
    the empty node name; a mapping with different numbers gives one value per node.
    """

    if not isinstance(setting, Mapping):
        return {"": setting}
    if len(set(setting.values())) == 1:
        return {"": next(iter(setting.values()))}
    return {str(node): setting[node] for node in sorted(setting)}


def _workbook_tables(path: Path) -> tuple[list[dict], dict[str, float], dict[str, str]]:
    """Return the corridor cost rows, the first-year residual capacity and the references."""

    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    costs: list[dict] = []
    residual: dict[str, float] = {}
    references: dict[str, str] = {}
    try:
        if "Transmission_Cost_Basis" in workbook.sheetnames:
            rows = list(workbook["Transmission_Cost_Basis"].iter_rows(values_only=True))
            header = [str(cell) if cell is not None else "" for cell in rows[0]]
            for row in rows[1:]:
                record = dict(zip(header, row))
                tech = record.get("Technology")
                if isinstance(tech, str) and tech.startswith("TRN"):
                    costs.append(
                        {
                            "tech": tech,
                            "corridor": record.get("Corridor") or "",
                            "capital": record.get("Capital_USD 2023_per_kW"),
                            "fixed": record.get("Fixed_USD 2023_per_kW_year"),
                            "low": record.get("Capital_low_USD 2023_per_kW"),
                            "high": record.get("Capital_high_USD 2023_per_kW"),
                            "design": record.get("Module_and_policy") or "",
                        }
                    )
        if "Interconnector_Params" in workbook.sheetnames:
            rows = list(workbook["Interconnector_Params"].iter_rows(values_only=True))
            header = [str(cell) if cell is not None else "" for cell in rows[0]]
            first_year = next(i for i, name in enumerate(header) if re.fullmatch(r"\d{4}", name))
            reference = header.index("IEEE Reference") if "IEEE Reference" in header else None
            for row in rows[1:]:
                tech, parameter = row[2], row[5]
                if not isinstance(tech, str):
                    continue
                if parameter == "ResidualCapacity" and row[0] == "BAU":
                    residual[tech] = float(row[first_year])
                if parameter == "CapitalCost" and reference is not None and row[reference]:
                    references.setdefault(tech, str(row[reference]))
    finally:
        workbook.close()
    return costs, residual, references


def _cost_units(path: Path) -> str:
    """Return the currency and year of the corridor costs, as the header of the cost basis sheet names them."""

    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Transmission_Cost_Basis" not in workbook.sheetnames:
            return ""
        header = next(workbook["Transmission_Cost_Basis"].iter_rows(values_only=True))
    finally:
        workbook.close()
    for cell in header:
        match = re.fullmatch(r"Capital_(.+)_per_kW", str(cell or ""))
        if match:
            return match.group(1)
    return ""


def _registry(manifest: SetupManifest, paths: ProjectPaths) -> tuple[dict, Path, Path]:
    """Return the raw scenario registry, its file and the folder of scenario files."""

    reference = manifest.authorities["scenario_registry"].reference
    path = resolve_authority_reference(reference, paths=paths, setup_root=manifest.root)
    raw = yaml_io.load_file(path)
    consolidated = "scenario_registry" in raw
    registry = raw.get("scenario_registry", raw)
    if "scenario_config_root" in manifest.authorities:
        folder = resolve_authority_reference(
            manifest.authorities["scenario_config_root"].reference,
            paths=paths,
            setup_root=manifest.root,
        )
    else:
        folder = path.parent.parent / "scenarios" if consolidated else path.parent
    return registry, path, folder


def setup_info(manifest: SetupManifest, paths: ProjectPaths) -> SetupInfo:
    """Read every value the site shows about one setup."""

    from ostram.pipeline.scenarios.registry import load_scenario_file

    root = paths.project_root
    authorities = manifest.source_paths(paths, require_exists=False)
    info = SetupInfo(
        setup_id=manifest.setup_id,
        title=str(manifest.metadata.get("title") or manifest.setup_id),
        manifest=relative(manifest.path, root),
    )
    info.sources["manifest"] = info.manifest

    taxonomy_path = authorities.get("interconnector_taxonomy")
    if taxonomy_path is not None and taxonomy_path.is_file():
        source = relative(taxonomy_path, root)
        info.sources["taxonomy"] = source
        rows = _taxonomy(taxonomy_path)
        for category, tech in rows:
            if category == INTERCONNECTOR_CATEGORY:
                info.interconnectors.append(tech)
            elif category in INTERNAL_CATEGORIES:
                info.families.setdefault(category, []).append(tech)
            elif category == DISPATCH_CATEGORY:
                info.dispatch.append(tech)
        nodes = sorted({_node_of(tech) for tech in info.families.get("TRANSMISSION_PLANNED_NONRENO", [])})
        info.nodes = nodes
        for category, tech in rows:
            if category in (INTERCONNECTOR_CATEGORY,):
                continue
            node = _node_of(tech)
            if node in nodes:
                info.technologies_by_node.setdefault(node, []).append(tech)
        info.interconnectors.sort()
        info.add("nodes", len(nodes), source, "nodes with an internal grid technology (PWRTRN)")
        info.add("interconnectors", len(info.interconnectors), source, "rows in the INTERCONNECTORS category")
        info.add(
            "internal_grid_technologies",
            sum(len(v) for v in info.families.values()),
            source,
            "rows in the six internal transmission categories",
        )
        info.add("internal_families", len(info.families), source, "internal transmission categories present")

    country_path = authorities.get("country_config")
    if country_path is not None and country_path.is_file():
        source = relative(country_path, root)
        info.sources["country_config"] = source
        country = yaml_io.load_file(country_path)
        countries = country.get("country_data") or {}
        info.node_labels = _node_labels(country_path, info.nodes, countries)
        for entry in manifest.metadata.get("country_regions") or []:
            if isinstance(entry, Mapping) and entry.get("region") in info.node_labels and entry.get("label"):
                info.node_labels[entry["region"]] = str(entry["label"])
        info.add("countries", len(countries), source, "entries in country_data")
        internal = country.get("internal_transmission") or {}
        info.internal = dict(internal)
        for name, value in internal.items():
            if name != "transmission_loss" or not isinstance(value, Mapping):
                info.add(f"internal.{name}", value, source, f"internal_transmission.{name}")
                continue
            for node, loss in node_losses(value).items():
                if node:
                    info.add(f"internal.{name}.{node}", loss, source, f"internal_transmission.{name}.{node}")
                else:
                    info.add(f"internal.{name}", loss, source, f"internal_transmission.{name}, the same for every node")
        if {"esmap_capital_cost_usd2010_per_kw", "muv_2023", "muv_2010"} <= set(internal):
            base = internal["esmap_capital_cost_usd2010_per_kw"] * internal["muv_2023"] / internal["muv_2010"]
            info.add("internal.base_capital_cost", base, source, "computed: ESMAP cost x MUV 2023 / MUV 2010")
            if "fixed_cost_share" in internal:
                info.add("internal.base_fixed_cost", internal["fixed_cost_share"] * base, source,
                         "computed: fixed cost share x base capital cost")
            if {"gec2_cost_rs_crore", "gec2_capacity_mw", "inr_per_usd_2022", "muv_2022"} <= set(internal):
                add_on = (
                    internal["gec2_cost_rs_crore"] * 1e7 / (internal["gec2_capacity_mw"] * 1e3)
                    / internal["inr_per_usd_2022"] * internal["muv_2023"] / internal["muv_2022"]
                )
                info.add("internal.re_add_on", add_on, source, "computed: GEC-II cost per kW in 2023 USD")
                info.add("internal.re_capital_cost", base + add_on, source, "computed: base plus RE add-on")
                info.add("internal.re_multiplier", (base + add_on) / base, source,
                         "computed: (base + RE add-on) / base")
        if "transmission_loss" in internal:
            for node, loss in node_losses(internal["transmission_loss"]).items():
                info.add("internal.efficiency" + (f".{node}" if node else ""), 1 - loss, source,
                         "computed: 1 - transmission_loss")
        for family in ("RNWTRN", "RNWNLI", "RNWRPO", "PWRTRN", "TRNNLI", "TRNRPO", "DSPTRN"):
            if isinstance(country.get(family), Mapping):
                info.family_parameters[family] = dict(country[family])
        info.add("first_year_config", country.get("first_year"), source, "first_year")

    compilation_path = authorities.get("compilation_config")
    if compilation_path is not None and compilation_path.is_file():
        source = relative(compilation_path, root)
        info.sources["compilation_config"] = source
        compilation = yaml_io.load_file(compilation_path)
        first, last = int(compilation["initial_year"]), int(compilation["final_year"])
        info.years = list(range(first, last + 1))
        info.add("first_year", first, source, "initial_year")
        info.add("last_year", last, source, "final_year")
        info.add("years", len(info.years), source, "final_year - initial_year + 1")
        sets = _timeslice_sets(compilation)
        info.timeslices = list(sets.get("Timeslices") or [])
        info.seasons = len(sets.get("Season") or [])
        info.day_types = len(sets.get("DayType") or [])
        info.brackets = len(sets.get("DailyTimeBracket") or [])
        info.add("timeslices", len(info.timeslices), source, "Timeslices")
        info.add("seasons", info.seasons, source, "Season")
        info.add("day_types", info.day_types, source, "DayType")
        info.add("daily_time_brackets", info.brackets, source, "DailyTimeBracket")

    if "scenario_registry" in manifest.authorities:
        registry, registry_path, folder = _registry(manifest, paths)
        source = relative(registry_path, root)
        info.sources["scenario_registry"] = source
        info.support = [str(name) for name in registry.get("support_scenarios", [])]
        root_names = []
        for entry in registry.get("root_scenarios", []):
            root_names.append(entry["name"])
            info.roots.append(
                {
                    "name": entry["name"],
                    "role": entry.get("role", ""),
                    "dependencies": [
                        f"{d.get('type')}: {d.get('scenario')}" for d in entry.get("dependencies") or []
                    ],
                }
            )
        info.decision = [str(name) for name in registry.get("decision_scenarios", [])]
        for name in info.decision:
            if name in root_names:
                continue
            scenario_file = folder / name / "scenario.yaml"
            block = load_scenario_file(scenario_file) if scenario_file.is_file() else {}
            info.derived.append(
                {
                    "name": name,
                    "base": block.get("base_scenario", ""),
                    "description": " ".join(str(block.get("description") or block.get("rule") or "").split()),
                    "edits": len(block.get("edits") or []),
                    "reuse": block.get("reuse_edits_from") or "",
                    "overlay": bool(block.get("direction_overlay")),
                    "gross_import_cap": block.get("gross_import_cap") or {},
                    "source": relative(scenario_file, root),
                }
            )
        info.add("scenarios", len(info.support) + len(info.decision), source,
                 "support scenarios plus decision scenarios")
        info.add("support_scenarios", len(info.support), source, "support_scenarios")
        info.add("root_scenarios", len(info.roots), source, "root_scenarios")
        info.add("decision_scenarios", len(info.decision), source, "decision_scenarios")
        info.add("derived_scenarios", len(info.derived), source, "decision scenarios that are not roots")

    policies = manifest.policies
    info.historical_trade_roots = list(policies.get("historical_trade_roots") or [])
    if isinstance(policies.get("link_import_cap"), Mapping):
        info.link_import_cap = dict(policies["link_import_cap"])
        cap = info.link_import_cap
        info.add("link_import_cap.first_year", cap.get("first_year"), info.manifest, "policies.link_import_cap.first_year")
        info.add("link_import_cap.links", len(cap.get("links") or []), info.manifest, "policies.link_import_cap.links")
    raw_manifest = yaml_io.load_file(manifest.path)
    ceilings = (raw_manifest.get("vre_ceilings") or {}).get("ceilings_gw") or {}
    info.vre_ceilings = {str(k): float(v) for k, v in ceilings.items()}
    if ceilings:
        info.add("vre_ceilings", len(ceilings), info.manifest, "vre_ceilings.ceilings_gw")
    effective = manifest.metadata.get("effective_values") or {}
    for name, value in effective.items():
        info.add(f"effective.{name}", value, info.manifest, f"metadata.effective_values.{name}")

    workbook = authorities.get("scenario_workbook")
    if workbook is not None and workbook.is_file():
        source = relative(workbook, root)
        info.sources["scenario_workbook"] = source
        info.corridor_costs, info.corridor_residual, info.corridor_references = _workbook_tables(workbook)
        info.corridor_units = _cost_units(workbook)
        if info.corridor_costs:
            info.add("corridor_costs", len(info.corridor_costs), source, "Transmission_Cost_Basis rows")
        for row in info.corridor_costs:
            for name in ("capital", "fixed", "low", "high"):
                if row[name] is not None:
                    info.add(f"corridor.{row['tech']}.{name}", row[name], source,
                             f"Transmission_Cost_Basis, {row['tech']}")
        for tech, residual in info.corridor_residual.items():
            info.add(f"corridor.{tech}.residual", residual, source,
                     f"Interconnector_Params, {tech} ResidualCapacity, first year")
    for tech, ceiling in info.vre_ceilings.items():
        info.add(f"vre_ceiling.{tech}", ceiling, info.manifest, f"vre_ceilings.ceilings_gw.{tech}")
    for derived in info.derived:
        cap = derived["gross_import_cap"]
        for name in ("fraction", "first_year", "last_year"):
            if name in cap:
                info.add(f"gross_import_cap.{derived['name']}.{name}", cap[name], derived["source"],
                         f"gross_import_cap.{name}")
    return info


def _timeslice_sets(compilation: Mapping) -> Mapping:
    """Return the mapping of the compilation configuration that holds the ``Timeslices`` list."""

    stack = [compilation]
    while stack:
        item = stack.pop(0)
        if isinstance(item, Mapping):
            if "Timeslices" in item:
                return item
            stack.extend(item.values())
    return {}


def rule_modules(paths: ProjectPaths) -> list[tuple[str, str]]:
    """Return each A3 rule module and the first line of its docstring."""

    folder = paths.package_root / "pipeline" / "scenarios" / "rules"
    modules = []
    for path in sorted(folder.glob("*.py")):
        if path.name == "__init__.py":
            continue
        tree = ast.parse(path.read_text(encoding="utf-8"))
        doc = (ast.get_docstring(tree) or "").strip().splitlines()
        modules.append((path.name, doc[0] if doc else ""))
    return modules


def literal_assignment(path: Path, name: str) -> Any:
    """Return the literal value assigned to ``name`` in a module, read without importing it."""

    tree = ast.parse(path.read_text(encoding="utf-8"))
    for node in tree.body:
        if isinstance(node, ast.Assign) and any(
            isinstance(target, ast.Name) and target.id == name for target in node.targets
        ):
            value = node.value
            if isinstance(value, ast.Call) and value.args:
                value = value.args[0]
            return ast.literal_eval(value)
    raise KeyError(f"{name} not assigned in {path}")


def code_counts(paths: ProjectPaths) -> list[dict]:
    """Return the Python files and lines of each package folder of ``ostram``."""

    package = paths.package_root
    rows = []
    folders = sorted(
        {p.parent for p in package.rglob("*.py") if "__pycache__" not in p.parts},
        key=lambda p: relative(p, paths.project_root),
    )
    for folder in folders:
        files = sorted(folder.glob("*.py"))
        lines = sum(len(f.read_text(encoding="utf-8").splitlines()) for f in files)
        rows.append({"folder": relative(folder, paths.project_root), "files": len(files), "lines": lines})
    return rows


def repository_values(paths: ProjectPaths, stages: list[dict], setups: int, registry: str) -> dict[str, Value]:
    """Return the counts of the code map, read by walking the repository."""

    root = paths.project_root
    found: dict[str, Value] = {}

    def add(name: str, value: int, source: str, note: str) -> None:
        found[f"repo.{name}"] = Value(f"repo.{name}", value, source, note)

    counts = code_counts(paths)
    for row in counts:
        add(f"code.{row['folder']}.files", row["files"], row["folder"], "Python files in the folder")
        add(f"code.{row['folder']}.lines", row["lines"], row["folder"], "lines of those files")
    for stage in stages:
        folder = stage["folder"]
        if folder.endswith(".py"):
            files, lines = 1, len((root / folder).read_text(encoding="utf-8").splitlines())
        else:
            inside = [r for r in counts if r["folder"] == folder or r["folder"].startswith(folder + "/")]
            files, lines = sum(r["files"] for r in inside), sum(r["lines"] for r in inside)
        add(f"stage.{stage['stage']}.files", files, folder, "Python files of the stage")
        add(f"stage.{stage['stage']}.lines", lines, folder, "lines of those files")
    add("tests", len(list((root / "tests").rglob("test_*.py"))), "tests", "files named test_*.py")
    add("docs_pages", len(list((root / "docs").rglob("*.md"))), "docs", "Markdown pages")
    add("rule_modules", len(rule_modules(paths)), "ostram/pipeline/scenarios/rules", "modules other than __init__.py")
    add("workflows", len(list((root / ".github" / "workflows").glob("*.yml"))), ".github/workflows", "workflow files")
    add("setups", setups, registry, "registered setups")
    return found
