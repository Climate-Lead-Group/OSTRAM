"""The source entry of each value the site shows, read from the repository's own records.

``content/sources.yaml`` names where the record of each value lives: a docs table
row, the comment lines above a YAML key, a workbook cell, a block of the base-year
value record (its sources resolved to its list of documents, its assumption labels
to their sentences), or the file a setting is read from. This module reads those records when the site is built and turns each
into one entry of four kinds:

- ``external``: citations with their URLs or DOIs, as the record gives them, and any
  ``ASSUMPTION:`` sentences of the same record;
- ``assumption``: the record's ``ASSUMPTION:`` sentences, without the label;
- ``computed``: the formula and the keys of its inputs, each with its own entry;
- ``setting``: a permalink to the file the value is read from, with a label.

No source text is written here or in the source map. A value with no record gets
no entry and is left out of the site.
"""

from __future__ import annotations

from fnmatch import fnmatchcase
from pathlib import Path
import re
import subprocess
from typing import Any

from ostram import yaml_io

MAP_FILE = Path(__file__).resolve().parent / "content" / "sources.yaml"
_URL = re.compile(r"https?://[^\s<>\"]+")
_DOI = re.compile(r"\bdoi:\s*(10\.\d{4,9}/\S+)", re.IGNORECASE)
_MARKER = re.compile(r"(?:^|(?<=\s))(ASSUMPTION|Source):\s")
_AUTHOR_YEAR = re.compile(r"([A-Z][A-Za-z&/.-]*(?:\s[A-Z][A-Za-z&/.-]*)*)\s\((\d{4})\)")
_SEE = re.compile(r"^ASSUMPTION;\s*see the (.+?)(?:\s+below)?\s*(?:\(|$)")


def git_commit(root: Path) -> str | None:
    """Return ``git rev-parse HEAD`` of the repository at ``root``, or None without Git."""

    if not (Path(root) / ".git").exists():
        return None
    try:
        result = subprocess.run(
            ["git", "-C", str(root), "rev-parse", "HEAD"],
            capture_output=True, text=True, check=True, timeout=30,
        )
    except (OSError, subprocess.SubprocessError):
        return None
    commit = result.stdout.strip()
    return commit if re.fullmatch(r"[0-9a-f]{40}", commit) else None


# Text of a record -------------------------------------------------------------------


def plain(text: str) -> str:
    """Return ``text`` without Markdown emphasis, code marks and autolink brackets."""

    text = re.sub(r"<(https?://[^>\s]+)>", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\((https?://[^)\s]+)\)", r"\1 (\2)", text)
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"(?<![\w*])\*(?!\s)(.+?)(?<!\s)\*(?![\w*])", r"\1", text)
    return text.replace("`", "")


def one_line(text: str) -> str:
    return " ".join(text.split())


def urls_of(text: str) -> list[str]:
    """Return the URLs and DOIs in ``text``, in order, without trailing punctuation."""

    found: list[str] = []
    for match in _URL.finditer(text):
        url = match.group(0)
        while url and (url[-1] in ".,;:]" or (url[-1] == ")" and url.count("(") < url.count(")"))):
            url = url[:-1]
        if url not in found:
            found.append(url)
    for match in _DOI.finditer(text):
        url = "https://doi.org/" + match.group(1).rstrip(".,;:)")
        if url not in found:
            found.append(url)
    return found


def _pieces(text: str) -> list[str]:
    """Split a record into paragraphs, and a numbered reference list into its items."""

    pieces: list[str] = []
    for paragraph in re.split(r"\n\s*\n", plain(text)):
        paragraph = one_line(paragraph)
        if not paragraph:
            continue
        if re.match(r"\[1\]\s", paragraph):
            pieces += [part.strip() for part in re.split(r"\s(?=\[\d+\]\s)", paragraph) if part.strip()]
        else:
            pieces.append(paragraph)
    return pieces


def parse_record(text: str) -> tuple[list[dict], list[str]]:
    """Return the citations and the ASSUMPTION sentences of one record.

    Text after ``Source:`` is a citation; text after ``ASSUMPTION:`` is an assumption
    sentence, up to the next marker. Text before the first marker is a citation,
    unless the record names its citation with ``Source:``; then it is a description.
    """

    pieces = _pieces(text)
    labelled = any(match.group(1) == "Source" for piece in pieces for match in _MARKER.finditer(piece))
    citations: list[dict] = []
    assumptions: list[str] = []
    for piece in pieces:
        marks = list(_MARKER.finditer(piece))
        lead = piece[: marks[0].start()] if marks else piece
        lead = re.sub(r"^\[\d+\]\s*", "", lead).strip()
        if lead and not labelled:
            citations.append({"text": lead, "urls": urls_of(lead)})
        for index, mark in enumerate(marks):
            end = marks[index + 1].start() if index + 1 < len(marks) else len(piece)
            body = piece[mark.end():end].strip()
            if not body:
                continue
            if mark.group(1) == "ASSUMPTION":
                assumptions.append(body)
            else:
                citations.append({"text": body, "urls": urls_of(body)})
    return citations, assumptions


def _tokens(text: str) -> list[str]:
    return [t for t in re.findall(r"[A-Za-z0-9]+", text) if len(t) > 1]


def _short_cites(text: str) -> list[list[str]]:
    found = _AUTHOR_YEAR.findall(text)
    if found:
        return [_tokens(f"{author} {year}") for author, year in found]
    return [_tokens(re.split(r"[(:]", part)[0]) for part in text.split(";") if _tokens(re.split(r"[(:]", part)[0])]


def resolve_short(text: str, pool: list[dict]) -> list[str] | None:
    """Return the URLs of the citations of ``pool`` that a short citation such as ``ESMAP (2026)`` names.

    Each short citation must match exactly one citation with a URL, holding all its words.
    """

    resolved: list[str] = []
    for tokens in _short_cites(text):
        if not tokens:
            return None
        matches = [
            c for c in pool
            if c["urls"] and all(re.search(rf"(?<![A-Za-z0-9]){re.escape(t)}(?![A-Za-z0-9])", c["text"]) for t in tokens)
        ]
        if len(matches) != 1:
            return None
        resolved += [url for url in matches[0]["urls"] if url not in resolved]
    return resolved or None


def entry_from_record(text: str, pool: list[dict] | None = None) -> dict | None:
    """Return the entry of one record: external, assumption, or None when it is empty."""

    citations, assumptions = parse_record(text or "")
    if any(c["urls"] for c in citations):
        return {"kind": "external", "citations": citations, "assumptions": assumptions}
    if citations and pool and not assumptions:
        text = " ".join(c["text"] for c in citations)
        urls = resolve_short(text, pool)
        if urls:
            return {"kind": "external", "citations": [{"text": text, "urls": urls}], "assumptions": []}
    if assumptions:
        return {"kind": "assumption", "assumptions": assumptions}
    return None


# Where records live ------------------------------------------------------------------


def markdown_table(path: Path, header: list[str]) -> list[dict[str, str]]:
    """Return the rows of the Markdown table in ``path`` whose header cells are ``header``."""

    rows: list[dict[str, str]] = []
    lines = path.read_text(encoding="utf-8").splitlines()
    for index, line in enumerate(lines):
        if _cells(line) != header:
            continue
        for row in lines[index + 2:]:
            if not row.strip().startswith("|"):
                break
            rows.append(dict(zip(header, _cells(row))))
        break
    return rows


def _cells(line: str) -> list[str]:
    line = line.strip()
    if not (line.startswith("|") and line.endswith("|")):
        return []
    return [cell.strip() for cell in line[1:-1].split("|")]


def yaml_comments(path: Path, block: str, key: str) -> str:
    """Return the comment lines directly above ``key`` inside the top-level ``block`` of a YAML file.

    A key with no comment of its own, directly below another key of the block, shares
    the comment of that key, as one comment documents a run of keys.
    """

    lines = path.read_text(encoding="utf-8").splitlines()
    start = next((i for i, line in enumerate(lines) if re.match(rf"^{re.escape(block)}:\s*(#.*)?$", line)), None)
    if start is None:
        return ""
    keys: dict[str, int] = {}
    indent = None
    for i in range(start + 1, len(lines)):
        line = lines[i]
        if line and not line[0].isspace() and not line.startswith("#"):
            break
        match = re.match(r"^(\s+)([A-Za-z0-9_]+):", line)
        if match and (indent is None or len(match.group(1)) == indent):
            indent = len(match.group(1))
            keys[match.group(2)] = i
    if key not in keys:
        return ""
    line_no = keys[key]
    comments: list[str] = []
    i = line_no - 1
    while i > start and lines[i].strip().startswith("#"):
        comments.insert(0, lines[i].strip()[1:].strip())
        i -= 1
    if comments:
        return "\n".join(comments)
    above = next((name for name, number in keys.items() if number == i), None)
    return yaml_comments(path, block, above) if above and i < line_no else ""


_WORKBOOKS: dict[tuple[str, str, tuple[str, ...]], dict] = {}


def workbook_records(path: Path, sheet: str, columns: list[str]) -> dict[tuple[str, str, str], str]:
    """Return the record of each (scenario, technology, parameter) row: the first non-empty preferred column."""

    cache_key = (str(path), sheet, tuple(columns))
    if cache_key in _WORKBOOKS:
        return _WORKBOOKS[cache_key]
    from openpyxl import load_workbook

    records: dict[tuple[str, str, str], str] = {}
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet in workbook.sheetnames:
            rows = workbook[sheet].iter_rows(values_only=True)
            header = [str(cell) if cell is not None else "" for cell in next(rows, [])]
            wanted = [header.index(name) for name in columns if name in header]
            if wanted and {"scenario", "Tech", "Parameter"} <= set(header):
                s, t, p = header.index("scenario"), header.index("Tech"), header.index("Parameter")
                for row in rows:
                    text = next((str(row[i]) for i in wanted if row[i] not in (None, "")), "")
                    records.setdefault((str(row[s]), str(row[t]), str(row[p])), text)
    finally:
        workbook.close()
    _WORKBOOKS[cache_key] = records
    return records


def _numbers(text: str) -> set[float]:
    return {float(n.replace(",", "")) for n in re.findall(r"\d[\d,]*(?:\.\d+)?", text)}


def _sections(text: str) -> list[tuple[int, str, str]]:
    """Return (level, heading, body) for every heading of a Markdown file."""

    parts = re.split(r"^(#{1,6}) (.+)$", text, flags=re.MULTILINE)
    return [(len(parts[i]), plain(parts[i + 1]).strip(), parts[i + 2]) for i in range(1, len(parts) - 2, 3)]


def _section(text: str, prefix: str) -> str:
    """Return a section and its subsections, by the start of its heading."""

    sections = _sections(text)
    for index, (level, heading, body) in enumerate(sections):
        if heading.startswith(prefix):
            out = [body]
            for later_level, later_heading, later_body in sections[index + 1:]:
                if later_level <= level:
                    break
                out.append(f"\n{'#' * later_level} {later_heading}\n{later_body}")
            return "".join(out)
    return ""


def _blocks(text: str) -> list[str]:
    """Return the bullets and paragraphs of Markdown text, each joined into one line."""

    out: list[str] = []
    current: list[str] = []
    for line in text.splitlines():
        if not line.strip() or line.startswith(("#", "|", "---")):
            if current:
                out.append(one_line(" ".join(current)))
            current = []
            continue
        if line.lstrip().startswith("- ") and not line.startswith(" "):
            if current:
                out.append(one_line(" ".join(current)))
            current = [line.strip()[2:]]
        else:
            current.append(line.strip())
    if current:
        out.append(one_line(" ".join(current)))
    return out


def _vre_note(text: str, name: str, node: str, label: str, tech_word: str) -> dict | None:
    """Return the ASSUMPTION sentence of the note or section ``name`` that names the node and technology."""

    name = plain(name).strip()
    candidates: list[str] = []
    for block in _blocks(text):
        if plain(block).lower().startswith(f"**{name}.**".lower()) or plain(block).lower().startswith(f"{name}.".lower()):
            candidates.append(block)
    if not candidates:
        for _, heading, body in _sections(text):
            if name.lower() in heading.lower():
                candidates += _blocks(body)
    found = []
    for block in candidates:
        _, assumptions = parse_record(block)
        if not assumptions:
            continue
        sentence = " ".join(assumptions)
        mentions_node = re.search(rf"(?<![A-Za-z]){re.escape(node)}(?![A-Za-z])", sentence) or (label and label in sentence)
        if mentions_node and tech_word.lower() in sentence.lower():
            found.append(assumptions)
    if len(found) != 1:
        return None
    return {"kind": "assumption", "assumptions": found[0]}


# The source map --------------------------------------------------------------------


class SourceMap:
    """The locations of ``content/sources.yaml`` and the entries read from them."""

    def __init__(self, root: Path, commit: str | None, repository_url: str, path: Path = MAP_FILE):
        self.root = Path(root)
        self.commit = commit
        self.repository_url = repository_url
        self.data = yaml_io.load_file(path)

    def blob(self, path: str) -> str:
        return f"{self.repository_url}/blob/{self.commit or 'main'}/{path}"

    def rule_of(self, name: str) -> tuple[str, Any]:
        """Return how the entry of the value ``name`` (a key without its setup id) is found."""

        for computed, inputs in (self.data.get("computed") or {}).items():
            if fnmatchcase(name, computed):
                if computed.endswith(".*"):
                    # A value per node: each input is the input of the same node.
                    node = name[len(computed) - 1:]
                    inputs = [n[:-1] + node if n.endswith(".*") else n for n in inputs]
                return "computed", inputs
        for spec in self.data.get("records") or []:
            if any(fnmatchcase(name, pattern) for pattern in spec["keys"]):
                return "record", spec
        for spec in self.data.get("settings") or []:
            if any(fnmatchcase(name, pattern) for pattern in spec["keys"]):
                return "setting", spec["label"]
        raise KeyError(f"no entry in {MAP_FILE.name} for site value {name!r}")

    def record_entry(self, spec: dict, name: str, value: Any, info) -> dict | None:
        """Return the entry of one value from the record locations of ``spec``."""

        if spec.get("none"):
            return None
        if "markdown_table" in spec:
            entry = self._table_entry(spec["markdown_table"], name, value, info)
            if entry:
                return entry
        if "yaml_comments" in spec:
            entry = self._comment_entry(spec["yaml_comments"], name, info)
            if entry:
                return entry
        if "workbook" in spec:
            entry = self._workbook_entry(spec["workbook"], name, info)
            if entry or "value_sources" not in spec:
                return entry
        if "value_sources" in spec:
            return self._value_sources_entry(spec["value_sources"], name, value, info)
        if "vre_table" in spec:
            return self._vre_entry(spec["vre_table"], name, value, info)
        return None

    def _file(self, role: str, info) -> str | None:
        return info.manifest if role == "manifest" else info.sources.get(role)

    def _table_entry(self, spec: dict, name: str, value: Any, info) -> dict | None:
        if info.sources.get("country_config") != spec["documents"]:
            return None
        key = name.split(".")[1]
        header = list(spec["header"])
        rows = markdown_table(self.root / spec["file"], header)
        pool = [c for row in rows for c in parse_record(row[header[2]])[0] if c["urls"]]
        for row in rows:
            if key in re.findall(r"`([^`]+)`", row[header[0]]):
                if float(value) not in _numbers(row[header[1]]):
                    raise ValueError(f"{spec['file']} gives {row[header[1]]!r} for {key}, the configuration {value!r}")
                entry = entry_from_record(row[header[2]], pool)
                if entry:
                    entry["record"] = self.blob(spec["file"])
                return entry
        return None

    def _comment_entry(self, spec: dict, name: str, info) -> dict | None:
        path = self._file(spec["role"], info)
        if path is None:
            return None
        key = spec.get("key") or name.split(".")[1]
        text = yaml_comments(self.root / path, spec["block"], key)
        entry = entry_from_record(text) if text else None
        if entry:
            entry["record"] = self.blob(path)
        return entry

    def _workbook_entry(self, spec: dict, name: str, info) -> dict | None:
        path = self._file(spec["role"], info)
        if path is None or not (self.root / path).is_file():
            return None
        tech = spec.get("technology") or name.split(".")[1]
        records = workbook_records(self.root / path, spec["sheet"], list(self.data.get("workbook_columns") or []))
        entry = entry_from_record(records.get((spec["scenario"], tech, spec["parameter"]), ""))
        if entry:
            entry["record"] = self.blob(path)
        return entry

    def _vre_entry(self, spec: dict, name: str, value: Any, info) -> dict | None:
        if info.manifest != spec["documents"]:
            return None
        tech = name.split(".", 1)[1]
        text = (self.root / spec["file"]).read_text(encoding="utf-8")
        table = _section(text, spec["table_section"])
        header = list(spec["header"])
        prefixes = {prefix: word for word, prefix in spec["technologies"].items()}
        word = prefixes.get(tech[:6])
        node = tech[6:]
        short = node[:-2] if node.endswith("XX") else node
        row = next(
            (r for r in markdown_table_text(table, header) if r[header[0]] == short and r[header[1]] == word),
            None,
        )
        if row is None:
            return None
        if float(value) not in _numbers(row[header[2]]):
            raise ValueError(f"{spec['file']} gives {row[header[2]]!r} for {tech}, the manifest {value!r}")
        cell = plain(row[header[3]]).strip()
        see = _SEE.match(cell)
        if see:
            label = info.node_labels.get(node, "")
            entry = _vre_note(text, see.group(1), short, label, word)
        else:
            pool = [
                c for block in _blocks(_section(text, spec["reference_section"]))
                for part in re.split(r"\s;\s", re.sub(r"^\*\*[^*]+\*\*:\s*", "", block))
                for c in [{"text": plain(part).strip(), "urls": urls_of(plain(part))}]
            ]
            entry = entry_from_record(cell, pool)
        if entry:
            entry["record"] = self.blob(spec["file"])
        return entry

    def _value_sources_entry(self, spec: dict, name: str, value: Any, info) -> dict | None:
        """Return the entry of the block for the technology, parameter and first model year, if its value is ``value``."""

        if info.manifest != spec["documents"] or not info.years:
            return None
        tech = name.split(".")[1]
        text = (self.root / spec["file"]).read_text(encoding="utf-8")
        blocks = [
            fields for techs, parameter, years, fields in value_blocks(text)
            if tech in techs and parameter == spec["parameter"] and years[0] <= info.years[0] <= years[-1]
        ]
        if len(blocks) != 1:
            return None
        fields = blocks[0]
        shown = re.fullmatch(r"\s*(-?\d[\d,]*(?:\.\d+)?)\s*", fields.get("Values", "").split(";")[0])
        unit = (fields.get("Values", "").split(";") + ["", ""])[1].strip()
        if shown is None or float(shown.group(1).replace(",", "")) != float(value) or unit != spec["unit"]:
            return None
        documents = value_documents(text)
        citations = []
        for label, part in re.findall(r"\[(R\d+)\]\s*(.*?)\s*(?=\[R\d+\]|$)", fields.get("Sources", "")):
            if label not in documents:
                return None
            where = part.strip().rstrip(";.").strip()
            cited = documents[label] + (f" Part used: {where}." if where else "")
            citations.append({"text": cited, "urls": urls_of(documents[label])})
        statements = value_assumptions(text)
        labels = re.match(r"\s*(A\d+(?:\s*,\s*A\d+)*)", fields.get("Assumptions", fields.get("Assumption", "")))
        assumptions = [statements.get(n) for n in re.findall(r"A\d+", labels.group(1))] if labels else []
        if not any(c["urls"] for c in citations) or any(a is None for a in assumptions):
            return None
        return {"kind": "external", "citations": citations, "assumptions": assumptions, "record": self.blob(spec["file"])}


def value_blocks(text: str) -> list[tuple[list[str], str, list[int], dict[str, str]]]:
    """Return the technologies, parameter, years and fields of each block of a base-year value record.

    A block is a level-3 heading such as ``### `TECH` Parameter, 2023 to 2024`` followed by
    bullets ``- Field: text``.
    """

    found = []
    parts = re.split(r"^### (.+)$", text, flags=re.MULTILINE)
    for heading, body in zip(parts[1::2], parts[2::2]):
        match = re.fullmatch(r"((?:`[^`]+`,?\s*)+)(\w+),\s*(\d{4})(?:\s+to\s+(\d{4}))?", heading.strip())
        if not match:
            continue
        fields = {}
        for block in _blocks(re.split(r"^#", body, maxsplit=1, flags=re.MULTILINE)[0]):
            field = re.match(r"^([A-Z][A-Za-z ]*):\s*(.*)$", block)
            if field:
                fields[field.group(1)] = field.group(2)
        years = [int(match.group(3)), int(match.group(4) or match.group(3))]
        found.append((re.findall(r"`([^`]+)`", match.group(1)), match.group(2), years, fields))
    return found


def value_documents(text: str) -> dict[str, str]:
    """Return the entry of each document ``Rn`` listed under the heading "Documents", in plain text."""

    documents = {}
    for block in _blocks(_section(text, "Documents")):
        match = re.match(r"^\*\*(R\d+)\*\*\s+(.+)$", block)
        if match:
            documents[match.group(1)] = plain(match.group(2)).strip()
    return documents


def value_assumptions(text: str) -> dict[str, str]:
    """Return each assumption ``An`` of a base-year value record with its label, as the record gives it."""

    statements = {}
    for block in _blocks(_section(text, "Assumptions")):
        match = re.match(r"^\*\*(A\d+)\.\*\*\s+(.+)$", block)
        if match:
            statements[match.group(1)] = f"{match.group(1)}. {plain(match.group(2)).strip()}"
    return statements


def markdown_table_text(text: str, header: list[str]) -> list[dict[str, str]]:
    """Return the rows of the table with ``header`` inside a piece of Markdown text."""

    rows: list[dict[str, str]] = []
    lines = text.splitlines()
    for index, line in enumerate(lines):
        if _cells(line) == header:
            for row in lines[index + 2:]:
                if not row.strip().startswith("|"):
                    break
                rows.append(dict(zip(header, _cells(row))))
            break
    return rows


def resolve(site) -> None:
    """Give every value of the site its entry, and leave out the values with no record."""

    sources = SourceMap(site.paths.project_root, site.commit, site.config["repository_url"])
    site.entries = {}
    site.left_out = {}
    for setup_id, info in site.setups.items():
        done: dict[str, dict | None] = {}

        def entry_of(key: str) -> dict | None:
            if key in done:
                return done[key]
            done[key] = None
            item = info.values.get(key)
            if item is None:
                return None
            name = key.split(".", 1)[1]
            rule, spec = sources.rule_of(name)
            if rule == "computed":
                inputs = [f"{setup_id}.{n}" for n in spec]
                entry = (
                    {"kind": "computed", "formula": re.sub(r"^computed:\s*", "", item.note), "inputs": inputs}
                    if all(entry_of(k) for k in inputs) else None
                )
            elif rule == "record":
                entry = sources.record_entry(spec, name, item.value, info)
            else:
                entry = {"kind": "setting", "label": spec, "path": item.source, "url": sources.blob(item.source)}
            done[key] = entry
            return entry

        for key in list(info.values):
            entry = entry_of(key)
            if entry is None:
                site.left_out[key] = info.values.pop(key)
            else:
                site.entries[key] = entry
    for key, item in site.repo.items():
        rule, spec = sources.rule_of(key)
        if rule != "setting":
            raise ValueError(f"repository value {key} must be a setting in {MAP_FILE.name}")
        site.entries[key] = {"kind": "setting", "label": spec, "path": item.source, "url": sources.blob(item.source)}


# How an entry is shown ---------------------------------------------------------------


def units(site, key: str) -> list[tuple]:
    """Return the distinct sources behind one value, in order: citations, sentences, settings."""

    entry = site.entries[key]
    if entry["kind"] == "computed":
        found: list[tuple] = []
        for name in entry["inputs"]:
            for unit in units(site, name):
                if unit not in found:
                    found.append(unit)
        return found
    if entry["kind"] == "setting":
        return [("setting", entry["label"], entry["path"], entry["url"])]
    found = [("citation", c["text"], tuple(c["urls"])) for c in entry.get("citations", [])]
    return found + [("assumption", sentence) for sentence in entry.get("assumptions", [])]


def _short(text: str) -> str:
    cut = re.split(r"\s\[Online\]|https?://", text, maxsplit=1)[0]
    return cut.strip().rstrip(" ,;:(")


def one_line_entry(site, key: str) -> str:
    """Return the entry of a value as one line of plain text, for the hover."""

    entry = site.entries[key]
    if entry["kind"] == "setting":
        return f"{entry['label'][0].upper()}{entry['label'][1:]}: {entry['path']}"
    if entry["kind"] == "computed":
        names = ", ".join(name.split(".", 1)[1] for name in entry["inputs"])
        return f"Computed: {entry['formula']}, from {names}"
    parts = []
    if entry.get("citations"):
        parts.append("Source: " + "; ".join(_short(c["text"]) for c in entry["citations"]))
    parts += [f"assumption: {sentence}" for sentence in entry.get("assumptions", [])]
    return " | ".join(parts)


__all__ = ["git_commit", "parse_record", "resolve", "units", "one_line_entry", "plain", "urls_of"]
